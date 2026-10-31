"""Начальная загрузка фонда — «заменить всё» (спецификация stays-fund-versions §9.3).

Переход с тестовых данных на реальные: все данные шахматки (фонд с версиями, брони, гости,
покрытие батлеров, OOO и блоки) удаляются, из файла фонда заводится исходная версия.
Только пока выключен переключатель «Шахматка в эксплуатации» (StaysSettings.in_production).
Перед удалением — полная копия данных модуля в Excel (§9.3; листы — Plan, Decision 4, плюс
по листу на каждую удаляемую таблицу). Настройки модуля
(KEEP) не трогаются. Ничего не коммитит — commit в роутере: удаление и загрузка — одна
транзакция."""
from __future__ import annotations

import datetime as dt
import io
import json
import os
import re
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from sqlalchemy import delete, func, select, update
from sqlalchemy.orm import Session

from app.core.audit import audit
from app.core.format import safe_sheet_cell
from app.core.media import BASE_DIR
from app.core.time import local_date, now_local

from ..adapters.orm import (
    Assignment,
    Category,
    CategoryKind,
    CategoryState,
    Companion,
    CoverageGuest,
    CoverageLine,
    CustomField,
    Escort,
    FundGroup,
    FundVersion,
    GroupState,
    Guest,
    GuestLink,
    Hold,
    Label,
    Logistics,
    Move,
    Occupant,
    Reservation,
    StayGroup,
    StaysSettings,
    Unit,
    UnitFlag,
    UnitLineage,
    UnitState,
    UnitStateFlag,
)
from . import allocation, fund_drafts
from .fund import FundConflict, FundError, get_settings
from .fund_drafts import _COLOR, _WEIGHT_MAX, FundFileError, _file_weight, unknown_kind_message
from .fund_versions import INITIAL_DATE, fund_on
from .fund_xlsx import (
    HEAD,
    PATH_SEP,
    SHEET_CATEGORIES,
    SHEET_FLAGS,
    SHEET_TREE,
    SHEET_UNITS,
    FileError,
    ParsedFund,
    export_fund,
)
from .locking import lock_module

CONFIRM_WORD = "ЗАМЕНИТЬ"
IDS_IGNORED = "ID в файле будут проигнорированы"
SHEET_RESERVATIONS = "Брони"
SHEET_GUESTS = "Гости"

# рядом с data/uploads (Plan, Decision 4); имя копии — строго по шаблону: его же проверяет скачивание
BACKUP_DIR = BASE_DIR / "data" / "backups" / "stays"
BACKUP_NAME = re.compile(r"stays-\d{8}-\d{6}\.xlsx")

# Таблицы stays_*: каждая — либо данные шахматки (удаляются, в порядке внешних ключей: сначала
# ссылающиеся), либо настройки модуля (остаются). Тест сверяет списки со всей схемой модуля.
WIPE = (CoverageGuest, CoverageLine, Occupant, Companion, Logistics, Move, Assignment, Reservation,
        StayGroup, GuestLink, Guest, Hold, UnitLineage, UnitStateFlag, UnitState, GroupState,
        CategoryState, FundVersion, UnitFlag, Unit, Category, FundGroup)
KEEP = (StaysSettings, Label, CustomField, Escort, CategoryKind)

# что будет удалено (предпросмотр) — тот же словарь служит отпечатком для 409 при повторе
COUNTED = {"units": Unit, "categories": Category, "groups": FundGroup, "flags": UnitFlag,
           "versions": FundVersion, "reservations": Reservation, "guests": Guest,
           "coverage": CoverageLine, "holds": Hold}


def reset_counts(db: Session) -> dict[str, int]:
    return {k: int(db.scalar(select(func.count()).select_from(m)) or 0) for k, m in COUNTED.items()}


def reset_state(db: Session) -> dict:
    """Отпечаток данных для подтверждения: счётчики и монотонные метки — наибольшие id таблиц
    шахматки и время последней версии фонда. Счётчики одни не отличили бы повторную загрузку
    файла того же размера; id и время после каждой замены только растут."""
    max_ids = {m.__tablename__: db.scalar(select(func.max(m.id)))
               for m in WIPE if "id" in m.__table__.c and m.__table__.c.id.primary_key}
    last = db.execute(select(func.max(FundVersion.created_at), func.max(FundVersion.published_at))).one()
    return {"counts": reset_counts(db), "max_ids": max_ids,
            "versions_at": [x.isoformat() if x else None for x in last]}


def fingerprint(state: dict) -> str:
    return json.dumps(state, sort_keys=True, separators=(",", ":"), default=str)


def parse_fingerprint(s: str) -> dict:
    try:
        v = json.loads(s)
    except ValueError:
        v = None
    if not isinstance(v, dict):
        raise FundError("Нет отпечатка предпросмотра — выполните предпросмотр заново")
    return v


# ── файл → данные исходной версии ──────────────────────────────────────────


def plan_load(parsed: ParsedFund, kinds: list[CategoryKind]) -> tuple[dict, list[FileError]]:
    """Разобранный файл → аргументы fund_drafts.seed_initial. ID и «Из чего» игнорируются.
    seed_initial не проверяет входные данные (демо-сид), поэтому всё, что упало бы на
    KeyError или ограничении БД, — здесь ошибкой файла с листом и строкой."""
    errors: list[FileError] = []

    def err(sheet: str, row: int, msg: str) -> None:
        errors.append(FileError(sheet, row, msg))

    def dup(sheet: str, seen: dict, key, row: int, what: str) -> bool:
        if key in seen:
            err(sheet, row, f"{what} повторяется (первый раз в строке {seen[key]})")
            return True
        seen[key] = row
        return False

    by_title = {k.title.casefold(): k.id for k in kinds if not k.archived}
    cats: list[dict] = []
    codes: dict[str, int] = {}
    for c in parsed.categories:
        if not c.code or not c.title:
            err(SHEET_CATEGORIES, c.row_no, "не указан код или название категории")
        dup(SHEET_CATEGORIES, codes, c.code, c.row_no, f"код «{c.code}»")
        if len(c.code) > 16:
            err(SHEET_CATEGORIES, c.row_no, "код категории длиннее 16 символов")
        if len(c.title) > 120:
            err(SHEET_CATEGORIES, c.row_no, "название категории длиннее 120 символов")
        kind = by_title.get(c.kind.strip().casefold())
        if not c.kind.strip():
            err(SHEET_CATEGORIES, c.row_no, "не указан вид категории")
        elif kind is None:
            err(SHEET_CATEGORIES, c.row_no, unknown_kind_message(c.kind, kinds))
        weight = _file_weight(c.weight)
        if weight is None:
            err(SHEET_CATEGORIES, c.row_no, f"вес {c.weight} — ожидается число больше 0 и не больше {_WEIGHT_MAX}")
        cats.append({"code": c.code, "title": c.title, "kind_id": kind, "escorted": c.escorted,
                     "weight": weight, "sort_order": c.sort_order})

    groups: list[dict] = []
    paths: dict[tuple[str, ...], int] = {}
    for g in sorted(parsed.groups, key=lambda g: len(g.path)):          # родитель раньше потомка
        if not g.path or any(not p or "/" in p for p in g.path):
            err(SHEET_TREE, g.row_no, "путь узла пуст, содержит пустое звено или «/» в названии")
            continue
        if dup(SHEET_TREE, paths, g.path, g.row_no, f"путь «{PATH_SEP.join(g.path)}»"):
            continue
        if len(g.path[-1]) > 120:
            err(SHEET_TREE, g.row_no, "название узла длиннее 120 символов")
        if len(g.path) > 1 and g.path[:-1] not in paths:
            err(SHEET_TREE, g.row_no, f"нет родителя «{PATH_SEP.join(g.path[:-1])}»")
        groups.append({"path": list(g.path), "sort_order": g.sort_order})

    flags: list[dict] = []
    titles: dict[str, int] = {}
    for i, f in enumerate(parsed.flags):
        sheet, row = (SHEET_FLAGS, f.row_no) if f.row_no else (SHEET_UNITS, 1)
        if not f.title or dup(sheet, titles, f.title, row, f"флаг «{f.title}»"):
            continue
        if len(f.title) > 80:
            err(sheet, row, f"название флага «{f.title[:20]}…» длиннее 80 символов")
        if f.color and not _COLOR.fullmatch(f.color):
            err(sheet, row, f"цвет «{f.color}» — ожидается #RRGGBB")
        flags.append({"title": f.title, "sort_order": i * 10, **({"color": f.color} if f.color else {})})

    units: list[dict] = []
    labels: dict[str, int] = {}
    for u in parsed.units:
        if not u.label:
            err(SHEET_UNITS, u.row_no, "не указан номер объекта")
        dup(SHEET_UNITS, labels, u.label, u.row_no, f"номер «{u.label}»")
        if len(u.label) > 40:
            err(SHEET_UNITS, u.row_no, "номер объекта длиннее 40 символов")
        if len(u.floor) > 8:
            err(SHEET_UNITS, u.row_no, "этаж длиннее 8 символов")
        if u.category_code not in codes:
            err(SHEET_UNITS, u.row_no, f"неизвестная категория «{u.category_code}»")
        if u.group_path is not None and tuple(u.group_path) not in paths:
            err(SHEET_UNITS, u.row_no, f"неизвестный узел «{PATH_SEP.join(u.group_path)}»")
        for t in sorted(u.flags):
            if t not in titles:
                err(SHEET_UNITS, u.row_no, f"неизвестный флаг «{t}»")
        units.append({"label": u.label, "category": u.category_code,
                      "path": list(u.group_path) if u.group_path else None,
                      "floor": u.floor, "flags": sorted(u.flags)})
    return {"categories": cats, "groups": groups, "flags": flags, "units": units}, errors


def _will_load(plan: dict) -> dict[str, int]:
    return {k: len(v) for k, v in plan.items()}


def _warnings(parsed: ParsedFund) -> list[str]:
    rows = [*parsed.categories, *parsed.groups, *parsed.flags, *parsed.units]
    if any(r.id is not None for r in rows) or any(u.predecessors for u in parsed.units):
        return [IDS_IGNORED]
    return []


def _not_in_production(db: Session) -> None:
    if get_settings(db).in_production:
        raise FundConflict("Шахматка в эксплуатации — начальная загрузка недоступна")


def preview(db: Session, parsed: ParsedFund) -> dict:
    """Что удалит и загрузит начальная загрузка; fingerprint — для подтверждения."""
    _not_in_production(db)
    plan, errors = plan_load(parsed, list(db.scalars(select(CategoryKind))))
    if errors:
        raise FundFileError(errors)
    state = reset_state(db)
    return {"counts": state["counts"], "will_load": _will_load(plan), "warnings": _warnings(parsed),
            "fingerprint": fingerprint(state)}


# ── копия ──────────────────────────────────────────────────────────────────


def _latest(db: Session, col, key, value) -> dict:
    """id → значение из последней (по дате) опубликованной версии, где оно есть."""
    out: dict = {}
    rows = db.execute(select(key, value).join(FundVersion, FundVersion.id == col)
                      .where(FundVersion.status == "published").order_by(FundVersion.valid_from))
    for k, v in rows:
        out[k] = v
    return out


def _sheet(wb, title: str, head: list[str], rows) -> None:
    ws = wb.create_sheet(title)
    for i, h in enumerate(head, start=1):
        ws.cell(row=1, column=i, value=h).font = HEAD
    for r, values in enumerate(rows, start=2):
        for i, v in enumerate(values, start=1):
            if isinstance(v, str):          # управляющие символы в заметках гостя openpyxl не пишет
                v = ILLEGAL_CHARACTERS_RE.sub("", v)
            ws.cell(row=r, column=i, value=safe_sheet_cell(v))
    ws.freeze_panes = "A2"


def _guest_name(g: Guest) -> str:
    return " ".join(x for x in (g.last_name, g.first_name, g.middle_name) if x)


def backup_bytes(db: Session) -> bytes:
    """Копия данных модуля: фонд на сегодня листами файла фонда (их же принимает загрузка),
    брони и гости — листами «Брони», «Гости»; затем по листу на каждую таблицу WIPE со всеми
    строками (_raw_sheets) — полная копия того, что удаляется (§9.3)."""
    wb = load_workbook(io.BytesIO(export_fund(fund_on(db, local_date()), list(db.scalars(select(UnitFlag))))))
    code = _latest(db, CategoryState.version_id, CategoryState.category_id, CategoryState.code)
    label = _latest(db, UnitState.version_id, UnitState.unit_id, UnitState.label)
    labels = dict(db.execute(select(Label.id, Label.title)).all())
    guests = {g.id: g for g in db.scalars(select(Guest).order_by(Guest.id))}
    units: dict[int, list[str]] = {}
    for rid, uid in db.execute(select(Assignment.reservation_id, Assignment.unit_id)
                               .order_by(Assignment.reservation_id, Assignment.seg_index)):
        if uid is not None:
            units.setdefault(rid, [])
            if (lb := label.get(uid, f"#{uid}")) not in units[rid]:
                units[rid].append(lb)
    occ: dict[int, list[str]] = {}
    for rid, gid, role in db.execute(select(Occupant.reservation_id, Occupant.guest_id, Occupant.role)
                                     .order_by(Occupant.reservation_id, Occupant.id)):
        name = _guest_name(guests[gid]) if gid in guests else f"#{gid}"
        occ.setdefault(rid, []).append(f"{name} (главный)" if role == "main" else name)
    yn = {True: "да", False: ""}
    _sheet(wb, SHEET_RESERVATIONS,
           ["ID", "номер", "статус", "заезд", "выезд", "РЗ", "ПВ", "категория (код)", "объекты", "гости",
            "метка", "заказчик", "оплата за (код)", "примечание", "без батлера", "источник", "внешний ID",
            "поля"],
           ([r.id, r.number, r.status, r.check_in, r.check_out, yn[bool(r.early)], yn[bool(r.late)],
             code.get(r.category_id, f"#{r.category_id}"), ", ".join(units.get(r.id, [])),
             "; ".join(occ.get(r.id, [])), labels.get(r.label_id, ""), r.customer, code.get(r.paid_category_id or r.category_id, ""),
             r.note, yn[bool(r.declined_butler)], r.source, r.external_id or "", r.fields_json]
            for r in db.scalars(select(Reservation).order_by(Reservation.check_in, Reservation.id))))
    _sheet(wb, SHEET_GUESTS,
           ["ID", "фамилия", "имя", "отчество", "телефон", "email", "дата рождения", "уровень",
            "предпочтения", "особые даты", "поля", "примечание", "источник", "внешний ID"],
           ([g.id, g.last_name, g.first_name, g.middle_name, g.phone, g.email, g.birth_date, g.level,
             g.preferences_json, g.special_dates_json, g.fields_json, g.note, g.source, g.external_id or ""]
            for g in guests.values()))
    _raw_sheets(wb, db)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _raw(v):
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, dt.date | dt.datetime):
        return v.isoformat()
    return v


def raw_sheet_names() -> dict[str, str]:
    """Таблица → имя листа сырой копии: имя таблицы, не длиннее 31 символа (предел Excel), без повторов."""
    out: dict[str, str] = {}
    for m in WIPE:
        name, n = m.__tablename__[:31], 1
        while name in out.values():
            n += 1
            name = f"{m.__tablename__[:31 - len(str(n)) - 1]}~{n}"
        out[m.__tablename__] = name
    return out


def _raw_sheets(wb, db: Session) -> None:
    """По листу на каждую удаляемую таблицу (WIPE): заголовок — колонки, строка — запись БД.
    Человекочитаемые листы выше — для просмотра; эти — полная копия всего, что удаляется
    (будущие версии и черновик фонда, блоки и OOO, покрытие, переезды, логистика…)."""
    names = raw_sheet_names()
    for m in WIPE:
        t = m.__table__
        cols = [c.name for c in t.columns]
        order = list(t.primary_key.columns)
        _sheet(wb, names[t.name], cols,
               ([_raw(v) for v in row] for row in db.execute(select(*t.columns).order_by(*order))))


class BackupError(FundConflict):
    """Копию записать не удалось (диск, права): замена не начиналась."""


def backup_module(db: Session) -> Path:
    """Пишет копию в BACKUP_DIR: сначала во временный файл, fsync, затем переименование —
    файл с именем по шаблону всегда полный. Имя уже занято (две замены в одну секунду) — 409."""
    data = backup_bytes(db)
    final = BACKUP_DIR / f"stays-{now_local():%Y%m%d-%H%M%S}.xlsx"
    tmp = final.with_name(final.name + ".part")
    try:
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        if final.exists():
            raise FundConflict("Копия с таким временем уже есть — повторите через секунду")
        with open(tmp, "wb") as f:
            f.write(data)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp, final)
    except OSError as e:
        tmp.unlink(missing_ok=True)
        raise BackupError(f"Не удалось сохранить копию — ничего не удалено ({e.strerror or e})") from e
    return final


def backup_path(name: str) -> Path | None:
    """Файл копии по имени: только имя по шаблону, без путей; иначе None."""
    if not BACKUP_NAME.fullmatch(name):
        return None
    root = BACKUP_DIR.resolve()
    path = (root / name).resolve()
    if path.parent != root or not path.is_file():
        return None
    return path


# ── замена ─────────────────────────────────────────────────────────────────


def wipe(db: Session) -> None:
    """Удаляет данные шахматки (WIPE) в порядке внешних ключей; KEEP не трогает."""
    db.execute(update(GroupState).values(parent_id=None))           # ссылки внутри таблицы
    db.execute(update(FundVersion).values(base_version_id=None))
    for model in WIPE:
        db.execute(delete(model))
    db.flush()
    db.expire_all()


# Сущности, чьи ID пишутся в файл фонда. SQLite (без AUTOINCREMENT) после удаления всех строк
# начал бы id заново с 1, и старый файл, загруженный потом черновиком, совпал бы ID с чужими
# объектами. Заглушка с прежним наибольшим id на время загрузки — новые id идут после старых.
_FILE_IDS = (Unit, Category, FundGroup, UnitFlag)
_PLACEHOLDER = "\u2063начальная загрузка"


def _hold_ids(db: Session, maxima: dict) -> list:
    rows = [m(id=n, title=_PLACEHOLDER) if m is UnitFlag else m(id=n) for m, n in maxima.items() if n]
    db.add_all(rows)
    db.flush()
    return rows


def _release_ids(db: Session, rows: list) -> None:
    for r in rows:
        db.delete(r)
    db.flush()


def initial_load(db: Session, parsed: ParsedFund, principal, expected_state: dict, *,
                 confirm: str, filename: str = "") -> dict:
    """«Начальная загрузка (заменить всё)». Под замком модуля: 409 — в эксплуатации или данные
    изменились после предпросмотра (expected_state ≠ reset_state); 422 — нет слова
    подтверждения или ошибка в файле. Затем копия, удаление, исходная версия из файла,
    распределение (броней нет — шахматка пустая) и аудит stays_reset."""
    lock_module(db)
    _not_in_production(db)
    if confirm != CONFIRM_WORD:
        raise FundError(f"Для замены введите слово «{CONFIRM_WORD}»")
    state = reset_state(db)
    if json.loads(fingerprint(state)) != expected_state:
        raise FundConflict("Данные шахматки изменились после предпросмотра — повторите предпросмотр")
    plan, errors = plan_load(parsed, list(db.scalars(select(CategoryKind))))
    if errors:
        raise FundFileError(errors)
    counts = state["counts"]
    path = backup_module(db)
    maxima = {m: db.scalar(select(func.max(m.id))) for m in _FILE_IDS}
    wipe(db)
    held = _hold_ids(db, maxima)
    now = now_local()
    db.add(FundVersion(valid_from=INITIAL_DATE, status="published", rev=1, note="Начальная загрузка",
                       author_id=principal.user.id if principal else None, created_at=now, published_at=now))
    db.flush()
    fund_drafts.seed_initial(db, **plan)
    _release_ids(db, held)
    allocation.run_allocation(db)
    loaded = _will_load(plan)
    audit(db, principal, "stays_reset", "stays", {"deleted": counts, "loaded": loaded,
                                                  "backup": path.name, "filename": filename})
    return {"deleted": counts, "loaded": loaded, "backup": path.name}


def set_in_production(db: Session, principal, value: bool) -> bool:
    lock_module(db)
    row = get_settings(db)
    before = bool(row.in_production)
    row.in_production = value
    db.flush()
    audit(db, principal, "stays_in_production", "stays_settings:in_production", {"from": before, "to": value})
    return value
