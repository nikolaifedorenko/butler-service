"""Файл фонда (§9): выгрузка снимка в Excel и разбор такого файла обратно. Чистый модуль —
БД не трогает: применение разобранного файла к черновику — отдельный шаг. Разбор никогда не
бросает исключение на плохом входе: всё уходит в список FileError (лист, строка, сообщение)."""
from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.format import safe_sheet_cell

from ..adapters.orm import UnitFlag
from .fund_versions import FundSnapshot

SHEET_CATEGORIES = "Категории"
SHEET_TREE = "Дерево"
SHEET_FLAGS = "Флаги"
SHEET_UNITS = "Объекты"
SHEET_LISTS = "Списки"      # скрытый: все пути дерева для выпадающих списков

PATH_SEP = " / "
ID_HEADER = "ID (не менять; пусто — новая запись)"
MAX_FILE_BYTES = 10 * 1024 * 1024          # сам файл
MAX_UNPACKED_BYTES = 100 * 1024 * 1024     # распакованное содержимое (защита от zip-бомбы)
MAX_ROWS = 20000                           # строк на лист
MAX_FLAG_COLUMNS = 200

HEAD = Font(bold=True)
GREY = PatternFill("solid", fgColor="D9D9D9")
_TRUE = {"да", "1", "true", "yes"}
_FALSE = {"нет", "0", "false", "no"}
_UNIT_FIXED = ["номер", "категория (код)", "узел", "этаж", "из чего (ID через запятую)"]


# ---------------------------------------------------------------- результат разбора

@dataclass(frozen=True)
class FileError:
    sheet: str
    row: int          # 0 — ошибка листа/файла целиком
    message: str


@dataclass
class ParsedCategory:
    row_no: int
    id: int | None
    code: str
    title: str
    kind: str             # название вида (справочник видов); пусто — не указан
    weight: float
    escorted: bool
    sort_order: int


@dataclass
class ParsedGroup:
    row_no: int
    id: int | None
    path: tuple[str, ...]
    sort_order: int


@dataclass
class ParsedFlag:
    row_no: int
    id: int | None
    title: str
    color: str


@dataclass
class ParsedUnit:
    row_no: int
    id: int | None
    label: str
    category_code: str
    group_path: tuple[str, ...] | None     # None — объект без узла
    floor: str
    predecessors: list[int] = field(default_factory=list)
    flags: frozenset[str] = frozenset()     # названия флагов с «да»


@dataclass
class ParsedFund:
    categories: list[ParsedCategory]
    groups: list[ParsedGroup]
    flags: list[ParsedFlag]
    units: list[ParsedUnit]
    flag_columns: list[str] = field(default_factory=list)   # заголовки колонок флагов «Объектов»


# ---------------------------------------------------------------- выгрузка

def _yn(v: bool) -> str:
    return "да" if v else "нет"


def _setup_sheet(ws, headers: list[str], widths: list[int]) -> None:
    for i, (h, w) in enumerate(zip(headers, widths, strict=True), start=1):
        c = ws.cell(row=1, column=i, value=safe_sheet_cell(h))
        c.font = HEAD
        c.protection = Protection(locked=(i == 1))
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = w
        ws.column_dimensions[letter].protection = Protection(locked=(i == 1))
    ws.freeze_panes = "A2"
    ws.protection.sheet = True          # заперт только столбец ID
    ws.protection.formatColumns = False
    ws.protection.formatCells = False


def _put_row(ws, row: int, values: list) -> None:
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=safe_sheet_cell(v))
        if i == 1:
            c.fill = GREY
            c.protection = Protection(locked=True)
        else:
            c.protection = Protection(locked=False)


def export_fund(snapshot: FundSnapshot, flags: list[UnitFlag]) -> bytes:
    used = {f for u in snapshot.units.values() for f in u.flags}
    cols = sorted((f for f in flags if not f.archived or f.id in used), key=lambda f: (f.sort_order, f.id))
    wb = Workbook()

    ws = wb.active
    ws.title = SHEET_CATEGORIES
    _setup_sheet(ws, [ID_HEADER, "код", "название", "вид", "вес", "сопровождение (да/нет)", "порядок"],
                 [16, 14, 28, 14, 8, 22, 10])
    for r, c in enumerate(sorted(snapshot.categories.values(), key=lambda c: (c.sort_order, c.id)), start=2):
        _put_row(ws, r, [c.id, c.code, c.title, c.kind_title, c.weight, _yn(c.escorted), c.sort_order])

    ws = wb.create_sheet(SHEET_TREE)
    tree_ws = ws
    _setup_sheet(ws, [ID_HEADER, "название", "родитель", "порядок"], [16, 28, 50, 10])
    groups = sorted(snapshot.groups.values(), key=lambda g: ([snapshot.groups[i].sort_order for i in g.path], g.id))
    for r, g in enumerate(groups, start=2):
        parent = PATH_SEP.join(snapshot.path_titles(g.parent_id)) if g.parent_id in snapshot.groups else ""
        _put_row(ws, r, [g.id, g.title, parent, g.sort_order])

    ws = wb.create_sheet(SHEET_FLAGS)
    _setup_sheet(ws, [ID_HEADER, "название", "цвет"], [16, 28, 12])
    for r, f in enumerate(cols, start=2):
        _put_row(ws, r, [f.id, f.title, f.color])

    ws = wb.create_sheet(SHEET_UNITS)
    _setup_sheet(ws, [ID_HEADER, *_UNIT_FIXED, *[f.title for f in cols]],
                 [16, 12, 18, 40, 8, 22, *[14] * len(cols)])
    for r, u in enumerate(sorted(snapshot.units.values(), key=lambda u: (u.label, u.id)), start=2):
        cat = snapshot.categories.get(u.category_id)
        path = PATH_SEP.join(snapshot.path_titles(u.group_id)) if u.group_id in snapshot.groups else ""
        # «Из чего» — история, а не состояние фонда: выгрузка оставляет колонку пустой (§9.1).
        _put_row(ws, r, [u.id, u.label, cat.code if cat else "", path, u.floor, "",
                         *[_yn(f.id in u.flags) for f in cols]])

    # Скрытый «Списки»: все пути в порядке дерева; на них ссылаются списки «Родитель» и «Узел».
    # Ссылка на диапазон, а не перечисление в формуле — лимит Excel в 255 знаков её не касается.
    lists = wb.create_sheet(SHEET_LISTS)
    lists.cell(row=1, column=1, value=safe_sheet_cell("путь")).font = HEAD
    for r, g in enumerate(groups, start=2):
        lists.cell(row=r, column=1, value=safe_sheet_cell(PATH_SEP.join(snapshot.path_titles(g.id))))
    lists.column_dimensions["A"].width = 60
    lists.protection.sheet = True
    lists.sheet_state = "hidden"
    source = f"={SHEET_LISTS}!$A$2:$A${max(len(groups) + 1, 2)}"
    for sheet, col in ((tree_ws, "C"), (ws, "D")):
        dv = DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=False)
        dv.add(f"{col}2:{col}{MAX_ROWS + 1}")
        sheet.add_data_validation(dv)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- разбор

class _Ctx:
    def __init__(self) -> None:
        self.errors: list[FileError] = []

    def err(self, sheet: str, row: int, msg: str) -> None:
        self.errors.append(FileError(sheet, row, msg))


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    if len(s) > 1 and s[0] == "'" and s[1] in "=+-@\t\r":   # снимаем защиту safe_sheet_cell
        s = s[1:]
    return s


def _int(raw, ctx: _Ctx, sheet: str, row: int, what: str, *, default: int | None = None) -> int | None:
    s = _text(raw)
    if not s:
        return default
    try:
        n = int(s)
    except ValueError:
        ctx.err(sheet, row, f"{what}: «{s}» — не целое число")
        return default
    if n < 0 or n > 2**31 - 1:
        ctx.err(sheet, row, f"{what}: «{s}» вне допустимого диапазона")
        return default
    return n


def _bool(raw, ctx: _Ctx, sheet: str, row: int, what: str) -> bool:
    s = _text(raw).lower()
    if not s or s in _FALSE:
        return False
    if s in _TRUE:
        return True
    ctx.err(sheet, row, f"{what}: «{_text(raw)}» — ожидается да или нет")
    return False


def _path(s: str) -> tuple[str, ...]:
    return tuple(p.strip() for p in s.split(PATH_SEP.strip()))


def _read_sheet(wb, name: str, width: int, ctx: _Ctx) -> tuple[list[str], list[tuple[int, tuple]]] | None:
    """(заголовок, [(номер строки, значения)]) без пустых строк; None — листа нет/слишком велик."""
    if name not in wb.sheetnames:
        ctx.err(name, 0, f"в файле нет листа «{name}»")
        return None
    ws = wb[name]
    rows = ws.iter_rows(min_row=1, max_col=width + MAX_FLAG_COLUMNS + 1, values_only=True)
    header: list[str] | None = None
    out: list[tuple[int, tuple]] = []
    for n, vals in enumerate(rows, start=1):
        if n == 1:
            header = [_text(v) for v in vals]
            continue
        if n - 1 > MAX_ROWS:
            ctx.err(name, 0, f"слишком много строк: больше {MAX_ROWS}")
            return None
        if all(v is None or _text(v) == "" for v in vals):
            continue
        out.append((n, vals + (None,) * max(0, width - len(vals))))
    if header is None or len(header) < width:
        ctx.err(name, 1, f"заголовок листа неполный: ожидается {width} колонок")
        return None
    return header, out


def _dup_ids(items: list, sheet: str, ctx: _Ctx) -> None:
    seen: dict[int, int] = {}
    for it in items:
        if it.id is None:
            continue
        if it.id in seen:
            ctx.err(sheet, it.row_no, f"ID {it.id} повторяется (первый раз в строке {seen[it.id]})")
        else:
            seen[it.id] = it.row_no


def _dup_keys(items: list, key, sheet: str, what: str, ctx: _Ctx) -> None:
    seen: dict = {}
    for it in items:
        k = key(it)
        if k in (None, "", ()):
            continue
        if k in seen:
            ctx.err(sheet, it.row_no, f"{what} «{k if isinstance(k, str) else PATH_SEP.join(k)}» повторяется "
                                      f"(первый раз в строке {seen[k]})")
        else:
            seen[k] = it.row_no


def _check_zip(data: bytes, ctx: _Ctx) -> bool:
    if len(data) > MAX_FILE_BYTES:
        ctx.err("", 0, f"файл больше {MAX_FILE_BYTES // 1024 // 1024} МБ")
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNPACKED_BYTES:
                ctx.err("", 0, "файл слишком велик после распаковки")
                return False
    except (zipfile.BadZipFile, OSError):
        ctx.err("", 0, "не файл Excel (.xlsx)")
        return False
    return True


def parse_fund(data: bytes, known_unit_ids: set[int] | None = None) -> tuple[ParsedFund | None, list[FileError]]:
    """Разбор файла фонда. known_unit_ids — id объектов текущего фонда: если передан, «Из чего»
    с неизвестным id — ошибка; иначе проверяется только формат."""
    ctx = _Ctx()
    if not _check_zip(data, ctx):
        return None, ctx.errors
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — любой мусор вместо xlsx это ошибка файла, а не сбой
        return None, [FileError("", 0, "не удалось прочитать файл Excel (.xlsx)")]
    try:
        return _parse(wb, known_unit_ids, ctx)
    except Exception:  # noqa: BLE001
        return None, [*ctx.errors, FileError("", 0, "файл повреждён или имеет неожиданную структуру")]
    finally:
        wb.close()


def _parse(wb, known_unit_ids: set[int] | None, ctx: _Ctx) -> tuple[ParsedFund | None, list[FileError]]:
    if SHEET_TREE in wb.sheetnames:       # прежний формат: «ID | путь узла | порядок» без «Родителя»
        head = next(wb[SHEET_TREE].iter_rows(min_row=1, max_row=1, max_col=4, values_only=True), ())
        if len(head) < 3 or not _text(head[2]).lower().startswith("родитель"):
            ctx.err(SHEET_TREE, 1, "старый формат файла: нет колонки «Родитель» — выгрузите фонд заново")
            return None, ctx.errors
    cats_raw = _read_sheet(wb, SHEET_CATEGORIES, 7, ctx)
    tree_raw = _read_sheet(wb, SHEET_TREE, 4, ctx)
    flags_raw = _read_sheet(wb, SHEET_FLAGS, 3, ctx)
    units_raw = _read_sheet(wb, SHEET_UNITS, 6, ctx)
    if cats_raw is None or tree_raw is None or flags_raw is None or units_raw is None:
        return None, ctx.errors

    cats: list[ParsedCategory] = []
    for n, v in cats_raw[1]:
        code, title = _text(v[1]), _text(v[2])
        if not code:
            ctx.err(SHEET_CATEGORIES, n, "не указан код категории")
        if not title:
            ctx.err(SHEET_CATEGORIES, n, "не указано название категории")
        w = _text(v[4]).replace(",", ".")
        try:
            weight = float(w) if w else 1.0
        except ValueError:
            ctx.err(SHEET_CATEGORIES, n, f"вес: «{w}» — не число")
            weight = 1.0
        cats.append(ParsedCategory(
            n, _int(v[0], ctx, SHEET_CATEGORIES, n, "ID"), code, title, _text(v[3]), weight,
            _bool(v[5], ctx, SHEET_CATEGORIES, n, "сопровождение"),
            _int(v[6], ctx, SHEET_CATEGORIES, n, "порядок", default=100) or 0))
    _dup_ids(cats, SHEET_CATEGORIES, ctx)
    _dup_keys(cats, lambda c: c.code, SHEET_CATEGORIES, "код", ctx)

    # Узел — «Название» + «Родитель» (полный путь родителя; пусто — корень). Строки в любом порядке:
    # родитель может стоять ниже ребёнка, поэтому наличие родителя проверяется по всему листу.
    groups: list[ParsedGroup] = []
    for n, v in tree_raw[1]:
        title, parent = _text(v[1]), _text(v[2])
        pp = _path(parent) if parent else ()
        if not title:
            ctx.err(SHEET_TREE, n, "не указано название узла")
            continue
        if "/" in title:
            ctx.err(SHEET_TREE, n, "в названии узла нельзя использовать «/»")
            continue
        if any(not x for x in pp):
            ctx.err(SHEET_TREE, n, "родитель содержит пустое звено")
            continue
        groups.append(ParsedGroup(n, _int(v[0], ctx, SHEET_TREE, n, "ID"), (*pp, title),
                                  _int(v[3], ctx, SHEET_TREE, n, "порядок", default=100) or 0))
    _dup_ids(groups, SHEET_TREE, ctx)
    _dup_keys(groups, lambda g: g.path, SHEET_TREE, "путь", ctx)
    paths = {g.path for g in groups}
    for g in groups:
        if len(g.path) > 1 and g.path[:-1] not in paths:
            ctx.err(SHEET_TREE, g.row_no, f"нет узла-родителя «{PATH_SEP.join(g.path[:-1])}»")

    flags: list[ParsedFlag] = []
    for n, v in flags_raw[1]:
        title = _text(v[1])
        if not title:
            ctx.err(SHEET_FLAGS, n, "не указано название флага")
            continue
        flags.append(ParsedFlag(n, _int(v[0], ctx, SHEET_FLAGS, n, "ID"), title, _text(v[2])))
    _dup_ids(flags, SHEET_FLAGS, ctx)
    _dup_keys(flags, lambda f: f.title, SHEET_FLAGS, "флаг", ctx)

    # Колонки флагов листа «Объекты»: новая колонка создаёт флаг (§9.1).
    uheader, urows = units_raw
    flag_cols = [(i, t) for i, t in enumerate(uheader) if i >= 6 and t]
    if len(uheader) > 6 + MAX_FLAG_COLUMNS and any(uheader[6 + MAX_FLAG_COLUMNS:]):
        ctx.err(SHEET_UNITS, 1, f"слишком много колонок флагов: больше {MAX_FLAG_COLUMNS}")
    known_titles = {f.title for f in flags}
    seen_cols: set[str] = set()
    for _, t in flag_cols:
        if t in seen_cols:
            ctx.err(SHEET_UNITS, 1, f"колонка флага «{t}» повторяется")
        seen_cols.add(t)
        if t not in known_titles:
            flags.append(ParsedFlag(0, None, t, ""))
            known_titles.add(t)
    # Заголовок «ID…» / фиксированные колонки — позиционные; проверим, что первая — ID.
    if not uheader[0].lower().startswith("id"):
        ctx.err(SHEET_UNITS, 1, "первая колонка должна быть ID")

    cat_codes = {c.code for c in cats}
    units: list[ParsedUnit] = []
    for n, v in urows:
        label = _text(v[1])
        if not label:
            ctx.err(SHEET_UNITS, n, "не указан номер объекта")
        code = _text(v[2])
        if not code or code not in cat_codes:
            ctx.err(SHEET_UNITS, n, f"неизвестная категория «{code}»" if code else "не указана категория")
        gs = _text(v[3])
        gpath = _path(gs) if gs else None
        if gpath is not None and gpath not in paths:
            ctx.err(SHEET_UNITS, n, f"неизвестный узел «{gs}»")
        preds: list[int] = []
        for part in _text(v[5]).replace(";", ",").split(","):
            if not part.strip():
                continue
            pid = _int(part, ctx, SHEET_UNITS, n, "«Из чего»")
            if pid is None:
                continue
            if known_unit_ids is not None and pid not in known_unit_ids:
                ctx.err(SHEET_UNITS, n, f"«Из чего»: неизвестный ID {pid}")
            preds.append(pid)
        uid = _int(v[0], ctx, SHEET_UNITS, n, "ID")
        if uid is not None and preds:
            ctx.err(SHEET_UNITS, n, "«Из чего» заполняется только у новых строк (ID пуст)")
        on = frozenset(t for i, t in flag_cols
                       if i < len(v) and _bool(v[i], ctx, SHEET_UNITS, n, f"флаг «{t}»"))
        units.append(ParsedUnit(n, uid, label, code, gpath,
                                _text(v[4]), preds, on))
    _dup_ids(units, SHEET_UNITS, ctx)
    _dup_keys(units, lambda u: u.label, SHEET_UNITS, "номер", ctx)

    if ctx.errors:
        return None, ctx.errors
    return ParsedFund(cats, groups, flags, units, [t for _, t in flag_cols]), []
