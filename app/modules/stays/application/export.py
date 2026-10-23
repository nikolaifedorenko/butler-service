"""Выгрузка в Excel (§8): обезличенная — для непрофильных служб (объект, даты, статус,
число гостей, категория); полная — с гостями, батлерами и логистикой. Данные берутся из
тех же видов, что и экраны, с уровнем видимости выгрузки — лишнего в файл не попадает."""
from __future__ import annotations

import datetime as dt
import io
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.format import safe_sheet_cell

from ..adapters.orm import Logistics, Reservation
from ..domain.halfday import PM, pos
from .context import load_context
from .reservations import STATUS_TITLES
from .views import BoardFilters, check_window, units_board

HEAD = Font(bold=True)
BUSY = PatternFill("solid", fgColor="DCE6F5")


def occupancy_counts(board: dict, days: list[dt.date]) -> dict[str, dict[str, int]]:
    """Категория → день → занято объектов (ночь дня: полоса покрывает PM даты)."""
    cat_of = {u["id"]: u["category"] for u in board["units"]}
    out: dict[str, dict[str, int]] = {}
    for d in days:
        p = pos(d, PM)
        busy = {b["unit_id"] for b in board["bars"] if b["unit_id"] and b["start_pos"] <= p <= b["end_pos"]}
        for uid in busy:
            out.setdefault(cat_of[uid], {}).setdefault(d.isoformat(), 0)
            out[cat_of[uid]][d.isoformat()] += 1
    return out


def _put(ws, row: int, values: list, bold: bool = False) -> None:
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=safe_sheet_cell(v))
        if bold:
            c.font = HEAD


def export_xlsx(db: Session, first: dt.date, last: dt.date, kind: Literal["full", "anon"]) -> bytes:
    days = check_window(first, last)
    level = "full" if kind == "full" else "occupancy"
    board = units_board(db, first, last, BoardFilters(), level)
    wb = Workbook()

    ws = wb.active
    ws.title = "Шахматка"
    _put(ws, 1, ["Объект", "Категория", *[d.strftime("%d.%m") for d in days]], bold=True)
    for r, u in enumerate(board["units"], start=2):
        _put(ws, r, [u["label"], u["category"]])
        for c, d in enumerate(days, start=3):
            p = pos(d, PM)
            bar = next((b for b in board["bars"] if b["unit_id"] == u["id"] and b["start_pos"] <= p <= b["end_pos"]), None)
            if bar:
                cell = ws.cell(row=r, column=c, value=safe_sheet_cell(bar.get("main_guest") or bar.get("number") or "занято")
                               if kind == "full" else "занято")
                cell.fill = BUSY
    for i in range(3, len(days) + 3):
        ws.column_dimensions[get_column_letter(i)].width = 14

    ws = wb.create_sheet("Брони")
    head = ["Объект", "Категория", "Заезд", "Выезд", "Статус", "Гостей"]
    if kind == "full":
        head += ["№ брони", "Гости", "Уровень", "Метка", "Заказчик", "Повод", "Примечание", "Батлеры", "Логистика"]
    _put(ws, 1, head, bold=True)
    labels = {u["id"]: u for u in board["units"]}
    rids = sorted({b["reservation_id"] for b in board["bars"]})
    res = {r.id: r for r in db.scalars(select(Reservation).where(Reservation.id.in_(rids)))} if rids else {}
    butlers, logistics, occupants = {}, {}, {}
    if kind == "full" and rids:
        ctx = load_context(db, first, last, None)
        for d in ctx.days():
            for (rid, _g), emps in ctx.cover(d).by_guest.items():
                for e in emps:
                    butlers.setdefault(rid, {}).setdefault(ctx.name(e), []).append(d.strftime("%d.%m"))
        for rid in rids:
            occupants[rid] = sorted({ctx.guest_name(o.guest_id) for o in ctx.occ.get(rid, [])})
        for lg in db.scalars(select(Logistics).where(Logistics.reservation_id.in_(rids))):
            logistics.setdefault(lg.reservation_id, []).append(
                f"{lg.date.strftime('%d.%m') if lg.date else ''} {lg.time} {lg.description}".strip())
    row = 2
    for rid in rids:
        r = res[rid]
        bars = [b for b in board["bars"] if b["reservation_id"] == rid]
        units = ", ".join(dict.fromkeys(labels[b["unit_id"]]["label"] for b in bars if b["unit_id"] in labels)) or "без объекта"
        cats = ", ".join(dict.fromkeys(labels[b["unit_id"]]["category"] for b in bars if b["unit_id"] in labels))
        values = [units, cats, r.check_in.strftime("%d.%m.%Y"), r.check_out.strftime("%d.%m.%Y"),
                  STATUS_TITLES[r.status], bars[0]["guests_count"]]
        if kind == "full":
            b0 = bars[0]
            values += [r.number, ", ".join(occupants.get(rid, [])), b0.get("level", ""),
                       (b0.get("label") or {}).get("title", ""), r.customer, r.purpose, r.note,
                       "; ".join(f"{n}: {', '.join(ds)}" for n, ds in butlers.get(rid, {}).items()),
                       "; ".join(logistics.get(rid, []))]
        _put(ws, row, values)
        row += 1

    ws = wb.create_sheet("Сводка")
    _put(ws, 1, ["Категория", "Объектов", *[d.strftime("%d.%m") for d in days]], bold=True)
    counts = occupancy_counts(board, days)
    totals: dict[str, int] = {}
    for u in board["units"]:
        totals[u["category"]] = totals.get(u["category"], 0) + 1
    for r_i, (cat, total) in enumerate(sorted(totals.items()), start=2):
        _put(ws, r_i, [cat, total, *[counts.get(cat, {}).get(d.isoformat(), 0) for d in days]])
    no_show = db.scalars(select(Reservation.check_in).where(Reservation.status == "no_show",
                                                            Reservation.check_in >= first, Reservation.check_in <= last))
    per_day: dict[dt.date, int] = {}
    for d in no_show:
        per_day[d] = per_day.get(d, 0) + 1
    _put(ws, len(totals) + 3, ["Не заехали", "", *[per_day.get(d, 0) for d in days]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
