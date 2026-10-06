"""Экспорт «Графика» и «Табеля» в PDF через DOCX.

Кнопки «PDF» рядом с «Печать»: сервер собирает DOCX с той же таблицей, что и
окно браузерной печати (альбомная A4, цветные подложки ячеек как на экране,
шапка повторяется на каждой странице), затем превращает его в PDF движком
из `pdf_render` (LibreOffice, а без него — встроенный рендер на reportlab).

python-docx входит в основные зависимости, поэтому сам DOCX собирается всегда;
PDF-движок опционален — без него эндпоинты отдают 501 с понятной подсказкой,
а «Печать» в браузере и выгрузка Excel продолжают работать.
"""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .groups import sorted_groups

# ── геометрия страницы: альбомная A4, узкие поля — влезает весь месяц ──
PAGE_W_MM, PAGE_H_MM = 297.0, 210.0
MARGIN_X_MM, MARGIN_Y_MM = 6.0, 9.0
USABLE_MM = PAGE_W_MM - 2 * MARGIN_X_MM          # 285 мм под таблицу

MUTED = RGBColor(0x5D, 0x6C, 0x85)
GREEN = RGBColor(0x1C, 0x6B, 0x41)
RED = RGBColor(0xC0, 0x39, 0x2B)

HEAD_FILL = "E8EEF7"      # шапка
WEEKEND_FILL = "F1F4FA"   # выходные в шапке
GROUP_FILL = "EEF2F8"     # строка блока / итога
INACTIVE_FILL = "ECEFF4"  # дни вне периодов работы
OUT_FILL = "D8D3E8"       # «чужие» дни при переходе между сменами
WORKED_OFF = "#7c3aed"    # работа в выходной

_hex_re = re.compile(r"^#?([0-9a-fA-F]{6})$")


def pdf_unavailable_detail(how: str) -> str:
    """Подсказка для 501: чем заменить PDF и как включить движок."""
    return (f"PDF на сервере собрать не удалось: {how} "
            "Установите PDF-движок: pip install -r requirements-pdf.txt "
            "(или LibreOffice — тогда PDF будет один в один как Word). "
            "Без него работают «Печать» в браузере и выгрузка Excel.")


def _fmt_hours(x) -> str:
    try:
        return f"{float(x or 0):g}"
    except (TypeError, ValueError):
        return "0"


def _hex6(color) -> str | None:
    m = _hex_re.match(color or "")
    return m.group(1).upper() if m else None


def _tint(color, alpha: float = 0.28, default: str = "FFFFFF") -> str:
    """Светлая подложка ячейки: цвет смены, смешанный с белым (как на экране)."""
    h = _hex6(color) or _hex6(default) or "FFFFFF"
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    mix = lambda c: max(0, min(255, round(c * alpha + 255 * (1 - alpha))))
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


# ── мелкие помощники python-docx ──────────────────────────────────────────
def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _cell_text(cell, text: str, size: float = 8.0, bold: bool = False,
               align=WD_ALIGN_PARAGRAPH.CENTER, color: RGBColor | None = None) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _name_cell(cell, name: str, sub: str = "", size: float = 7.5) -> None:
    """ФИО + должность второй строкой (левый край, как в экранной сетке)."""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(name or "")
    run.font.size = Pt(size)
    run.font.bold = True
    if sub:
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(sub)
        r2.font.size = Pt(size - 1.5)
        r2.font.color.rgb = MUTED


def _new_doc(footer_text: str = "") -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(PAGE_W_MM), Mm(PAGE_H_MM)
    sec.left_margin = sec.right_margin = Mm(MARGIN_X_MM)
    sec.top_margin = sec.bottom_margin = Mm(MARGIN_Y_MM)
    if footer_text:
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run(footer_text)
        run.font.size = Pt(7.5)
        run.font.color.rgb = MUTED
    return doc


def _title(doc: Document, text: str, sub: str = "") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(13)
    run.font.bold = True
    if sub:
        ps = doc.add_paragraph()
        ps.paragraph_format.space_after = Pt(5)
        rs = ps.add_run(sub)
        rs.font.size = Pt(8)
        rs.font.color.rgb = MUTED


def _fixed_widths(table, widths_mm: list[float]) -> None:
    """Фиксированные ширины колонок + узкие поля ячеек (иначе 31 день не влезает)."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("top", 15), ("left", 28), ("bottom", 15), ("right", 28)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)
    for i, col in enumerate(table.columns):
        if i < len(widths_mm):
            col.width = Mm(widths_mm[i])
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_mm):
                cell.width = Mm(widths_mm[i])


def _repeat_header(row) -> None:
    """Шапка таблицы повторяется на каждой странице (для LibreOffice и Word)."""
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


# ── график сменности ──────────────────────────────────────────────────────
def schedule_docx(data: dict) -> bytes:
    """Тот же вид, что и «Печать»/Excel: блоки смен, строки сотрудников,
    код смены в ячейке (звёздочка — согласованное отсутствие части смены)."""
    year = data["year"]
    days = data["days"]
    ndays = max(1, len(days))
    doc = _new_doc(f"График сменности · {data['month_name']} {year} · «Батлер Сервис»")
    _title(doc, f"График сменности, {data['month_name']} {year} г.",
           "План по блокам. «—» — день другой смены или вне периода работы; "
           "×2 в шапке — день двойной оплаты; РАБ — работа в свой выходной.")

    name_w, hours_w = 29.0, 11.0
    day_w = (USABLE_MM - name_w - hours_w) / ndays
    widths = [name_w] + [day_w] * ndays + [hours_w]

    table = doc.add_table(rows=1, cols=ndays + 2)
    table.style = "Table Grid"
    hdr = table.rows[0]
    _repeat_header(hdr)
    _cell_text(hdr.cells[0], "Сотрудник", size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    _shade(hdr.cells[0], HEAD_FILL)
    for i, d in enumerate(days):
        c = hdr.cells[i + 1]
        _cell_text(c, f"{d['day']:02d}\n{d['weekday']}" + ("\n×2" if d.get("double_scope") else ""),
                   size=5.8, bold=True)
        _shade(c, WEEKEND_FILL if d.get("is_weekend") else HEAD_FILL)
    _cell_text(hdr.cells[-1], "Часов", size=7.5, bold=True)
    _shade(hdr.cells[-1], HEAD_FILL)

    groups: dict[str, list[dict]] = {}
    for row in data["rows"]:
        for b in row.get("blocks") or []:
            groups.setdefault(b["group"] or "Без группы", []).append(row)

    for gname in sorted_groups(groups.keys()):
        g_row = table.add_row()
        for c in g_row.cells:
            _shade(c, GROUP_FILL)
        _cell_text(g_row.cells[0], gname.upper(), size=8, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)

        seen: set[int] = set()
        for row in groups[gname]:
            if row["employee"]["id"] in seen and gname not in [b["group"] for b in row["blocks"]]:
                continue
            seen.add(row["employee"]["id"])
            block = next((b for b in row["blocks"] if b["group"] == gname), None)
            ranges = (block or {}).get("ranges") or []
            tr = table.add_row()
            emp = row["employee"]
            _name_cell(tr.cells[0], emp["full_name"], emp.get("position") or "")
            for i, d in enumerate(days):
                iso = d["date"]
                cell = row["cells"].get(iso) or {}
                c = tr.cells[i + 1]
                out = bool(ranges) and not any(
                    rg["start"] <= iso and (not rg["end"] or iso <= rg["end"]) for rg in ranges)
                sh = cell.get("shift")
                if not cell.get("employed", True):
                    _cell_text(c, "—", size=6, color=MUTED)
                    _shade(c, INACTIVE_FILL)
                elif out:
                    _cell_text(c, "—", size=6, color=MUTED)
                    _shade(c, OUT_FILL)
                elif cell.get("worked_off"):
                    _cell_text(c, f"РАБ {_fmt_hours(cell.get('fact_hours'))}", size=5.8, bold=True)
                    _shade(c, _tint(WORKED_OFF, 0.32))
                elif sh:
                    text = (sh.get("display_code") or sh.get("name") or "").replace("\u2013", "-")
                    size = 5.6
                    if cell.get("partial"):
                        text += "*"
                        size = 5.2   # иначе метка не влезает и переносится в 31-дневных месяцах
                    _cell_text(c, text, size=size, bold=sh.get("kind") == "work")
                    _shade(c, _tint(sh.get("color"), 0.28))
                else:
                    _cell_text(c, "", size=6)
                    if d.get("is_weekend"):
                        _shade(c, WEEKEND_FILL)
            _cell_text(tr.cells[-1], _fmt_hours(row["totals"]["planned_hours"]), size=7.5, bold=True)

    _fixed_widths(table, widths)

    # расшифровка кодов — из смен, реально встречающихся в месяце
    seen_codes: dict[str, dict] = {}
    for row in data["rows"]:
        for cell in row["cells"].values():
            sh = cell.get("shift")
            if sh and sh.get("code") and sh["code"] not in seen_codes:
                seen_codes[sh["code"]] = sh
    if seen_codes:
        parts = []
        for sh in seen_codes.values():
            code = sh.get("display_code") or sh.get("name") or sh.get("code")
            if sh.get("kind") == "work":
                parts.append(f"{code} = {sh.get('start', '')}–{sh.get('end', '')}")
            else:
                parts.append(f"{code} — {sh.get('name', '')}")
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        run = p.add_run("Коды: " + " · ".join(parts))
        run.font.size = Pt(7)
        run.font.color.rgb = MUTED

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── табель: сводка план/факт (как в окне браузерной печати) ──────────────
