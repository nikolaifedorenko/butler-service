"""Генерация PDF на сервере: встроенный рендер на reportlab, без внешних программ.

Сервер разбирает уже заполненный .docx (параграфы, таблицы, выравнивание, жирный/курсив,
фоны ячеек, поля страницы, колонтитулы текстом, картинки) и собирает PDF. Кириллица —
шрифты DejaVu из `assets/fonts/` (можно заменить своими: TTF в ту же папку или путь в
`PDF_FONT`). Один путь на Mac, сервере и в Docker; точный вид бланка даёт DOCX.

До 22.10.2026 сначала пробовали LibreOffice: в Docker-образе его не было, и PDF там без
предупреждения собирался встроенным рендером, который тогда пропускал картинки (фото
ночного отчёта). Путь LibreOffice удалён, картинки рендер вставляет сам.
"""
from __future__ import annotations

import contextlib
import io
import os
from pathlib import Path
from typing import Optional

BASE_DIR = Path(__file__).resolve().parent.parent
FONTS_DIR = BASE_DIR / "assets" / "fonts"

# имена шрифтов, которыми рисуем PDF (кириллица обязательна)
F_REGULAR = "DocFont"
F_BOLD = "DocFont-Bold"
F_ITALIC = "DocFont-Italic"
F_BOLD_ITALIC = "DocFont-BoldItalic"

_font_ready = False
_font_error = ""


def _candidates(regular: str) -> list[Path]:
    """Где искать шрифт с кириллицей: переменная окружения, папка проекта, системные."""
    env = os.environ.get("PDF_FONT")
    sys_paths = [
        # Linux
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/freefont/FreeSans.ttf",
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
        # macOS
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    out = []
    if env:
        out.append(Path(env))
    out.append(FONTS_DIR / regular)
    out += [Path(p) for p in sys_paths]
    return out


def _register_fonts() -> tuple[bool, str]:
    """Регистрируем обычный/жирный/курсив (чего нет — заменяем обычным)."""
    global _font_ready, _font_error
    if _font_ready:
        return True, ""
    try:
        from reportlab.lib.fonts import addMapping
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        _font_error = ("Не установлен PDF-движок (reportlab). Выполните: pip install -r requirements.txt. "
                       "DOCX и обычная печать работают всегда.")
        return False, _font_error

    regular = next((p for p in _candidates("DejaVuSans.ttf") if p.exists()), None)
    if regular is None:
        _font_error = ("Не найден шрифт с кириллицей (TTF). Положите DejaVuSans.ttf в "
                       f"{FONTS_DIR} или задайте переменную окружения PDF_FONT=/путь/шрифт.ttf")
        return False, _font_error

    def find(names: list[str], fallback: Path) -> Path:
        for name in names:
            for cand in _candidates(name):
                if cand.exists():
                    return cand
        return fallback

    bold = find(["DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf"], regular)
    italic = find(["DejaVuSans-Oblique.ttf", "DejaVuSans.ttf", "LiberationSans-Italic.ttf"], regular)
    bold_italic = find(["DejaVuSans-BoldOblique.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-Bold.ttf"], bold)

    try:
        pdfmetrics.registerFont(TTFont(F_REGULAR, str(regular)))
        pdfmetrics.registerFont(TTFont(F_BOLD, str(bold)))
        pdfmetrics.registerFont(TTFont(F_ITALIC, str(italic)))
        pdfmetrics.registerFont(TTFont(F_BOLD_ITALIC, str(bold_italic)))
        addMapping(F_REGULAR, 0, 0, F_REGULAR)
        addMapping(F_REGULAR, 1, 0, F_BOLD)
        addMapping(F_REGULAR, 0, 1, F_ITALIC)
        addMapping(F_REGULAR, 1, 1, F_BOLD_ITALIC)
    except Exception as exc:  # pragma: no cover - зависит от окружения
        _font_error = f"Не удалось подключить шрифт: {exc}"
        return False, _font_error
    _font_ready = True
    return True, ""


def native_available() -> bool:
    ok, _ = _register_fonts()
    return ok


def native_error() -> str:
    _, err = _register_fonts()
    return err


def pdf_mode() -> str:
    """native | none — может ли сервер делать PDF."""
    return "native" if native_available() else "none"


# ─────────────────────────── разбор заполненного .docx ───────────────────────────
def _escape(text: str) -> str:
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def _spaces(text: str) -> str:
    """Табы и подряд идущие пробелы → неразрывные (reportlab их иначе схлопывает)."""
    text = text.replace("\t", " " * 6)
    out, run = [], 0
    for ch in text:
        if ch == " ":
            run += 1
        else:
            out.append("&nbsp;" * run if run > 1 else " " * run)
            run = 0
            out.append(ch)
    out.append("&nbsp;" * run if run > 1 else " " * run)
    return "".join(out)


def _para_markup(p, default_size: float) -> tuple[str, float, str, bool]:
    """Текст параграфа с разметкой <b>/<i>, размер, выравнивание, «это пустая строка»."""
    parts = []
    size = default_size
    for run in p.runs:
        txt = run.text or ""
        if not txt:
            continue
        chunk = _spaces(_escape(txt))
        bold = bool(run.bold)
        italic = bool(run.italic)
        if run.font and run.font.size is not None:
            with contextlib.suppress(Exception):
                size = float(run.font.size.pt)
        if run.underline:
            chunk = f"<u>{chunk}</u>"
        if bold and italic:
            chunk = f'<font name="{F_BOLD_ITALIC}">{chunk}</font>'
        elif bold:
            chunk = f'<font name="{F_BOLD}">{chunk}</font>'
        elif italic:
            chunk = f'<font name="{F_ITALIC}">{chunk}</font>'
        parts.append(chunk)
    text = "".join(parts) or _spaces(_escape(p.text or ""))
    align = str(getattr(p.alignment, "name", "") or "").upper() if p.alignment is not None else ""
    return text, size, align, not bool((p.text or "").strip())


def _align_of(name: str):
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    if "CENTER" in name:
        return TA_CENTER
    if "RIGHT" in name:
        return TA_RIGHT
    if "JUSTIFY" in name:
        return TA_JUSTIFY
    return TA_LEFT


def _run_look(p, default_size: float = 9.0) -> tuple[float, bool, str | None]:
    """Оформление первого содержательного рана параграфа: размер, жирный, цвет."""
    size: float | None = None
    bold = False
    color: str | None = None
    for run in p.runs:
        font = getattr(run, "font", None)
        if font is None:
            continue
        if size is None and font.size is not None:
            with contextlib.suppress(Exception):
                size = max(4.0, min(20.0, float(font.size.pt)))
        if not bold and font.bold:
            bold = True
        if color is None:
            try:
                rgb = font.color.rgb if font.color is not None else None
                if rgb is not None:
                    color = "#" + str(rgb)
            except Exception:
                pass
        if size is not None and bold and color:
            break
    return size or default_size, bold, color


def _cell_markup(p) -> str:
    """Текст параграфа ячейки таблицы в разметке reportlab (переводы строк → <br/>)."""
    return _spaces(_escape(p.text or "")).replace("\n", "<br/>")


def _grid_widths(tbl_el, avail: float) -> list[float]:
    """Ширины колонок пропорционально tblGrid документа (плотные сетки графика иначе не
    влезали: все колонки получались одинаковыми). [] — сетки нет."""
    from docx.oxml.ns import qn

    try:
        grid_el = tbl_el.find(qn("w:tblGrid"))
        ws = [float(gc.get(qn("w:w")) or 0)
              for gc in grid_el.findall(qn("w:gridCol"))] if grid_el is not None else []
    except Exception:
        return []
    total = sum(ws)
    return [avail * w / total for w in ws] if total > 0 else []


def _pictures(p, doc, max_w: float, max_h: float) -> list:
    """Картинки абзаца (w:drawing → a:blip) — размер из wp:extent документа, не шире max_w
    и не выше max_h, пропорции сохраняются. Нечитаемую картинку пропускаем: документ важнее."""
    from docx.oxml.ns import qn
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image

    out = []
    for drawing in p._p.iter(qn("w:drawing")):
        blip = next(drawing.iter(qn("a:blip")), None)
        rid = blip.get(qn("r:embed")) if blip is not None else None
        part = doc.part.related_parts.get(rid) if rid else None
        if part is None:
            continue
        try:
            nat_w, nat_h = ImageReader(io.BytesIO(part.blob)).getSize()   # заодно проверка файла
        except Exception:
            continue
        extent = next(drawing.iter(qn("wp:extent")), None)
        w, h = ((int(extent.get("cx")) / 12700, int(extent.get("cy")) / 12700)   # EMU → pt
                if extent is not None else (float(nat_w), float(nat_h)))
        if w <= 0 or h <= 0:
            continue
        k = min(1.0, max_w / w, max_h / h)
        out.append(Image(io.BytesIO(part.blob), width=w * k, height=h * k))
    return out


def docx_to_pdf_native(docx_bytes: bytes, filename: str = "document.pdf") -> Optional[bytes]:
    """Собрать PDF из заполненного .docx собственными силами (без внешних программ)."""
    ok, err = _register_fonts()
    if not ok:
        return None
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph as DocxParagraph
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            BaseDocTemplate,
            Frame,
            PageTemplate,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            TableStyle,
        )
        from reportlab.platypus import Table as RLTable
    except ImportError:
        return None

    try:
        doc = Document(io.BytesIO(docx_bytes))
    except Exception:
        return None

    section = doc.sections[0] if doc.sections else None
    if section is not None:
        page_w = max(60 * mm, float(section.page_width.mm) * mm) if section.page_width else A4[0]
        page_h = max(80 * mm, float(section.page_height.mm) * mm) if section.page_height else A4[1]
        margins = (float(section.left_margin.mm) * mm, float(section.right_margin.mm) * mm,
                   float(section.top_margin.mm) * mm, float(section.bottom_margin.mm) * mm)
        landscape = bool(section.orientation) and page_w < page_h
        if landscape:
            page_w, page_h = page_h, page_w
    else:
        page_w, page_h = A4
        margins = (20 * mm, 15 * mm, 20 * mm, 20 * mm)
    left, right, top, bottom = [m if m and m > 0 else 18 * mm for m in margins]

    base_style = ParagraphStyle(
        "DocBody", fontName=F_REGULAR, fontSize=11, leading=15,
        spaceAfter=4, alignment=_align_of(""))

    header_text = ""
    footer_text = ""
    if section is not None:
        try:
            header_text = "\n".join(p.text for p in section.header.paragraphs if p.text.strip())
            footer_text = "\n".join(p.text for p in section.footer.paragraphs if p.text.strip())
        except Exception:
            pass

    avail = page_w - left - right
    frame_h = page_h - top - bottom - 12
    story = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = DocxParagraph(child, doc)
            text, size, align, empty = _para_markup(p, base_style.fontSize)
            pics = _pictures(p, doc, avail, frame_h)
            story.extend(pics)
            if empty and pics:
                continue
            if empty:
                story.append(Spacer(1, max(4, size * 0.6)))
                continue
            style = ParagraphStyle(
                f"p{len(story)}", parent=base_style, fontSize=size, leading=size * 1.35,
                alignment=_align_of(align))
            try:
                story.append(Paragraph(text, style))
            except Exception:
                story.append(Paragraph(_escape(p.text or ""), base_style))
        elif tag == "tbl":
            tbl = DocxTable(child, doc)
            grid_w = _grid_widths(child, avail)
            rows = []
            cell_styles = []      # фоны ячеек из w:shd (цветные сетки графика/табеля)
            for r_i, row in enumerate(tbl.rows):
                cells = []
                for c_i, cell in enumerate(row.cells):
                    paras = []
                    cap = (grid_w[c_i] if c_i < len(grid_w) else avail) - 6   # минус поля ячейки
                    for pp in cell.paragraphs:
                        pics = _pictures(pp, doc, cap, frame_h)
                        paras.extend(pics)
                        txt = _cell_markup(pp)
                        if pics and not txt:
                            continue
                        size, bold, color = _run_look(pp)
                        align = _align_of(str(getattr(pp.alignment, "name", "") or "").upper())
                        paras.append(Paragraph(txt or "&nbsp;", ParagraphStyle(
                            f"cell{r_i}_{c_i}", parent=base_style, fontSize=size,
                            leading=size * 1.28, spaceAfter=0, alignment=align,
                            fontName=F_BOLD if bold else F_REGULAR,
                            textColor=colors.HexColor(color) if color else colors.black)))
                    cells.append(paras or [Paragraph("&nbsp;", base_style)])
                    tc_pr = cell._tc.tcPr
                    if tc_pr is not None:
                        shd = tc_pr.find(qn("w:shd"))
                        fill = shd.get(qn("w:fill")) if shd is not None else None
                        if fill and fill.lower() not in ("auto", "ffffff"):
                            with contextlib.suppress(Exception):
                                cell_styles.append(("BACKGROUND", (c_i, r_i), (c_i, r_i),
                                                    colors.HexColor("#" + fill)))
                rows.append(cells)
            if not rows:
                continue
            n_cols = max(len(r) for r in rows)
            for r in rows:
                while len(r) < n_cols:
                    r.append([Paragraph("&nbsp;", base_style)])
            col_w = grid_w if len(grid_w) == n_cols else [avail / n_cols] * n_cols
            t = RLTable(rows, colWidths=col_w, repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9aa5b5")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f8")),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ] + cell_styles))
            story.append(Spacer(1, 4))
            story.append(t)
            story.append(Spacer(1, 6))

    if not story:
        story = [Paragraph("&nbsp;", base_style)]

    buf = io.BytesIO()

    def on_page(canvas, doc_):
        canvas.saveState()
        if header_text:
            canvas.setFont(F_REGULAR, 8.5)
            canvas.setFillColor(colors.HexColor("#5c6b84"))
            canvas.drawRightString(page_w - right, page_h - top * 0.55, header_text.replace("\n", " · "))
        if footer_text:
            canvas.setFont(F_REGULAR, 8.5)
            canvas.setFillColor(colors.HexColor("#5c6b84"))
            canvas.drawString(left, bottom * 0.45, footer_text.replace("\n", " · "))
            canvas.drawRightString(page_w - right, bottom * 0.45, f"стр. {canvas.getPageNumber()}")
        canvas.restoreState()

    doc_tpl = BaseDocTemplate(buf, pagesize=(page_w, page_h),
                              leftMargin=left, rightMargin=right,
                              topMargin=top, bottomMargin=bottom,
                              title=filename, author="Батлер Сервис")
    frame = Frame(left, bottom, page_w - left - right, page_h - top - bottom, id="main")
    doc_tpl.addPageTemplates([PageTemplate(id="page", frames=[frame], onPage=on_page)])
    try:
        doc_tpl.build(story)
    except Exception:
        # запасной вариант: простой документ без колонтитулов
        buf = io.BytesIO()
        SimpleDocTemplate(buf, pagesize=(page_w, page_h), leftMargin=left, rightMargin=right,
                          topMargin=top, bottomMargin=bottom, title=filename).build(story)
    return buf.getvalue()


def docx_to_pdf(docx_bytes: bytes, filename: str = "document.pdf") -> tuple[Optional[bytes], str]:
    """PDF из заполненного DOCX. Возвращает (байты или None, чем сделано/почему не вышло)."""
    native = docx_to_pdf_native(docx_bytes, filename)
    if native:
        return native, "native"
    return None, native_error() or "встроенный рендер PDF не смог собрать документ"
