"""Вкладка «Управленческий табель» (16.6): единственный сценарий тяжёлого конвейера.

Период — календарный месяц либо половина месяца по настройке period_granularity
(v7.1/R36). Открытие: месяц не закрыт → view_period (живой расчёт, ничего не
сохраняется; при незакрытом предыдущем — типизированное отклонение); закрыт → чтение
сохранённых SettleResult и снимка FlagsResult без пересчёта. «Закрыть месяц» →
close_period; «Пересчитать» → recalculate (только последний закрытый период).
Периоды вне трудоустройства не предлагаются и показываются отказом (R35)."""
from __future__ import annotations

import datetime as dt
import io
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..auth import Principal, audit
from ..db import get_db
from ..engine.application.close_period import close_period
from ..engine.application.errors import ApplicationError
from ..engine.application.months import (
    half_month_period,
    month_period,
    next_period,
    period_containing,
)
from ..engine.application.recalculate import recalculate
from ..engine.application.view_period import view_period
from ..engine.domain.settings.types import GRANULARITY_HALF_MONTH
from ..engine.domain.time.local_midnight import next_midnight
from ..engine.domain.types.entities import Employment, Period
from ..engine.domain.types.errors import DomainError
from ..engine.interface.serialize import plain
from ..engine.interface.views import period_view_dict
from ..permissions import principal_permissions, require_perm
from .engine_common import domain_error_text, emp_brief, engine_ports, now_utc, period_employees

router = APIRouter(prefix="/api/mgmt", tags=["mgmt-timesheet"])


def _period(payload_year: int, payload_month: int, half: Optional[int], granularity: str) -> Period:
    if granularity == GRANULARITY_HALF_MONTH:
        return half_month_period(payload_year, payload_month, 2 if half == 2 else 1)
    return month_period(payload_year, payload_month)


def _row(db: Session, ports, emp, period: Period, now: dt.datetime, detail: bool) -> dict:
    eid = str(emp.id)
    base = {"employee": emp_brief(emp), "error": None}
    try:
        state = ports.store.period_state(eid, period)
        view = view_period(eid, period, now, ports)
        employment = view.employment
        effective_end = _effective_end(period, employment)
        can_close_now = now >= next_midnight(effective_end, ports.settings.load(period).tz)
        meta = (ports.store.saved_meta(eid, period)
                if state.closed and hasattr(ports.store, "saved_meta") else {}) or {}
        return {**base, **meta,
                **period_view_dict(view, period, detail=detail),
                "state": {"closed": state.closed, "previous_closed": state.previous_closed,
                          "next_closed": state.next_closed,
                          "opening_bank_set": state.opening_bank_set,
                          "can_close": (not state.closed)
                          and (state.previous_closed or state.opening_bank_set) and can_close_now,
                          "can_recalculate": state.closed and not state.next_closed},
                "source": "saved" if state.closed else "view"}
    except (DomainError, ApplicationError, ValueError) as exc:
        return {**base, "error": domain_error_text(exc)}


def _effective_end(period: Period, employment: Employment) -> dt.date:
    return period.end if employment.end is None else min(period.end, employment.end)


def _context(db: Session, principal: Optional[Principal], year: int, month: int,
             half: Optional[int], ids: Optional[list[int]]):
    """Контекст вкладки: порты, гранулярность, период, сотрудники.

    Повреждённая конфигурация (ConfigError, уровень 1 валидации) — типизированный отказ
    409 с человеческим текстом, а не 500: вкладка должна показать, что именно сломано
    и где это чинить. Раньше исключение всплывало до обработчика и отдавало
    «Internal Server Error» на весь период.
    """
    now = now_utc()
    ports = engine_ports(db, principal, now)
    try:
        granularity = ports.settings.load(month_period(year, month)).period_granularity
    except (DomainError, ValueError) as exc:
        raise HTTPException(status_code=409,
                            detail=f"{domain_error_text(exc)}. Исправьте в разделе "
                                   "«Настройки» — расчёт периода не запустится, "
                                   "пока конфигурация невалидна") from exc
    period = _period(year, month, half, granularity)
    emps = period_employees(db, period.start, period.end, ids)
    return now, ports, granularity, period, emps


@router.get("")
def get_mgmt(year: int, month: int, half: Optional[int] = None, employee_id: Optional[int] = None,
             detail: bool = False, principal: Principal = Depends(require_perm("mgmt.view")),
             db: Session = Depends(get_db)):
    now, ports, granularity, period, emps = _context(db, principal, year, month, half,
                                                     [employee_id] if employee_id else None)
    if employee_id and not emps:
        raise HTTPException(status_code=404,
                            detail="Сотрудник не трудоустроен в этом периоде")
    try:
        st = ports.settings.load(period)
    except (DomainError, ValueError) as exc:
        raise HTTPException(status_code=409, detail=domain_error_text(exc)) from exc
    rows = [_row(db, ports, e, period, now, detail or bool(employee_id)) for e in emps]
    perms = principal_permissions(db, principal)
    return {"year": year, "month": month, "half": half, "granularity": granularity,
            "period": {"start": period.start.isoformat(), "end": period.end.isoformat()},
            "days": [(period.start + dt.timedelta(days=i)).isoformat()
                     for i in range((period.end - period.start).days + 1)],
            "tariffs": list(st.placement_order), "step_minutes": st.step_minutes,
            "settings_versions": list(st.versions), "now": now.isoformat(),
            "can_close": "mgmt.close" in perms, "can_recalculate": "mgmt.recalculate" in perms,
            "rows": rows}


class PeriodAction(BaseModel):
    year: int
    month: int
    half: Optional[int] = None
    employee_ids: Optional[list[int]] = None


def _run_action(payload: PeriodAction, principal: Principal, db: Session, action) -> dict:
    now, ports, _granularity, period, emps = _context(db, principal, payload.year, payload.month,
                                                      payload.half, payload.employee_ids)
    done, errors = [], []
    for emp in emps:
        # у каждого сотрудника своя транзакция закрытия (UnitOfWork внутри use case)
        try:
            view = action(str(emp.id), period, now, ports)
            done.append({"employee_id": emp.id,
                         "bank_closed": round(view.settle.bank_closed_minutes / 60, 2),
                         "deferred": len(view.settle.deferred_blocks)})
        except (DomainError, ApplicationError, ValueError) as exc:
            errors.append({"employee_id": emp.id, "name": emp.display_name,
                           "error": domain_error_text(exc)})
    return {"period": {"start": period.start.isoformat(), "end": period.end.isoformat()},
            "done": done, "errors": errors}


@router.post("/close")
def close(payload: PeriodAction, principal: Principal = Depends(require_perm("mgmt.close")),
          db: Session = Depends(get_db)):
    out = _run_action(payload, principal, db, close_period)
    audit(db, principal, "period_close", f"{payload.year}-{payload.month:02d}",
          {"employees": [d["employee_id"] for d in out["done"]], "errors": out["errors"]})
    db.commit()
    return out


@router.post("/recalculate")
def recalc(payload: PeriodAction, principal: Principal = Depends(require_perm("mgmt.recalculate")),
           db: Session = Depends(get_db)):
    out = _run_action(payload, principal, db, recalculate)
    audit(db, principal, "period_recalculate", f"{payload.year}-{payload.month:02d}",
          {"employees": [d["employee_id"] for d in out["done"]], "errors": out["errors"]})
    db.commit()
    return out


@router.get("/periods")
def periods(principal: Principal = Depends(require_perm("mgmt.view")),
            db: Session = Depends(get_db)):
    """Цепочка периодов сотрудника для навигации (16.1 п.7, R35/R36): granularity +
    границы по трудоустройству каждого сотрудника."""
    now = now_utc()
    utc_today = now.astimezone(dt.timezone.utc).date()
    ports = engine_ports(db, principal, now)
    try:
        st = ports.settings.load(month_period(utc_today.year, utc_today.month))
        today = now.astimezone(st.tz).date()
        st = ports.settings.load(month_period(today.year, today.month))
    except (DomainError, ValueError) as exc:
        raise HTTPException(status_code=409, detail=domain_error_text(exc)) from exc
    granularity = st.period_granularity
    current = period_containing(today, granularity)
    emps = period_employees(db, current.start, current.end)
    out = []
    for emp in emps:
        employment = ports.employment.employment(str(emp.id), current)
        out.append({"employee": emp_brief(emp),
                    "employment": {"start": employment.start.isoformat(),
                                   "end": employment.end.isoformat() if employment.end else None},
                    "state": plain(ports.store.period_state(str(emp.id), current))})
    return {"granularity": granularity, "current": {"start": current.start.isoformat(),
                                                    "end": current.end.isoformat()},
            "employees": out,
            "next": {"start": next_period(current, granularity).start.isoformat()}}


@router.get("/xlsx")
def export_xlsx(year: int, month: int, half: Optional[int] = None,
                principal: Principal = Depends(require_perm("mgmt.view")),
                db: Session = Depends(get_db)):
    """УТ в Excel: «ФИО × дни», ячейка «ДЯ2 3 ДН 1», итоги по кодам и банк (интерфейс-слой)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    data = get_mgmt(year, month, half, None, False, principal, db)
    wb = Workbook()
    ws = wb.active
    ws.title = f"УТ {month:02d}.{year}"
    days = data["days"]
    tariffs = data["tariffs"]
    ws.append(["ФИО", "Табельный №"] + [d_[8:] for d_ in days] + tariffs
              + ["Недостача", "Банк на закрытие", "Статус"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in data["rows"]:
        if r.get("error"):
            ws.append([r["employee"]["full_name"], r["employee"]["tab_number"], r["error"]])
            continue
        cells = [" ".join(f"{t} {h / 60:g}" for t, h in r["ut"].get(d_, {}).items())
                 for d_ in days]
        status = "закрыт" if r["state"]["closed"] else "просмотр"
        ws.append([r["employee"]["full_name"], r["employee"]["tab_number"], *cells,
                   *[r["totals"].get(t, 0) / 60 for t in tariffs],
                   r["debt"]["residual_h"], r["bank"]["closed_h"], status])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 34
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"upr_tabel_{year}_{month:02d}{'_h' + str(half) if half else ''}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'})
