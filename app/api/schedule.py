"""Вкладка «График» (16.3, движок v7.1).

График — производный документ: генерируется из Табеля по окнам CFG-WINDOWS
(source="generated"), ручные смены (source="manual") заменяют генерацию дня целиком;
день редактируем ⇔ accepts_ut(дня-якоря) = Да. «Разрез» смены — логический (v7/R17):
кнопка «пропуск/отлучка» создаёт запись согласованного отсутствия (UTC-интервал, может
пересекать полночь), хранимая смена физически не делится; удаление отсутствия
восстанавливает отображение. Правка Графика не меняет часы (RUL-SCHEDULE-ONLY-FLAGS).

Циклы/блоки — инструмент заполнения Табеля (решение владельца 11.10.2026):
fill-pattern/continue пишут коды и планы в Табель, График перегенерируется."""
from __future__ import annotations

import calendar
import datetime as dt
import re
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..auth import Principal, audit, require_manager
from ..db import get_db
from ..deps import TZ, now_local, safe_sheet_cell, utc_to_local
from ..doc_render import hours_text
from ..doublepay import SCOPE_TITLES, calendar_map
from ..employment import employed_on, periods_of, ranges_from_periods
from ..engine.adapters.sql_settings import SqlSettingsAdapter
from ..engine.application import schedule_view
from ..engine.application.errors import ApplicationError
from ..engine.application.months import month_period
from ..engine.application.schedule_view import rebuild_generated_schedule
from ..engine.domain.day_card.fmt_hours import fmt_hours
from ..engine.domain.settings.types import Settings
from ..engine.domain.types.entities import Period, Shift
from ..engine.domain.types.errors import ConfigError, DomainError
from ..groups import GROUP_ORDER, group_index, group_kind, normalize_group
from ..models import (
    BlockAssignment,
    Employee,
    EmploymentPeriod,
    ScheduleAbsence,
    ScheduleShift,
    ShiftType,
    TabelDay,
)
from ..names import suggest_genitive
from ..permissions import require_perm
from ..schedule_patterns import pattern_list, validate_pattern
from ..services.engine_common import day_code_colors, domain_error_text, engine_ports, now_utc
from ..tabel_fill import fill_timesheet_from_cycles

router = APIRouter(prefix="/api/schedule", tags=["schedule"])
_HHMM = re.compile(r"^([01]?\d|2[0-3]):[0-5]\d$")

OUT_OF_BLOCK_COLOR = "#efe6d7"
WORKED_OFF_COLOR = "#e7f0e2"
INACTIVE_COLOR = "#f1f1f1"


def _month_bounds(year: int, month: int) -> tuple[dt.date, dt.date]:
    last = calendar.monthrange(year, month)[1]
    return dt.date(year, month, 1), dt.date(year, month, last)


def _local_hm(value: dt.datetime) -> str:
    return value.astimezone(TZ).strftime("%H:%M")


def _to_utc(day: dt.date, hhmm: str, next_day: bool = False) -> dt.datetime:
    hh, mm = (int(x) for x in hhmm.split(":"))
    local = dt.datetime.combine(day + (dt.timedelta(days=1) if next_day else dt.timedelta()),
                                dt.time(hh, mm), TZ)
    return local.astimezone(dt.timezone.utc)


def _settings(db: Session, period: Period) -> Settings:
    try:
        return SqlSettingsAdapter(db).load(period)
    except ConfigError as exc:
        raise HTTPException(status_code=409, detail=domain_error_text(exc))


def _shift_brief(row: ScheduleShift) -> dict:
    start, end = utc_to_local(row.start_utc.replace(tzinfo=dt.timezone.utc)), \
        utc_to_local(row.end_utc.replace(tzinfo=dt.timezone.utc))
    return {"source": row.source, "anchor": row.date.isoformat(),
            "start": start.isoformat(), "end": end.isoformat(),
            "start_hm": start.strftime("%H:%M"), "end_hm": end.strftime("%H:%M"),
            "note": row.note or "", "id": row.id}


def _absence_brief(row: ScheduleAbsence, kind: Optional[ShiftType],
                   whole_shift: bool = False) -> dict:
    start = utc_to_local(row.start_utc.replace(tzinfo=dt.timezone.utc))
    end = utc_to_local(row.end_utc.replace(tzinfo=dt.timezone.utc))
    return {"id": row.id, "start": start.isoformat(), "end": end.isoformat(),
            "start_hm": start.strftime("%H:%M"), "end_hm": end.strftime("%H:%M"),
            "note": row.note or "",
            # покрывает всё рабочее окно дня — «выходной за часы» (EX-29): сетка
            # рисует плашку на весь день вместо разрезанных отрезков смены
            "whole_shift": whole_shift,
            "kind": None if kind is None else
            {"id": kind.id, "code": kind.code, "name": kind.name, "color": kind.color,
             "doc_type": kind.doc_type or ""}}


def _cell_color(shifts: list, absences: list, trow: Optional[TabelDay], work_types: dict,
                absence_types: dict, code_colors: dict) -> Optional[str]:
    """Цвет ячейки Графика — из словаря «Смены и отсутствия» (16.10.2026).

    Разрешается СЕРВЕРОМ, чтобы экран, браузерная печать, PDF и Excel красили
    одинаково: плашка «выходной за часы» — цвет вида отсутствия; смена — цвет
    словарной смены с тем же окном (первая по id); код дня без смены — цвет вида
    по tzh_code; запасной вариант — цвет кода Табеля из presentation настроек."""
    off = next((a for a in absences if a.get("whole_shift")), None)
    if off is not None:
        return (off.get("kind") or {}).get("color") or "#0e9aa7"
    if shifts:
        wt = work_types.get((shifts[0]["start_hm"], shifts[0]["end_hm"]))
        if wt is not None and (wt.color or "").strip():
            return wt.color
    elif trow is not None:
        at = absence_types.get(trow.code)
        if at is not None and (at.color or "").strip():
            return at.color
    if trow is not None:
        return code_colors.get(trow.code)
    return None


def _visible_segments(shift_rows: list, absence_rows: list) -> list[list[str]]:
    """Видимые части смен дня: смена МИНУС интервалы согласованных отсутствий.

    Разрез логический (v7/R17): записи смен не меняются, сетка показывает остаток —
    каждый отрезок с новой строки («08:00–12:00», «14:00–16:00», «17:00–20:00»).
    Моменты — реальные UTC-датавремя, поэтому отрезки через полночь режутся без
    особых случаев. Без отсутствий результат равен интервалам самих смен."""
    out: list[list[str]] = []
    for sh in shift_rows:
        segs = [(sh.start_utc, sh.end_utc)]
        for a in sorted(absence_rows, key=lambda r: r.start_utc):
            nxt = []
            for ss, se in segs:
                if a.end_utc <= ss or a.start_utc >= se:
                    nxt.append((ss, se))
                    continue
                if a.start_utc > ss:
                    nxt.append((ss, a.start_utc))
                if a.end_utc < se:
                    nxt.append((a.end_utc, se))
            segs = nxt
        hm = lambda v: utc_to_local(v.replace(tzinfo=dt.timezone.utc)).strftime("%H:%M")  # noqa: E731
        out.extend([[hm(ss), hm(se)] for ss, se in segs])
    return out


# единая сущность «отпросился (выходной за часы)» (Q3, 12.10.2026): один вид словаря,
# две точки ввода — интервал части смены либо вся смена. Заявление выбирается по точке
# ввода: на всю смену — «выходной за ранее отработанные часы», интервал — «об отсутствии
# на рабочем месте». Оба вида заявлений уже есть в словаре видов (DOC_TYPES).
DOC_WHOLE_SHIFT = "day_off_hours"
DOC_PARTIAL_SHIFT = "time_off_request"


def _statement_kind_exists(db: Session, code: str) -> bool:
    from ..models import StatementKind

    return db.scalar(select(StatementKind.id).where(StatementKind.code == code)) is not None


def doc_type_of_absence(db: Session, kind: ShiftType, whole_shift: bool) -> str:
    """Какое заявление печатать по отсутствию.

    * вид с кодом Табеля (ОТ/ДО/ОВ/Б…) — заявление определяется самим видом и печатается
      из Табеля (Q4: отпуск, за свой счёт, оплачиваемый выходной — только слой 0);
    * вид без кода Табеля — это «отпросился (выходной за часы)» из Графика: документ
      зависит от точки ввода (вся смена / интервал). Если нужный вид заявления в
      словаре отсутствует (владелец его удалил), остаётся привязка самого вида смены.
    """
    own = (kind.doc_type or "").strip()
    if (kind.tzh_code or "").strip():
        return own
    wanted = DOC_WHOLE_SHIFT if whole_shift else DOC_PARTIAL_SHIFT
    return wanted if _statement_kind_exists(db, wanted) else own


def _position_rank(position: Optional[str]) -> int:
    """Порядок внутри блока (решение владельца 16.10.2026): старшие («Старший
    батлер» и т.п.) впереди остальных, дальше — алфавит. Работает на всех стоках
    сетки разом (экран, печать, PDF, Excel), потому что порядок задаётся здесь."""
    return 0 if position and "старш" in position.lower() else 1


def _visible_employees(db: Session, first: dt.date,
                       last: dt.date) -> tuple[list[Employee], dict[int, list]]:
    """Сотрудники месяца + их периоды работы (батч, без N+1)."""
    candidates = db.scalars(select(Employee).where(Employee.deleted_at.is_(None))).all()
    all_periods = db.scalars(select(EmploymentPeriod).where(
        EmploymentPeriod.employee_id.in_([e.id for e in candidates]))).all() if candidates else []
    per_emp: dict[int, list] = {}
    for p in all_periods:
        per_emp.setdefault(p.employee_id, []).append(p)

    def visible(emp: Employee) -> bool:
        recs = per_emp.get(emp.id)
        if not recs:
            return bool(emp.active)
        return any(p.start_date <= last and (p.end_date is None or p.end_date >= first)
                   for p in recs)

    out = [e for e in candidates if visible(e)]
    out.sort(key=lambda e: (group_index(normalize_group(e.schedule_group)),
                            _position_rank(e.position), e.full_name))
    return out, per_emp


def _strip_to_plan(payload: dict) -> dict:
    """Батлеры видят только план: без заметок, id записей и служебных полей."""
    for row in payload["rows"]:
        for cell in row["cells"].values():
            cell["absences"] = [{k: a[k] for k in ("start_hm", "end_hm", "whole_shift")
                                 if k in a} for a in cell["absences"]]
            if cell.get("tabel"):
                cell["tabel"].pop("note", None)
            for sh in cell["shifts"]:
                sh.pop("note", None)
                sh.pop("id", None)
    return payload


@router.get("")
def get_grid(year: int, month: int, principal: Principal = Depends(require_perm("schedule.view")),
             db: Session = Depends(get_db)):
    """Сетка Графика: ФИО × дни месяца; ячейка — Табель-код, отрезки смен (generated/
    manual) и согласованные отсутствия (логический разрез)."""
    if not (1 <= month <= 12):
        raise HTTPException(status_code=422, detail="month должен быть 1..12")
    if not (1 <= year <= 9999):
        raise HTTPException(status_code=422, detail="year должен быть 1..9999")
    first, last = _month_bounds(year, month)
    period = month_period(year, month)
    days = [first + dt.timedelta(days=i) for i in range((last - first).days + 1)]
    st = _settings(db, period)

    employees, per_emp_periods = _visible_employees(db, first, last)
    emp_ids = [e.id for e in employees]
    lo_utc = dt.datetime.combine(first, dt.time(0), TZ).astimezone(dt.timezone.utc)
    hi_utc = dt.datetime.combine(last + dt.timedelta(days=1), dt.time(0),
                                 TZ).astimezone(dt.timezone.utc)

    shift_rows = db.scalars(select(ScheduleShift).where(
        ScheduleShift.employee_id.in_(emp_ids), ScheduleShift.start_utc < hi_utc,
        ScheduleShift.end_utc > lo_utc).order_by(ScheduleShift.start_utc)).all() if emp_ids else []
    absence_rows = db.scalars(select(ScheduleAbsence).where(
        ScheduleAbsence.employee_id.in_(emp_ids), ScheduleAbsence.start_utc < hi_utc,
        ScheduleAbsence.end_utc > lo_utc).order_by(ScheduleAbsence.start_utc)).all() \
        if emp_ids else []
    tabel_rows = {(r.employee_id, r.date): r for r in db.scalars(select(TabelDay).where(
        TabelDay.employee_id.in_(emp_ids), TabelDay.date >= first,
        TabelDay.date <= last))} if emp_ids else {}
    block_rows = db.scalars(select(BlockAssignment).where(
        BlockAssignment.employee_id.in_(emp_ids),
        BlockAssignment.start_date <= last)).all() if emp_ids else []
    blocks_by_emp: dict[int, list] = {}
    for b in block_rows:
        blocks_by_emp.setdefault(b.employee_id, []).append(b)
    kinds = {k.id: k for k in db.scalars(select(ShiftType).where(
        ShiftType.id.in_({a.kind_id for a in absence_rows if a.kind_id})))} \
        if absence_rows else {}
    # цвета ячеек — из словаря «Смены и отсутствия»: смены матчатся по окну,
    # коды отсутствий — по tzh_code; запасной — цвет кода Табеля (16.10.2026)
    work_types: dict[tuple[str, str], ShiftType] = {}
    for x in db.scalars(select(ShiftType).where(ShiftType.kind == "work")
                        .order_by(ShiftType.id)):
        work_types.setdefault((x.start_time, x.end_time), x)
    absence_types: dict[str, ShiftType] = {}
    for x in db.scalars(select(ShiftType).where(ShiftType.kind == "absence")
                        .order_by(ShiftType.id)):
        if (x.tzh_code or "").strip():
            absence_types.setdefault(x.tzh_code.strip(), x)
    code_colors = day_code_colors(db)
    shifts_by_emp_day: dict[tuple[int, str], list] = {}
    for row in shift_rows:
        start_local = row.start_utc.replace(tzinfo=dt.timezone.utc).astimezone(TZ)
        end_local = row.end_utc.replace(tzinfo=dt.timezone.utc).astimezone(TZ)
        day = start_local.date()
        while day <= min(end_local.date(), last):
            if first <= day <= last:
                shifts_by_emp_day.setdefault((row.employee_id, day.isoformat()), []).append(row)
            day += dt.timedelta(days=1)
    absences_by_emp_day: dict[tuple[int, str], list] = {}
    for row in absence_rows:
        start_local = row.start_utc.replace(tzinfo=dt.timezone.utc).astimezone(TZ)
        end_local = row.end_utc.replace(tzinfo=dt.timezone.utc).astimezone(TZ)
        day = start_local.date()
        while day <= min(end_local.date(), last):
            if first <= day <= last:
                absences_by_emp_day.setdefault((row.employee_id, day.isoformat()), []).append(row)
            day += dt.timedelta(days=1)

    today = now_local().date()
    dbl_cal = calendar_map(db, first, last)
    rows = []
    for emp in employees:
        periods = per_emp_periods.get(emp.id) or []
        cells = {}
        planned_total = 0.0
        work_days = 0
        for d in days:
            employed = employed_on(periods, d) if periods else True
            trow = tabel_rows.get((emp.id, d))
            sh_rows = shifts_by_emp_day.get((emp.id, d.isoformat()), [])
            ab_rows = absences_by_emp_day.get((emp.id, d.isoformat()), [])
            policy = st.day_code_policy.get(trow.code) if trow else None
            value = None
            if trow:
                value = (f"{trow.code} {fmt_hours(trow.plan_minutes)}"
                         if policy and policy.carries_hours else trow.code)
            planned = (trow.plan_minutes / 60) if trow else 0.0
            planned_total += planned
            work_days += 1 if (policy and policy.carries_hours) else 0
            shifts_brief = [_shift_brief(r) for r in sh_rows]
            absences_brief = [_absence_brief(r, kinds.get(r.kind_id),
                                             _absence_covers_day(st, tabel_rows, emp.id, r))
                              for r in ab_rows]
            cells[d.isoformat()] = {
                "tabel": None if trow is None else
                {"code": trow.code, "value": value, "hours": trow.plan_minutes / 60,
                 "source": trow.source, "note": trow.note or ""},
                "shifts": shifts_brief,
                # видимые отрезки смены (минус отсутствия) — по строке на отрезок
                "segments": _visible_segments(sh_rows, ab_rows),
                "absences": absences_brief,
                "color": _cell_color(shifts_brief, absences_brief, trow,
                                     work_types, absence_types, code_colors),
                "planned_hours": planned,
                "employed": employed,
                # день редактируем (ручная смена) ⇔ accepts_ut кода Табеля (16.3)
                "editable": bool(policy and policy.accepts_ut),
                "is_today": d == today,
                "is_future": d > today,
            }
        emp_blocks = _block_ranges(blocks_by_emp.get(emp.id, []), emp, first, last)
        rows.append({
            "employee": {
                "id": emp.id, "full_name": emp.full_name, "short_name": emp.display_name,
                "full_name_genitive": emp.full_name_genitive or suggest_genitive(emp.full_name),
                "position": emp.position, "tab_number": emp.tab_number or "",
                "schedule_group": normalize_group(emp.schedule_group) or "",
                "group_color": emp.group_color or "#8a94a6",
                "schedule_pattern": emp.schedule_pattern or "",
                "active": emp.active,
                "hired_at": emp.hired_at.isoformat() if emp.hired_at else None,
                "dismissed_at": emp.dismissed_at.isoformat() if emp.dismissed_at else None,
            },
            "blocks": emp_blocks,
            "employed_ranges": ranges_from_periods(periods, first, last),
            "cells": cells,
            "totals": {"planned_hours": round(planned_total, 2), "work_days": work_days},
        })

    from ..deps import month_name
    payload = {
        "year": year, "month": month, "month_name": month_name(month),
        "days": [{"date": d.isoformat(), "day": d.day,
                  "weekday": ["пн", "вт", "ср", "чт", "пт", "сб", "вс"][d.weekday()],
                  "is_weekend": d.weekday() >= 5, "is_today": d == today,
                  "double_scope": dbl_cal.get(d, ""),
                  "double_title": SCOPE_TITLES.get(dbl_cal.get(d, ""), "") if dbl_cal.get(d) else ""}
                 for d in days],
        "rows": rows,
        "today": today.isoformat(),
        "readonly": not principal.is_manager,
        "colors": {"out_of_block": OUT_OF_BLOCK_COLOR, "worked_off": WORKED_OFF_COLOR,
                   "inactive": INACTIVE_COLOR},
    }
    if not principal.is_manager:
        _strip_to_plan(payload)
    return payload


def _block_ranges(blocks, emp: Employee, first: dt.date, last: dt.date) -> list[dict]:
    per_group: dict[str, list[dict]] = {}
    labels: dict[str, str] = {}
    for b in blocks:
        if b.end_date and b.end_date < first:
            continue
        g = normalize_group(b.group)
        per_group.setdefault(g, []).append(
            {"start": b.start_date.isoformat(),
             "end": b.end_date.isoformat() if b.end_date else None})
        import json as _json
        try:
            lbl = (_json.loads(b.pattern_json or "{}").get("label") or "") if b.pattern_json else ""
        except ValueError:
            lbl = ""
        if lbl:
            labels[g] = lbl
    return [{"group": g, "ranges": rs, "pattern_label": labels.get(g, "")}
            for g, rs in per_group.items()] or \
        [{"group": normalize_group(emp.schedule_group), "ranges": [], "pattern_label": ""}]


# ── ручные смены и отсутствия (16.3) ──────────────────────────────────────────

class ShiftIntervalIn(BaseModel):
    start: str                       # "HH:MM"
    end: str                         # "HH:MM"
    end_next_day: bool = False
    note: str = ""


class ManualShiftIn(BaseModel):
    employee_id: int
    date: dt.date
    shifts: list[ShiftIntervalIn]
    shift_type_id: Optional[int] = None     # пресет из словаря (время берётся из него)


@router.put("/shift")
def put_manual_shift(payload: ManualShiftIn, principal: Principal = Depends(require_perm("schedule.edit")),
                     db: Session = Depends(get_db)):
    intervals = payload.shifts
    if payload.shift_type_id is not None and not intervals:
        kind = db.get(ShiftType, payload.shift_type_id)
        if kind is None or kind.kind != "work":
            raise HTTPException(status_code=422, detail="Смена словаря не найдена")
        intervals = [ShiftIntervalIn(start=kind.start_time, end=kind.end_time,
                                     end_next_day=bool(kind.overnight))]
    if not intervals:
        raise HTTPException(status_code=422, detail="Укажите интервалы смены")
    shifts = []
    for iv in intervals:
        if not (_HHMM.match(iv.start) and _HHMM.match(iv.end)):
            raise HTTPException(status_code=422, detail="Время в формате ЧЧ:ММ")
        start = _to_utc(payload.date, iv.start)
        end = _to_utc(payload.date, iv.end, iv.end_next_day)
        if end <= start:
            raise HTTPException(status_code=422,
                                detail="Конец смены должен быть позже начала (или отметьте "
                                       "«заканчивается на следующие сутки»)")
        shifts.append(Shift(str(payload.employee_id), start, end))
    ports = engine_ports(db, principal, now_utc())
    try:
        schedule_view.save_schedule_day(str(payload.employee_id), payload.date, shifts, ports)
    except (DomainError, ApplicationError) as exc:
        raise HTTPException(status_code=422, detail=domain_error_text(exc))
    for row, iv in zip(db.scalars(select(ScheduleShift).where(
            ScheduleShift.employee_id == payload.employee_id,
            ScheduleShift.date == payload.date, ScheduleShift.source == "manual")),
                     intervals, strict=False):
        row.note = iv.note[:255]
        row.updated_by = principal.user.id
    audit(db, principal, "schedule_manual_shift", f"employee:{payload.employee_id}",
          {"date": payload.date.isoformat(),
           "shifts": [[i.start, i.end, i.end_next_day] for i in intervals]})
    db.commit()
    return {"ok": True}


@router.delete("/shift")
def delete_manual_shift(employee_id: int, date: dt.date,
                        principal: Principal = Depends(require_perm("schedule.edit")),
                        db: Session = Depends(get_db)):
    """Удаление ручной смены — день возвращается к генерации из Табеля (16.3)."""
    ports = engine_ports(db, principal, now_utc())
    schedule_view.drop_manual_day(str(employee_id), date, ports)
    # генерация дня была вытеснена ручной сменой при сохранении — восстанавливаем
    schedule_view.rebuild_generated_schedule(str(employee_id),
                                             month_period(date.year, date.month), ports)
    audit(db, principal, "schedule_drop_manual", f"employee:{employee_id}",
          {"date": date.isoformat()})
    db.commit()
    return {"ok": True}


class AbsenceIn(BaseModel):
    employee_id: int
    date: dt.date
    from_time: str                   # "HH:MM"
    until_time: str                  # "HH:MM"
    next_day: bool = False
    kind_id: Optional[int] = None    # тип отсутствия из словаря (заявление, режим отметок)
    note: str = ""


@router.post("/absence")
def add_absence(payload: AbsenceIn, principal: Principal = Depends(require_perm("schedule.edit")),
                db: Session = Depends(get_db)):
    """«Пропуск/отлучка» — согласованное отсутствие (логический разрез Графика, v7/R17).
    На часы и банк не влияет (RUL-ABSENCE-ONLY-FLAGS): недостача возникает из отсутствия
    присутствия в окне и гасится банком/кодами по периоду."""
    for value in (payload.from_time, payload.until_time):
        if not _HHMM.match(value or ""):
            raise HTTPException(status_code=422, detail="Время в формате ЧЧ:ММ")
    ports = engine_ports(db, principal, now_utc())
    try:
        absence_id = schedule_view.save_absence(str(payload.employee_id), payload.date,
                                                payload.from_time, payload.until_time,
                                                payload.next_day, ports)
    except (DomainError, ApplicationError) as exc:
        raise HTTPException(status_code=422, detail=domain_error_text(exc))
    row = db.get(ScheduleAbsence, absence_id)
    if row is not None:
        row.kind_id = payload.kind_id
        row.note = payload.note[:255]
        row.created_by = principal.user.id
    audit(db, principal, "schedule_absence", f"employee:{payload.employee_id}",
          {"date": payload.date.isoformat(), "from": payload.from_time,
           "until": payload.until_time, "next_day": payload.next_day,
           "kind_id": payload.kind_id})
    db.commit()
    return {"ok": True, "id": absence_id}


@router.delete("/absence/{absence_id}")
def remove_absence(absence_id: int, principal: Principal = Depends(require_perm("schedule.edit")),
                   db: Session = Depends(get_db)):
    ports = engine_ports(db, principal, now_utc())
    schedule_view.drop_absence(absence_id, ports)
    audit(db, principal, "schedule_absence_remove", f"absence:{absence_id}", {})
    db.commit()
    return {"ok": True}


# ── отсутствие на период: коды Табеля (первичный документ) ────────────────────

class RangeAbsenceIn(BaseModel):
    employee_ids: list[int]
    first: dt.date
    last: dt.date
    shift_type_id: Optional[int] = None   # вид отсутствия из словаря (код — из его Т-13)
    code: str = ""                        # либо код Табеля напрямую (ОТ/ДО/Б/ОВ/НН…)
    note: str = ""
    workdays_only: bool = False           # только дни, где по Табелю была работа


def standard_day_hours(db: Session, st, emp, day: dt.date) -> tuple[Optional[float], str]:
    """Часы «выходного за ранее отработанные часы» — стандартное окно смен сотрудника.

    Порядок вывода (норматив от 12.10.2026):

    1. `window_of_day` — окно значения Табеля этого дня (или дня, который сотрудник
       меняет на выходной): сколько часов он обычно работает, столько и стоит в
       заявлении;
    2. `timesheet_mode` — мода длин окон в Табеле сотрудника за последние 90 дней:
       день ещё не заполнен, но норма сотрудника понятна;
    3. `default` — глобальная настройка объекта (`timeoff_default_hours`).

    Число информационное: код отсутствия не несёт плановых часов, согласованное
    отсутствие не влияет на план, часы и банк (RUL-ABSENCE-ONLY-FLAGS), недостачу дня
    гасит банк (EX-29).
    """
    from collections import Counter

    from ..schedule_helpers import DEFAULT_RULES, load_rules

    def window_hours(row) -> Optional[float]:
        if row is None:
            return None
        policy = st.day_code_policy.get(row.code)
        if policy is None or not policy.carries_hours or not row.plan_minutes:
            return None
        window = st.windows.get(f"{row.code} {fmt_hours(row.plan_minutes)}")
        return None if window is None else window.length / 60

    own = window_hours(db.scalar(select(TabelDay).where(TabelDay.employee_id == emp.id,
                                                        TabelDay.date == day)))
    if own is not None:
        return own, "window_of_day"
    # день, который меняют на выходной: ближайший предыдущий рабочий день сотрудника
    for back in range(1, 15):
        prev = window_hours(db.scalar(select(TabelDay).where(
            TabelDay.employee_id == emp.id, TabelDay.date == day - dt.timedelta(days=back))))
        if prev is not None:
            return prev, "window_of_day"
    since = day - dt.timedelta(days=90)
    lengths = [window_hours(r) for r in db.scalars(select(TabelDay).where(
        TabelDay.employee_id == emp.id, TabelDay.date >= since, TabelDay.date <= day))]
    known = [v for v in lengths if v is not None]
    if known:
        return Counter(known).most_common(1)[0][0], "timesheet_mode"
    rules = load_rules(db)
    try:
        return float(rules.get("timeoff_default_hours", DEFAULT_RULES["timeoff_default_hours"])), "default"
    except (TypeError, ValueError):
        return float(DEFAULT_RULES["timeoff_default_hours"]), "default"


def st_day(db: Session, day: dt.date):
    """Настройки периода, содержащего день (для вывода часов отсутствия)."""
    return _settings(db, month_period(day.year, day.month))


def _covers_window(st, row: Optional[TabelDay], day: dt.date,
                   start_local: dt.datetime, end_local: dt.datetime) -> bool:
    """Отсутствие покрывает рабочее окно значения Табеля дня целиком (чистая функция).

    ЕДИНСТВЕННОЕ правило «выходного за часы» (EX-29) — его используют заявление
    (absence-period) и сетка Графика (плашка на весь день); второй копии на
    фронтенде нет (дух антипаттерна 12.23)."""
    if (end_local - start_local) >= dt.timedelta(hours=23):
        return True
    if row is None:
        return False
    policy = st.day_code_policy.get(row.code)
    if policy is None or not policy.carries_hours or not row.plan_minutes:
        return False
    window = st.windows.get(f"{row.code} {fmt_hours(row.plan_minutes)}")
    if window is None:
        return False
    midnight = dt.datetime.combine(day, dt.time(0), TZ)
    win_start = midnight + dt.timedelta(minutes=window.start_minute)
    win_end = midnight + dt.timedelta(minutes=window.end_minute)
    return start_local <= win_start and end_local >= win_end


def _covers_whole_day(db: Session, st, employee_id: int, day: dt.date,
                      start_local: dt.datetime, end_local: dt.datetime) -> bool:
    """Отсутствие покрывает рабочую часть дня целиком.

    Признак «дня за часы»: сравниваем с окном значения Табеля этого дня (12 часовая
    смена 08:00–20:00 покрыта записью 08:00–20:00), а не с календарными сутками —
    иначе заявление печаталось бы с длиной записи вместо нормы смены.
    """
    row = db.scalar(select(TabelDay).where(TabelDay.employee_id == employee_id,
                                           TabelDay.date == day))
    return _covers_window(st, row, day, start_local, end_local)


def _absence_covers_day(st, tabel_rows: dict, employee_id: int,
                        row: ScheduleAbsence) -> bool:
    """whole_shift для строки сетки: окно берём у дня-якоря отсутствия (у смены
    через полночь якорь — день начала)."""
    start_local = row.start_utc.replace(tzinfo=dt.timezone.utc).astimezone(TZ)
    end_local = row.end_utc.replace(tzinfo=dt.timezone.utc).astimezone(TZ)
    day = start_local.date()
    return _covers_window(st, tabel_rows.get((employee_id, day)), day, start_local, end_local)


def _day_window_hhmm(db: Session, st, employee_id: int, day: dt.date) -> tuple[str, str, bool]:
    """Интервал согласованного отсутствия на день: окно значения Табеля, иначе сутки.

    Отсутствие покрывает смену целиком — так в Табеле остаётся явка, отсутствие снимает
    ожидание и флаги, а недостачу гасит банк (паттерн EX-29).
    """
    row = db.scalar(select(TabelDay).where(TabelDay.employee_id == employee_id,
                                           TabelDay.date == day))
    if row is not None:
        policy = st.day_code_policy.get(row.code)
        if policy is not None and policy.carries_hours:
            value = f"{row.code} {fmt_hours(row.plan_minutes)}" if row.plan_minutes else row.code
            window = st.windows.get(value)
            if window is not None:
                return (f"{window.start_minute // 60:02d}:{window.start_minute % 60:02d}",
                        f"{window.end_minute // 60:02d}:{window.end_minute % 60:02d}", False)
    return "00:00", "00:00", True


def _range_absence_as_schedule(payload: RangeAbsenceIn, kind: ShiftType, principal: Principal,
                               db: Session) -> dict:
    """Вид отсутствия БЕЗ кода Табеля: пишем согласованные отсутствия в График.

    Норматив от 12.10.2026: «выходной за ранее отработанные часы» и «отпросился» —
    данные Графика, а не код первичного документа. В Табеле остаётся явка («Я»/«К»),
    отсутствие не меняет план (RUL-ABSENCE-ONLY-FLAGS), недостачу гасит банк часов —
    то есть ранее накопленные переработки (EX-29).
    """
    st = _settings(db, month_period(payload.first.year, payload.first.month))
    ports = engine_ports(db, principal, now_utc())
    written, skipped, errors = 0, 0, []
    for employee_id in payload.employee_ids:
        periods = periods_of(db, employee_id)
        day = payload.first
        while day <= payload.last:
            if not employed_on(periods, day):
                skipped += 1
                day += dt.timedelta(days=1)
                continue
            trow = db.scalar(select(TabelDay).where(TabelDay.employee_id == employee_id,
                                                    TabelDay.date == day))
            carries = bool(trow is not None and st.day_code_policy.get(trow.code)
                           and st.day_code_policy[trow.code].carries_hours)
            if payload.workdays_only and not carries:
                skipped += 1
                day += dt.timedelta(days=1)
                continue
            start_hhmm, end_hhmm, next_day = _day_window_hhmm(db, st, employee_id, day)
            try:
                absence_id = schedule_view.save_absence(str(employee_id), day, start_hhmm,
                                                        end_hhmm, next_day, ports)
            except (DomainError, ApplicationError) as exc:
                errors.append({"employee_id": employee_id, "date": day.isoformat(),
                               "error": domain_error_text(exc)})
                day += dt.timedelta(days=1)
                continue
            row = db.get(ScheduleAbsence, absence_id)
            if row is not None:
                row.kind_id = kind.id
                row.note = payload.note[:255]
                row.created_by = principal.user.id
            written += 1
            day += dt.timedelta(days=1)
    audit(db, principal, "schedule_range_absence", f"{payload.first}..{payload.last}",
          {"employees": payload.employee_ids, "kind": kind.code, "written": written,
           "mode": "schedule_absence"})
    db.commit()
    return {"ok": not errors, "written": written, "skipped": skipped, "errors": errors[:20],
            "mode": "schedule_absence"}


@router.post("/range-absence")
def range_absence(payload: RangeAbsenceIn, principal: Principal = Depends(require_perm("schedule.edit")),
                  db: Session = Depends(get_db)):
    """«Назначить отсутствие» на период.

    Два способа — по виду отсутствия из словаря:

    * **есть код Табеля** (ОТ, ДО, Б, ОВ, НН…) — отсутствие на весь день это код
      первичного документа; График дня перегенерируется (смены в эти дни нет →
      сотрудник не ожидается, флагов нет);
    * **кода Табеля нет** («Выходной за часы», «Отпросился») — пишется согласованное
      отсутствие в График, Табель не трогается: явка остаётся, недостачу гасит банк
      (норматив от 12.10.2026, паттерн EX-29).

    Период может пересекать границы месяцев.
    """
    if payload.last < payload.first:
        raise HTTPException(status_code=422, detail="Дата окончания раньше даты начала")
    if (payload.last - payload.first).days > 366:
        raise HTTPException(status_code=422,
                            detail="Период отсутствия не может быть больше 366 дней")
    if len(payload.employee_ids) > 500:
        raise HTTPException(status_code=422, detail="Слишком много сотрудников (максимум 500)")
    from ..engine.application import timesheet_edit
    code = payload.code.strip()
    kind = db.get(ShiftType, payload.shift_type_id) if payload.shift_type_id else None
    if payload.shift_type_id and kind is None:
        raise HTTPException(status_code=422, detail="Вид отсутствия не найден")
    if not code and kind is not None:
        code = (kind.tzh_code or "").strip()
    if not code and kind is not None:
        return _range_absence_as_schedule(payload, kind, principal, db)
    st = _settings(db, month_period(payload.first.year, payload.first.month))
    policy = st.day_code_policy.get(code)
    if policy is None:
        raise HTTPException(status_code=422,
                            detail=f"Код «{code or '—'}» отсутствует в CFG-DAY-CODES "
                                   "(у вида отсутствия не задан код Т-13)")
    if policy.carries_hours:
        raise HTTPException(status_code=422,
                            detail="Код несёт плановые часы — это не отсутствие")
    ports = engine_ports(db, principal, now_utc())
    written, skipped, errors = 0, 0, []
    for employee_id in payload.employee_ids:
        periods = periods_of(db, employee_id)
        day = payload.first
        while day <= payload.last:
            if not employed_on(periods, day):
                skipped += 1
                day += dt.timedelta(days=1)
                continue
            if payload.workdays_only:
                trow = db.scalar(select(TabelDay).where(TabelDay.employee_id == employee_id,
                                                        TabelDay.date == day))
                carries = (trow is not None and st.day_code_policy.get(trow.code)
                           and st.day_code_policy[trow.code].carries_hours)
                if not carries:
                    skipped += 1
                    day += dt.timedelta(days=1)
                    continue
            try:
                timesheet_edit.save_timesheet_day(str(employee_id), day, code, None, ports)
                row = db.scalar(select(TabelDay).where(TabelDay.employee_id == employee_id,
                                                       TabelDay.date == day))
                if row is not None and payload.note:
                    row.note = payload.note[:255]
                written += 1
            except (DomainError, ApplicationError) as exc:
                errors.append({"employee_id": employee_id, "date": day.isoformat(),
                               "error": domain_error_text(exc)})
            day += dt.timedelta(days=1)
    audit(db, principal, "schedule_range_absence", f"{payload.first}..{payload.last}",
          {"employees": payload.employee_ids, "code": code, "written": written})
    db.commit()
    return {"ok": not errors, "written": written, "skipped": skipped, "errors": errors[:20]}


@router.get("/timeoff-hours")
def timeoff_hours(employee_id: int, date: dt.date,
                  principal: Principal = Depends(require_perm("schedule.view")), db: Session = Depends(get_db)):
    """Часы «выходного за ранее отработанные часы» для сотрудника и дня.

    Значение выводится, а не задаётся константой (норматив от 12.10.2026):
    стандартное окно смен сотрудника из Табеля → мода окон за 90 дней → глобальный
    fallback. `hours_source` показывает, какое правило сработало.
    """
    emp = db.get(Employee, employee_id)
    if emp is None:
        raise HTTPException(status_code=404, detail="Сотрудник не найден")
    st = _settings(db, month_period(date.year, date.month))
    value, source = standard_day_hours(db, st, emp, date)
    return {"employee_id": employee_id, "date": date.isoformat(), "hours": value,
            "hours_source": source}


def _interval_info(db: Session, emp: Optional[Employee], row: ScheduleAbsence) -> Optional[dict]:
    """Интервал согласованного отсутствия как фрагмент заявления.

    None — запись без вида словаря либо у вида нет заявления (doc_type пуст):
    в печатный период такой интервал не входит. Часы: «вся смена» — стандартное
    окно смен сотрудника (норматив 12.10.2026, EX-29), интервал — длина записи."""
    if not row.kind_id:
        return None
    kind = db.get(ShiftType, row.kind_id)
    if kind is None or not kind.doc_type:
        return None
    start_local = row.start_utc.replace(tzinfo=dt.timezone.utc).astimezone(TZ)
    end_local = row.end_utc.replace(tzinfo=dt.timezone.utc).astimezone(TZ)
    day = start_local.date()
    day_settings = st_day(db, day)
    whole_shift = _covers_whole_day(db, day_settings, row.employee_id, day,
                                    start_local, end_local)
    if whole_shift:
        hours_value, hours_source = standard_day_hours(db, day_settings, emp, day)
    else:
        hours_value = round((end_local - start_local).total_seconds() / 3600, 2)
        hours_source = "absence_interval"
    return {"id": row.id,
            "start": day.isoformat(),
            "end": (end_local - dt.timedelta(seconds=1)).date().isoformat(),
            "from_time": start_local.strftime("%H:%M"),
            "until_time": end_local.strftime("%H:%M"),
            "hours": hours_value, "hours_source": hours_source,
            "whole_shift": whole_shift, "kind": _kind_brief(kind),
            "note": row.note or "",
            "doc_type": doc_type_of_absence(db, kind, whole_shift)}


@router.get("/absence-period")
def absence_period(employee_id: int, date: dt.date, absence_id: Optional[int] = None,
                   principal: Principal = Depends(require_perm("schedule.view")),
                   db: Session = Depends(get_db)):
    """Непрерывный период одного отсутствия вокруг даты — для печати заявления.

    Источник — Табель (коды отсутствий) либо записи согласованного отсутствия
    («отпросился»): заявление находит весь период, даже если он пересекает границы
    месяцев.

    Отлучек в одном дне может быть несколько (решение владельца 16.10.2026):
    по умолчанию заявление ОДНО на все интервалы дня (`intervals`, `intervals_text`
    — перечень «с 14:00 до 16:00 (2 ч) и с 18:00 до 20:00 (2 ч)», `hours` — сумма);
    с `absence_id` — заявление по одному выбранному интервалу."""
    emp = db.get(Employee, employee_id)
    employee_payload = None if emp is None else {
        "id": emp.id, "full_name": emp.full_name, "short_name": emp.display_name,
        "full_name_genitive": emp.full_name_genitive or suggest_genitive(emp.full_name),
        "position": emp.position, "tab_number": emp.tab_number or "",
        "nationality": emp.nationality or "", "subdivision": emp.subdivision or "",
        "department": emp.department.name if emp.department else "",
    }
    # 1) согласованные отсутствия дня («отпросился») — записи Графика с типом-заявлением
    day_lo = dt.datetime.combine(date, dt.time(0), TZ).astimezone(dt.timezone.utc)
    day_hi = day_lo + dt.timedelta(days=1)
    day_rows = list(db.scalars(select(ScheduleAbsence).where(
        ScheduleAbsence.employee_id == employee_id, ScheduleAbsence.start_utc < day_hi,
        ScheduleAbsence.end_utc > day_lo).order_by(ScheduleAbsence.start_utc)))
    intervals = [iv for iv in (_interval_info(db, emp, r) for r in day_rows) if iv]
    if absence_id is not None:
        intervals = [iv for iv in intervals if iv["id"] == absence_id]
        if not intervals:
            raise HTTPException(status_code=404,
                                detail="Интервал отсутствия не найден в этот день")
    if intervals:
        first, last = intervals[0], intervals[-1]
        total = round(sum(iv["hours"] or 0 for iv in intervals), 2)
        one = len(intervals) == 1
        doc_types = {iv["doc_type"] for iv in intervals}
        if len(doc_types) == 1:
            doc_type = doc_types.pop()
        else:
            # смешанные точки ввода (вся смена + интервал): перечень интервалов
            # печатается заявлением «об отсутствии на рабочем месте»
            doc_type = (DOC_PARTIAL_SHIFT if _statement_kind_exists(db, DOC_PARTIAL_SHIFT)
                        else first["doc_type"])
        intervals_text = " и ".join(
            f"с {iv['from_time']} до {iv['until_time']} ({hours_text(iv['hours'] or 0)} ч)"
            for iv in intervals)
        return {"employee": employee_payload,
                "source": "schedule_absence",
                "whole_shift": all(iv["whole_shift"] for iv in intervals),
                "doc_type": doc_type,
                "start": first["start"], "end": last["end"], "days": 1,
                "shift": first["kind"], "note": first["note"],
                "hours": total,
                "hours_source": first["hours_source"] if one else "combined",
                "partial": {"from_time": first["from_time"],
                            "until_time": last["until_time"],
                            "hours": total,
                            "hours_source": first["hours_source"] if one else "combined",
                            "kind": first["kind"]},
                "intervals": intervals, "intervals_text": intervals_text,
                "notes": sorted({iv["note"] for iv in intervals if iv["note"]})}
    # 2) отсутствие на период — непрерывный диапазон одного кода Табеля
    rows = {r.date: r for r in db.scalars(select(TabelDay).where(
        TabelDay.employee_id == employee_id).order_by(TabelDay.date))}
    anchor = rows.get(date)
    if anchor is None:
        raise HTTPException(status_code=404, detail="В эту дату нет назначенного отсутствия")
    start = end = date
    d = date - dt.timedelta(days=1)
    while d in rows and rows[d].code == anchor.code:
        start = d
        d -= dt.timedelta(days=1)
    d = date + dt.timedelta(days=1)
    while d in rows and rows[d].code == anchor.code:
        end = d
        d += dt.timedelta(days=1)
    kind = db.scalar(select(ShiftType).where(ShiftType.tzh_code == anchor.code,
                                             ShiftType.kind == "absence",
                                             ShiftType.doc_type != "").limit(1))
    if kind is None:
        raise HTTPException(status_code=404,
                            detail=f"Для кода «{anchor.code}» не настроен тип заявления")
    st = _settings(db, month_period(start.year, start.month))
    hours_value, hours_source = standard_day_hours(db, st, emp, start)
    return {"employee": employee_payload,
            "source": "tabel_code", "whole_shift": True,
            "doc_type": doc_type_of_absence(db, kind, True),
            "start": start.isoformat(), "end": end.isoformat(),
            "days": (end - start).days + 1, "shift": _kind_brief(kind),
            "note": anchor.note or "", "partial": None,
            "hours": hours_value, "hours_source": hours_source,
            "notes": sorted({r.note for d_, r in rows.items() if start <= d_ <= end and r.note})}


def _kind_brief(kind: ShiftType) -> dict:
    return {"id": kind.id, "code": kind.code, "name": kind.name, "kind": kind.kind,
            "color": kind.color, "doc_type": kind.doc_type or "",
            "display_code": kind.display_code, "tzh_code": kind.tzh_code}


# ── циклы: инструмент заполнения Табеля (решение владельца) ───────────────────

class PatternIn(BaseModel):
    employee_ids: list[int]
    start_date: dt.date
    end_date: dt.date
    work_shift_id: int
    off_shift_id: Optional[int] = None
    pattern: str = "2/2"
    cycle: str = ""
    off_weekdays: list[int] = [5, 6]
    offset: int = 0
    only_empty: bool = False
    sync_blocks: bool = True


@router.get("/patterns")
def get_patterns(principal: Principal = Depends(require_perm("schedule.edit"))):
    return {"patterns": pattern_list(), "groups": GROUP_ORDER}


@router.post("/fill-pattern")
def fill_pattern(payload: PatternIn, principal: Principal = Depends(require_perm("schedule.edit")),
                 db: Session = Depends(get_db)):
    """Шаблон блока/цикла: сохраняет назначение блока и ЗАПОЛНЯЕТ ТАБЕЛЬ по циклу;
    График перегенерируется из Табеля (16.3)."""
    work = db.get(ShiftType, payload.work_shift_id)
    if work is None:
        raise HTTPException(status_code=422, detail="Смена шаблона не найдена")
    try:
        validate_pattern(payload.pattern, payload.cycle)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    emps = [db.get(Employee, i) for i in payload.employee_ids]
    emps = [e for e in emps if e is not None and e.deleted_at is None]
    if payload.sync_blocks:
        for emp in emps:
            _assign_block(db, emp, payload, work)
    res = _fill_and_rebuild(db, emps, payload.start_date, payload.end_date,
                            payload.only_empty, principal)
    audit(db, principal, "schedule_fill_pattern",
          f"{payload.start_date}..{payload.end_date}",
          {"employees": payload.employee_ids, "pattern": payload.pattern,
           "cycle": payload.cycle, **{k: res[k] for k in ("written", "skipped")}})
    db.commit()
    return res


def _assign_block(db: Session, emp: Employee, payload: PatternIn, work: ShiftType) -> None:
    import json
    group = normalize_group(emp.schedule_group) or "Другие смены"
    kind = group_kind(group)
    if payload.cycle or payload.pattern == "custom":
        cycle = (payload.cycle or "").strip()
        if cycle and not re.fullmatch(r"\d+\s*/\s*\d+", cycle):
            cycle = "custom:" + cycle          # произвольная последовательность «ВВРРРРВВ»
        pattern_json = json.dumps({"kind": "cycle",
                                   "cycle": cycle or payload.pattern,
                                   "shift_code": work.code,
                                   "anchor": payload.start_date.isoformat(),
                                   "label": f"{payload.cycle or payload.pattern} {work.display_code}"},
                                  ensure_ascii=False)
    elif kind == "week5" or payload.pattern == "5/2":
        pattern_json = json.dumps({"kind": "week5", "off_weekdays": payload.off_weekdays,
                                   "shift_code": work.code,
                                   "label": f"{payload.pattern}, вых. "
                                            f"{'/'.join(str(x) for x in payload.off_weekdays)}"},
                                  ensure_ascii=False)
    else:
        pattern_json = json.dumps({"kind": "base", "invert": bool(payload.offset)},
                                  ensure_ascii=False)
    existing = db.scalar(select(BlockAssignment).where(
        BlockAssignment.employee_id == emp.id, BlockAssignment.group == group,
        BlockAssignment.start_date == payload.start_date).limit(1))
    if existing is not None:
        existing.pattern_json = pattern_json          # повторное заполнение той же даты
        for ba in db.scalars(select(BlockAssignment).where(
                BlockAssignment.employee_id == emp.id, BlockAssignment.end_date.is_(None),
                BlockAssignment.id != existing.id)):
            ba.end_date = payload.start_date - dt.timedelta(days=1)
    else:
        for ba in db.scalars(select(BlockAssignment).where(
                BlockAssignment.employee_id == emp.id, BlockAssignment.end_date.is_(None))):
            ba.end_date = payload.start_date - dt.timedelta(days=1)
        db.add(BlockAssignment(employee_id=emp.id, group=group,
                               start_date=payload.start_date, end_date=None,
                               pattern_json=pattern_json))
    emp.schedule_pattern = payload.cycle or payload.pattern
    emp.schedule_anchor = payload.start_date
    db.flush()


def _fill_and_rebuild(db: Session, emps: list[Employee], first: dt.date, last: dt.date,
                      only_empty: bool, principal: Principal) -> dict:
    st = _settings(db, month_period(first.year, first.month))
    res = fill_timesheet_from_cycles(db, emps, first, last, st, overwrite=not only_empty,
                                     actor_id=principal.user.id)
    ports = engine_ports(db, principal, now_utc())
    months = _months_between(first, last)
    for emp in emps:
        for m in months:
            rebuild_generated_schedule(str(emp.id), m, ports)
    return res


def _months_between(first: dt.date, last: dt.date) -> list[Period]:
    out = []
    y, m = first.year, first.month
    while (y, m) <= (last.year, last.month):
        out.append(month_period(y, m))
        y, m = (y + (m == 12), m % 12 + 1)
    return out


class ContinueIn(BaseModel):
    employee_ids: Optional[list[int]] = None
    year: int
    month: int
    only_empty: bool = True


@router.post("/continue")
def continue_schedule(payload: ContinueIn, principal: Principal = Depends(require_perm("schedule.edit")),
                      db: Session = Depends(get_db)):
    """«Продолжить график»: заполняет Табель следующего месяца по циклам и
    перегенерирует График."""
    first, last = _month_bounds(payload.year, payload.month)
    emps = [db.get(Employee, i) for i in (payload.employee_ids or [])]
    emps = [e for e in emps if e is not None and e.deleted_at is None]
    if not emps:
        emps = _visible_employees(db, first, last)[0]
    res = _fill_and_rebuild(db, emps, first, last, payload.only_empty, principal)
    audit(db, principal, "schedule_continue", f"{payload.year}-{payload.month:02d}",
          {"employees": len(emps), **{k: res[k] for k in ("written", "skipped")}})
    db.commit()
    return res


class CopyIn(BaseModel):
    employee_ids: Optional[list[int]] = None
    src_year: int
    src_month: int
    dst_year: int
    dst_month: int
    only_empty: bool = True


@router.post("/copy")
def copy_period(payload: CopyIn, principal: Principal = Depends(require_perm("schedule.edit")),
                db: Session = Depends(get_db)):
    """Копирование ячеек Табеля месяц-в-месяц (по дням месяца); закрытые периоды
    не трогаются (16.5)."""
    from ..engine.application import timesheet_edit
    src_first, src_last = _month_bounds(payload.src_year, payload.src_month)
    dst_first, dst_last = _month_bounds(payload.dst_year, payload.dst_month)
    ports = engine_ports(db, principal, now_utc())
    emps = [db.get(Employee, i) for i in (payload.employee_ids or [])]
    emps = [e for e in emps if e is not None and e.deleted_at is None] or \
        _visible_employees(db, dst_first, dst_last)[0]
    src_rows = {(r.employee_id, r.date): r for r in db.scalars(select(TabelDay).where(
        TabelDay.employee_id.in_([e.id for e in emps]), TabelDay.date >= src_first,
        TabelDay.date <= src_last))} if emps else {}
    copied, skipped, errors = 0, 0, []
    for emp in emps:
        if timesheet_edit.lock_state(str(emp.id), dst_first, ports) == "locked":
            skipped += 1
            continue
        for offset in range(min((src_last - src_first).days, (dst_last - dst_first).days) + 1):
            src = src_rows.get((emp.id, src_first + dt.timedelta(days=offset)))
            if src is None:
                continue
            dst_day = dst_first + dt.timedelta(days=offset)
            dst_row = db.scalar(select(TabelDay).where(TabelDay.employee_id == emp.id,
                                                       TabelDay.date == dst_day))
            if dst_row is not None and payload.only_empty:
                skipped += 1
                continue
            try:
                timesheet_edit.save_timesheet_day(
                    str(emp.id), dst_day, src.code,
                    src.plan_minutes or None, ports)
                copied += 1
            except (DomainError, ApplicationError) as exc:
                errors.append({"employee_id": emp.id, "date": dst_day.isoformat(),
                               "error": domain_error_text(exc)})
    audit(db, principal, "schedule_copy",
          f"{payload.src_year}-{payload.src_month:02d}→{payload.dst_year}-{payload.dst_month:02d}",
          {"copied": copied})
    db.commit()
    return {"ok": not errors, "copied": copied, "skipped": skipped, "errors": errors[:20]}


class ClearMonthIn(BaseModel):
    employee_ids: Optional[list[int]] = None
    year: int
    month: int
    manual_too: bool = False


@router.post("/clear-month")
def clear_month(payload: ClearMonthIn, principal: Principal = Depends(require_perm("schedule.edit")),
                db: Session = Depends(get_db)):
    """Очистка Табеля месяца (и сгенерированного Графика); manual-строки — только с
    флагом manual_too. Закрытые периоды не трогаются."""
    from ..engine.application import timesheet_edit
    first, last = _month_bounds(payload.year, payload.month)
    ports = engine_ports(db, principal, now_utc())
    emps = [db.get(Employee, i) for i in (payload.employee_ids or [])]
    emps = [e for e in emps if e is not None and e.deleted_at is None] or \
        _visible_employees(db, first, last)[0]
    cleared, skipped = 0, 0
    for emp in emps:
        if timesheet_edit.lock_state(str(emp.id), first, ports) == "locked":
            skipped += 1
            continue
        q = select(TabelDay).where(TabelDay.employee_id == emp.id, TabelDay.date >= first,
                                   TabelDay.date <= last)
        if not payload.manual_too:
            q = q.where(TabelDay.source != "manual")
        rows = list(db.scalars(q))
        for row in rows:
            db.delete(row)
            cleared += 1
        db.flush()      # rebuild читает ту же сессию (autoflush off) — не должен видеть удалённое
        schedule_view.rebuild_generated_schedule(str(emp.id), month_period(payload.year,
                                                                          payload.month), ports)
    audit(db, principal, "schedule_clear_month", f"{payload.year}-{payload.month:02d}",
          {"cleared": cleared})
    db.commit()
    return {"ok": True, "cleared": cleared, "skipped_employees": skipped}


@router.post("/reset-month")
def reset_month(payload: ClearMonthIn, principal: Principal = Depends(require_perm("schedule.edit")),
                db: Session = Depends(get_db)):
    """«Сбросить месяц» в Графике (решения владельца 16.10.2026): убрать РУЧНЫЕ
    смены и согласованные ОТСУТСТВИЯ месяца (отлучки-интервалы и «выходные за
    часы») и перегенерировать График из Табеля — слой 1 возвращается к чистому
    плану из слоя 0.

    Отличие от clear-month: Табель НЕ трогается (он сбрасывается из «Слой 0 ·
    Табель»). Отсутствие принадлежит месяцу по перекрытию интервала (как в сетке:
    отлучка через полночь на границе месяцев уходит в сброс своего месяца).
    Записи удаляются с аудитом; напечатанные ранее заявления остаются в истории
    журнала. Закрытые периоды не трогаются (RUL-CLOSED-PERIOD-IMMUTABLE)."""
    from ..engine.application import timesheet_edit
    if not (1 <= payload.month <= 12):
        raise HTTPException(status_code=422, detail="month должен быть 1..12")
    if not (1 <= payload.year <= 9999):
        raise HTTPException(status_code=422, detail="year должен быть 1..9999")
    first, last = _month_bounds(payload.year, payload.month)
    ports = engine_ports(db, principal, now_utc())
    emps = [db.get(Employee, i) for i in (payload.employee_ids or [])]
    emps = [e for e in emps if e is not None and e.deleted_at is None] or \
        _visible_employees(db, first, last)[0]
    lo_utc = dt.datetime.combine(first, dt.time(0), TZ).astimezone(dt.timezone.utc)
    hi_utc = dt.datetime.combine(last + dt.timedelta(days=1), dt.time(0),
                                 TZ).astimezone(dt.timezone.utc)
    cleared, absences_cleared, skipped = 0, 0, 0
    for emp in emps:
        if timesheet_edit.lock_state(str(emp.id), first, ports) == "locked":
            skipped += 1
            continue
        manual = list(db.scalars(select(ScheduleShift).where(
            ScheduleShift.employee_id == emp.id, ScheduleShift.source == "manual",
            ScheduleShift.date >= first, ScheduleShift.date <= last)))
        for row in manual:
            db.delete(row)
        cleared += len(manual)
        absences = list(db.scalars(select(ScheduleAbsence).where(
            ScheduleAbsence.employee_id == emp.id, ScheduleAbsence.start_utc < hi_utc,
            ScheduleAbsence.end_utc > lo_utc)))
        for row in absences:
            db.delete(row)
        absences_cleared += len(absences)
        db.flush()      # rebuild не должен видеть удалённые строки (autoflush off)
        schedule_view.rebuild_generated_schedule(str(emp.id),
                                                 month_period(payload.year, payload.month),
                                                 ports)
    audit(db, principal, "schedule_reset_month", f"{payload.year}-{payload.month:02d}",
          {"manual_shifts_cleared": cleared, "absences_cleared": absences_cleared})
    db.commit()
    return {"ok": True, "cleared": cleared, "absences_cleared": absences_cleared,
            "skipped_employees": skipped}


# ── печать и экспорт ──────────────────────────────────────────────────────────

@router.get("/xlsx")
def export_xlsx(year: int, month: int, principal: Principal = Depends(require_manager),
                db: Session = Depends(get_db)):
    """График в Excel: тот же вид ячейки, что на экране (16.10.2026) — время смены
    (отрезки по строкам), плашка «Выходной за часы», цвет из словаря «Смены и
    отсутствия»; итоговой колонки «Часов» нет (решение владельца)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from ..grid_docx import cell_plan_text, tint_hex
    data = get_grid(year, month, principal, db)
    wb = Workbook()
    ws = wb.active
    ws.title = f"График {month:02d}.{year}"
    days = [d["date"] for d in data["days"]]
    ws.append(["Сотрудник"] + [d[8:] for d in days])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in data["rows"]:
        cells = []
        for iso in days:
            cell = row["cells"].get(iso) or {}
            cells.append(safe_sheet_cell(cell_plan_text(cell)))
        ws.append([safe_sheet_cell(row["employee"]["full_name"]), *cells])
    center = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for ri, row in enumerate(data["rows"], start=2):
        for ci, iso in enumerate(days, start=2):
            cell = row["cells"].get(iso) or {}
            wsc = ws.cell(row=ri, column=ci)
            wsc.alignment = center
            tint = tint_hex(cell.get("color"))
            if tint:
                wsc.fill = PatternFill("solid", start_color=tint, end_color=tint)
        ws.cell(row=ri, column=1).alignment = Alignment(vertical="center")
    ws.column_dimensions["A"].width = 32
    for ci in range(2, len(days) + 2):
        ws.column_dimensions[get_column_letter(ci)].width = 13.5
    import io
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    from fastapi.responses import StreamingResponse
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="schedule_{year}_{month:02d}.xlsx"'})


@router.get("/pdf")
def export_pdf(year: int, month: int, principal: Principal = Depends(require_perm("schedule.view")),
               db: Session = Depends(get_db)):
    """График на месяц в PDF: та же таблица, что в «Печать»/Excel. Собирается через
    DOCX → PDF (LibreOffice или встроенный движок); без движка — 501 с подсказкой."""
    from fastapi.responses import Response

    from ..grid_docx import pdf_unavailable_detail, schedule_docx
    from ..pdf_render import docx_to_pdf
    data = get_grid(year, month, principal, db)
    pdf, how = docx_to_pdf(schedule_docx(data), f"grafik_{year}-{month:02d}.pdf")
    if pdf is None:
        raise HTTPException(status_code=501, detail=pdf_unavailable_detail(how))
    audit(db, principal, "schedule_pdf", f"{year}-{month:02d}", {"rows": len(data["rows"])})
    db.commit()
    return Response(content=pdf, media_type="application/pdf",
                    headers={"Content-Disposition":
                             f'attachment; filename="grafik_{year}-{month:02d}.pdf"'})
