"""Сквозная проверка всех 11 требований заказчика на чистой демо-базе (TestClient)."""
from __future__ import annotations

import datetime as dt
import io
import os
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

tmp_db = os.path.join(tempfile.gettempdir(), "tt_verify.db")
if os.path.exists(tmp_db):
    os.remove(tmp_db)
os.environ["DATABASE_URL"] = f"sqlite:///{tmp_db}"
os.environ["SECRET_KEY"] = "verify-secret"

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app, init_db  # noqa: E402

init_db()
client = TestClient(app)
results: list[tuple[int, str, bool, str]] = []


def check(num: int, title: str, ok: bool, evidence: str = "") -> None:
    results.append((num, title, ok, evidence))
    print(f"{'✓' if ok else '✗'} [{num:>2}] {title}" + (f" — {evidence}" if evidence else ""))


def login(username: str) -> TestClient:
    c = TestClient(app)
    r = c.post("/api/auth/login", json={"username": username, "password": "demo1234"})
    assert r.status_code == 200, r.text
    return c


admin = login("admin")
manager = login("gromova")
employee = login("ivanov")
today = dt.date.today()
YM = {"year": today.year, "month": today.month}

# ── 1. светлая тема ──
css = client.get("/static/css/app.css").text
light = "--bg: #eef1f6" in css and "--panel: #ffffff" in css and "#0f1520" not in css
check(1, "Интерфейс в светлых тонах", light, "app.css: --bg #eef1f6, панели #ffffff")

# ── 2. произвольные циклы N/M ──
emps = manager.get("/api/employees").json()
types = manager.get("/api/shift-types").json()
day12 = next(t for t in types if t["code"] == "DAY12")
off = next(t for t in types if t["code"] == "OFF")
first = today.replace(day=1)
r = manager.post("/api/schedule/fill-pattern", json={
    "employee_ids": [emps[-1]["id"]], "start_date": first.isoformat(),
    "end_date": (first + dt.timedelta(days=6)).isoformat(),
    "work_shift_id": day12["id"], "off_shift_id": off["id"], "pattern": "4/3"})
grid = manager.get("/api/schedule", params=YM).json()
row = next(x for x in grid["rows"] if x["employee"]["id"] == emps[-1]["id"])
seq = [row["cells"][(first + dt.timedelta(days=i)).isoformat()]["shift"]["code"] for i in range(7)]
check(2, "Произвольные графики 2/2, 3/3, 4/3…", r.status_code == 200 and seq == ["DAY12"] * 4 + ["OFF"] * 3,
      f"4/3 → {' '.join(seq)}")

# ── 3. произвольные часы работы ──
r = admin.post("/api/shift-types", json={
    "code": "DAY8_17", "name": "08:00–17:00 (9 ч)", "display_code": "08–17", "tzh_code": "Я",
    "kind": "work", "start_time": "08:00", "end_time": "17:00", "color": "#4cae52"})
listed = {t["code"]: t for t in admin.get("/api/shift-types").json()}
ok3 = listed.get("DAY8_17", {}).get("planned_hours") == 9.0     # 08:00–17:00 уже в словаре
r2 = admin.post("/api/shift-types", json={
    "code": "DAY7_16", "name": "07:00–16:00 (9 ч)", "display_code": "07–16", "tzh_code": "Я",
    "kind": "work", "start_time": "07:00", "end_time": "16:00", "color": "#2bb673"})
ok3 = ok3 and r2.status_code == 200 and r2.json()["planned_hours"] == 9.0
check(3, "Произвольные часы 08:00–17:00, 09:00–18:00…", ok3, "словарь смен: POST /api/shift-types, 9 ч без перерыва")

# ── 4. переходы смена-смена (произвольный цикл) ──
r = manager.post("/api/schedule/fill-pattern", json={
    "employee_ids": [emps[-1]["id"]], "start_date": first.isoformat(),
    "end_date": (first + dt.timedelta(days=7)).isoformat(),
    "work_shift_id": day12["id"], "off_shift_id": off["id"], "pattern": "custom", "cycle": "ВВРРРРВВ"})
grid = manager.get("/api/schedule", params=YM).json()
row = next(x for x in grid["rows"] if x["employee"]["id"] == emps[-1]["id"])
seq = [row["cells"][(first + dt.timedelta(days=i)).isoformat()]["shift"]["code"] for i in range(8)]
check(4, "Переходы: ВВРРРРВВ / РРВВВВРР и т.п.", r.status_code == 200 and
      seq == ["OFF", "OFF", "DAY12", "DAY12", "DAY12", "DAY12", "OFF", "OFF"], f"{' '.join(seq)}")

# ── 5. 5/2 с выбором выходных ──
d0 = first
while d0.weekday() != 0:
    d0 += dt.timedelta(days=1)
r = manager.post("/api/schedule/fill-pattern", json={
    "employee_ids": [emps[-1]["id"]], "start_date": d0.isoformat(),
    "end_date": (d0 + dt.timedelta(days=6)).isoformat(),
    "work_shift_id": day12["id"], "off_shift_id": off["id"], "pattern": "5/2", "off_weekdays": [0, 1]})
grid = manager.get("/api/schedule", params=YM).json()
row = next(x for x in grid["rows"] if x["employee"]["id"] == emps[-1]["id"])
seq = [row["cells"][(d0 + dt.timedelta(days=i)).isoformat()]["shift"]["code"] for i in range(7)]
check(5, "5/2 с выбором выходных (пн+вт)", r.status_code == 200 and
      seq == ["OFF", "OFF", "DAY12", "DAY12", "DAY12", "DAY12", "DAY12"], f"пн..вс: {' '.join(seq)}")

# ── 6. описание переработки ──
punches = employee.get(f"/api/punches?year={today.year}&month={today.month}").json()
out_p = next((p for p in punches if p["kind"] == "OUT"), None)
ok6 = False
ev6 = "нет отметки ухода в демо-данных"
if out_p:
    r = employee.put(f"/api/punches/{out_p['id']}/note",
                     json={"note": "Проверка: оставался с Громовой, готовили зал"})
    ts = manager.get("/api/timesheet", params={**YM, "employee_id": employee.get("/api/auth/me").json()["user"]["employee_id"]}).json()
    notes = [d["ot_note"] for rr in ts["rows"] for d in rr["days"] if d["ot_note"]]
    ok6 = r.status_code == 200 and any("готовили зал" in n for n in notes)
    ev6 = "PUT /api/punches/{id}/note → поле ot_note в табеле и лист Excel «Переработки»"
check(6, "Описание переработки: с кем работал, что делал", ok6, ev6)

# ── 7. без перерывов и денег ──
t0 = next(t for t in manager.get("/api/shift-types").json() if t["code"] == "DAY12")
rules = manager.get("/api/settings").json()["rules"]
ts = manager.get("/api/timesheet", params=YM).json()
no_break = "break_minutes" not in t0 and t0["planned_hours"] == 12.0
no_money = "break_placement" not in rules and "ot_split_hours" not in rules \
    and "ot15" not in ts["grand"] and "hourly_rate" not in (emps[0] or {})
check(7, "Перерывы и денежные расчёты исключены", no_break and no_money,
      "08:00–20:00 = 12 ч; в правилах и табеле нет перерывов/ставок/×1,5")

# ── 8. Excel-табель сеткой «ДЯ/ДН» ──
raw = manager.get("/api/timesheet/xlsx", params=YM).content
ok8 = raw[:2] == b"PK"
ev8 = ""
if ok8:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw))
    ws = wb.active
    header = [c.value for c in ws[4]]
    cells = [c.value for row in ws.iter_rows(min_row=5, max_row=8) for c in row]
    ok8 = header[1] == f"{first.day:02d}.{first.month:02d}" and \
        any(isinstance(v, str) and re.fullmatch(r"ДЯ \d+( ДН \d+)?", v) for v in cells) and \
        "Переработки" in wb.sheetnames
    sample = next(v for v in cells if isinstance(v, str) and v.startswith("ДЯ"))
    ev8 = f"шапка ФИО|01.MM|…, ячейки вида «{sample}», лист «Переработки»"
check(8, "Выгрузка Excel-табеля сеткой ФИО × дни", ok8, ev8)

# ── 9. ночь = официальный день: день+ночь в одну дату ──
import unittest  # noqa: E402

from tests.test_engine import TestShiftDateAttribution  # noqa: E402

suite = unittest.TestSuite()
suite.addTest(TestShiftDateAttribution("test_night_planned_plus_day_work_same_date"))
suite.addTest(TestShiftDateAttribution("test_day_plus_night_marathon_belongs_to_start_day"))
res = unittest.TextTestRunner(verbosity=0).run(suite)
check(9, "Ночь 22.09 и «день+ночь» → переработка в 22.09", res.wasSuccessful(),
      "22.09 20:00–24:00 + 23.09 00:00–08:00 и марафон 24 ч относятся к дате начала")

# ── 10. блоки: Смена 1 / Смена 2 / Администрация ──
groups = sorted({e["schedule_group"] for e in emps if e["schedule_group"]})
order_ok = groups == ["Администрация", "Смена 1", "Смена 2"] or set(groups) >= {"Смена 1", "Смена 2", "Администрация"}
check(10, "Блоки графика: Смена 1, Смена 2, Администрация", order_ok, f"группы: {', '.join(groups)}")

# ── 11. цвета должностей ──
by_pos = {}
for e in emps:
    by_pos.setdefault(e["position"], set()).add(e["group_color"])
gray = by_pos.get("Батлер", set())
orange = by_pos.get("Старший батлер", set())
yellow = by_pos.get("Менеджер объекта", set()) | by_pos.get("Документооборот", set())
check(11, "Цвета должностей: серый / оранжевый / жёлтый",
      gray == {"#8a94a6"} and orange == {"#e8842c"} and yellow == {"#d9a514"},
      f"батлер {sorted(gray)}, старший {sorted(orange)}, менеджеры/ДО {sorted(yellow)}")

print("\nИТОГО:", sum(1 for *_ , ok, _ in results if ok), "из", len(results), "требований подтверждены")
raise SystemExit(0 if all(ok for *_ , ok, _ in results) else 1)
