"""Скриншоты фичи корпоративных docx-шаблонов: панель настроек + кнопки в редакторе ячейки."""
from __future__ import annotations

import io
from pathlib import Path

import httpx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from playwright.sync_api import sync_playwright

BASE = "http://127.0.0.1:8000"
OUT = Path(__file__).parent / "screens"
OUT.mkdir(exist_ok=True)


def make_demo_template() -> bytes:
    """Макет «корпоративного» бланка: свой шрифт, шапка справа, разбитый плейсхолдер."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = brand.add_run("ООО «Батлер Сервис»")
    r.bold = True
    r.font.size = Pt(14)

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    head.add_run("Генеральному директору\n{director}\nот {position}\n").italic = True
    head.add_run("{full_")          # плейсхолдер специально разбит на два run
    head.add_run("name}\nтабельный № {tab_number}")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("ЗАЯВЛЕНИЕ")
    tr.bold = True

    body = doc.add_paragraph()
    body.paragraph_format.first_line_indent = Pt(24)
    body.add_run("Прошу предоставить мне ежегодный оплачиваемый отпуск "
                 "продолжительностью {days_word} ({days}) календарных дней "
                 "с {date_from_ru} по {date_to_ru}.")

    sign = doc.add_paragraph()
    sign.add_run("{today} г. ______________ / {short_name} /")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload_demo_template() -> None:
    data = make_demo_template()
    with httpx.Client(base_url=BASE) as c:
        r = c.post("/api/auth/login", json={"username": "admin", "password": "demo1234"})
        r.raise_for_status()
        r = c.post("/api/docs/templates", params={"type": "vacation_paid"},
                   files={"file": ("blank_zayavlenie.docx", data,
                                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        r.raise_for_status()
        print("upload:", r.json())


def main() -> None:
    upload_demo_template()
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        ctx = browser.new_context(viewport={"width": 1560, "height": 950}, locale="ru-RU")
        page = ctx.new_page()
        page.goto(BASE, wait_until="networkidle")
        page.fill("#login-username", "admin")
        page.fill("#login-password", "demo1234")
        page.click("#login-form button[type=submit]")
        page.wait_for_timeout(1500)

        # 1) панель шаблонов в настройках
        page.click('.nav-item[data-view="settings"]')
        page.wait_for_selector("#view-settings .panel", timeout=10000)
        page.wait_for_timeout(700)
        panel = page.locator("#view-settings .panel",
                             has_text="Шаблоны документов компании").first
        panel.scroll_into_view_if_needed()
        page.wait_for_timeout(300)
        panel.screenshot(path=str(OUT / "docs_templates_panel.png"))
        print("saved docs_templates_panel.png")
        # раскрыть справку плейсхолдеров и снять ещё раз
        try:
            panel.locator("details summary").click()
            page.wait_for_timeout(300)
            panel.screenshot(path=str(OUT / "docs_templates_placeholders.png"))
            print("saved docs_templates_placeholders.png")
        except Exception as e:
            print("details skip:", e)

        # 2) редактор ячейки с кнопками DOCX/PDF
        page.click('.nav-item[data-view="schedule"]')
        page.wait_for_selector("#grid-host .cellbtn:not([disabled])", timeout=10000)
        page.click("#grid-host .cellbtn:not([disabled])")
        page.wait_for_selector(".modal .shift-opt", timeout=8000)
        vac = page.locator(".modal .shift-opt", has_text="тпуск").first
        vac.click()
        page.wait_for_timeout(400)
        page.locator(".modal").screenshot(path=str(OUT / "docs_cell_buttons.png"))
        print("saved docs_cell_buttons.png")

        browser.close()


if __name__ == "__main__":
    main()
