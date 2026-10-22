"""Картинки во встроенном PDF-рендере (22.10.2026): путь LibreOffice удалён, и встроенный
рендер обязан сам вставлять фото — раньше в Docker (без LibreOffice) PDF ночного отчёта
выходил без фотографий. Фото ночного отчёта лежат в ячейках таблицы, бланки — в абзацах.
"""
from __future__ import annotations

import io

import pytest

pytest.importorskip("reportlab")


def _png(w: int, h: int) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (w, h), (200, 40, 40)).save(buf, "PNG")
    return buf.getvalue()


def test_pictures_from_paragraph_and_table_cell_reach_pdf():
    from docx import Document
    from docx.shared import Cm
    from pypdf import PdfReader

    from app.core.render import docx_to_pdf, pdf_mode

    if pdf_mode() == "none":
        pytest.skip("нет шрифта с кириллицей для встроенного рендера")
    doc = Document()
    doc.add_paragraph("Логотип бланка")
    doc.add_paragraph().add_run().add_picture(io.BytesIO(_png(400, 200)), width=Cm(30))  # шире страницы
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Пункт обхода"
    table.cell(0, 1).add_paragraph().add_run().add_picture(io.BytesIO(_png(300, 300)), width=Cm(6))
    buf = io.BytesIO()
    doc.save(buf)

    pdf, how = docx_to_pdf(buf.getvalue(), "probe.pdf")
    assert pdf and how == "native"
    page = PdfReader(io.BytesIO(pdf)).pages
    images = [img for p in page for img in p.images]
    assert len(images) == 2, "картинка из абзаца или из ячейки таблицы не попала в PDF"
    assert "Пунктобхода" in "".join("".join(p.extract_text() or "" for p in page).split())


def test_broken_picture_does_not_break_document():
    from docx import Document
    from docx.shared import Cm

    from app.core.render import docx_to_pdf, pdf_mode

    if pdf_mode() == "none":
        pytest.skip("нет шрифта с кириллицей для встроенного рендера")
    doc = Document()
    doc.add_paragraph("Текст остаётся")
    doc.add_paragraph().add_run().add_picture(io.BytesIO(_png(50, 50)), width=Cm(2))
    buf = io.BytesIO()
    doc.save(buf)
    raw = buf.getvalue()
    # портим картинку внутри архива docx: рендер должен пропустить её, а не упасть
    import zipfile

    src, out = zipfile.ZipFile(io.BytesIO(raw)), io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item)
            dst.writestr(item, b"not an image" if item.filename.startswith("word/media/") else data)
    pdf, _ = docx_to_pdf(out.getvalue(), "broken.pdf")
    assert pdf and pdf.startswith(b"%PDF")


def test_unbuildable_document_gives_none_not_exception(monkeypatch):
    """Если reportlab не собрал документ ни основной, ни запасной сборкой, рендер возвращает None
    (вызывающий отдаст 501 с подсказкой про DOCX), а не роняет запрос в 500. Сбой вызываем
    сами: высокую строку таблицы reportlab 5 уже умеет делить по страницам."""
    from docx import Document
    from reportlab.platypus import BaseDocTemplate

    from app.core.render import docx_to_pdf, pdf_mode

    if pdf_mode() == "none":
        pytest.skip("нет шрифта с кириллицей для встроенного рендера")

    def broken_build(self, *args, **kwargs):
        raise RuntimeError("Flowable too large")

    monkeypatch.setattr(BaseDocTemplate, "build", broken_build)   # SimpleDocTemplate — наследник
    doc = Document()
    doc.add_paragraph("Текст")
    buf = io.BytesIO()
    doc.save(buf)
    pdf, how = docx_to_pdf(buf.getvalue(), "broken.pdf")
    assert pdf is None and how
