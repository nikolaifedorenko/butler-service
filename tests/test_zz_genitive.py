"""Родительный падеж ФИО: правила склонения, поле карточки, плейсхолдер {full_name_genitive}.

Файл назван на «zz», чтобы импортироваться после test_api (общая тестовая БД):
модуль test_api удаляет и пересоздаёт БД при импорте, поэтому свои данные создаём
только на этапе выполнения тестов, а при отдельном запуске файла — сами чистим БД.
"""
from __future__ import annotations

import io
import os
import sys
import tempfile
import unittest

tmp_db = os.path.join(tempfile.gettempdir(), "tt_test.db")
if "app.main" not in sys.modules and os.path.exists(tmp_db):
    os.remove(tmp_db)
os.environ.setdefault("DATABASE_URL", f"sqlite:///{tmp_db}")
os.environ.setdefault("SECRET_KEY", "test-secret-key")

from fastapi.testclient import TestClient  # noqa: E402

from app.deps import local_date  # noqa: E402
from app.main import app, init_db  # noqa: E402
from app.names import suggest_genitive  # noqa: E402

init_db()
client = TestClient(app)


class TestSuggestGenitive(unittest.TestCase):
    """Чистые правила склонения — без API."""

    def test_male_typical(self):
        self.assertEqual(suggest_genitive("Иванов Иван Иванович"), "Иванова Ивана Ивановича")
        self.assertEqual(suggest_genitive("Федоренко Николай Сергеевич"),
                         "Федоренко Николая Сергеевича")          # -ко не склоняется
        self.assertEqual(suggest_genitive("Кузнецов Никита Сергеевич"),
                         "Кузнецова Никиты Сергеевича")            # мужское имя на -а
        self.assertEqual(suggest_genitive("Гоголь Николай Васильевич"), "Гоголя Николая Васильевича")
        self.assertEqual(suggest_genitive("Седых Иван Петрович"), "Седых Ивана Петровича")  # -их
        self.assertEqual(suggest_genitive("Толстой Лев"), "Толстого Льва")                  # -й

    def test_female_typical(self):
        self.assertEqual(suggest_genitive("Смирнова Анна Петровна"), "Смирновой Анны Петровны")
        self.assertEqual(suggest_genitive("Толстая Мария Ильинична"), "Толстой Марии Ильиничны")
        self.assertEqual(suggest_genitive("Иванова Любовь Петровна"), "Ивановой Любови Петровны")
        self.assertEqual(suggest_genitive("Гоголь Ольга Николаевна"), "Гоголь Ольги Николаевны")

    def test_edge(self):
        self.assertEqual(suggest_genitive(""), "")
        self.assertEqual(suggest_genitive("Иванов"), "Иванова")       # одна фамилия (муж.)
        self.assertEqual(suggest_genitive("Иванова"), "Ивановой")     # одна фамилия (жен.)
        self.assertEqual(suggest_genitive("Ольга"), "Ольги")          # -га → -ги


class TestGenitiveApi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        r = client.post("/api/auth/login", json={"username": "gromova", "password": "demo1234"})
        assert r.status_code == 200, r.text
        cls.admin = TestClient(app)
        r = cls.admin.post("/api/auth/login", json={"username": "admin", "password": "demo1234"})
        assert r.status_code == 200, r.text
        cls.emp_client = TestClient(app)
        r = cls.emp_client.post("/api/auth/login", json={"username": "ivanov", "password": "demo1234"})
        assert r.status_code == 200, r.text

    def _create(self, full_name, genitive=""):
        r = self.admin.post("/api/employees", json={
            "full_name": full_name, "full_name_genitive": genitive,
            "position": "Батлер", "schedule_group": "Смена 1"})
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def test_01_field_and_hint(self):
        emp = self._create("Федоренко Николай Сергеевич", "Федоренко Николая Сергеевича")
        self.assertEqual(emp["full_name_genitive"], "Федоренко Николая Сергеевича")
        self.assertEqual(emp["full_name_genitive_hint"], "Федоренко Николая Сергеевича")

        emp2 = self._create("Тестов Пётр Ильич")     # поле пустое → hint = автоподстановка
        self.assertEqual(emp2["full_name_genitive"], "")
        self.assertEqual(emp2["full_name_genitive_hint"], "Тестова Петра Ильича")

        # обновление поля
        body = {k: emp2.get(k, "") for k in ("full_name", "short_name", "position", "phone",
                                             "tab_number", "telegram", "email")}
        body.update({"full_name": "Тестов Пётр Ильич", "full_name_genitive": "Тестова Петра",
                     "hired_at": emp2["hired_at"], "schedule_group": emp2["schedule_group"],
                     "group_color": emp2["group_color"], "balance_hours": emp2["balance_hours"],
                     "role": emp2["role"] or "employee", "username": emp2["username"]})
        r = self.admin.put(f"/api/employees/{emp2['id']}", json=body)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["full_name_genitive"], "Тестова Петра")

        # список и карточка содержат поле
        emps = client.get("/api/employees").json()
        me = next(e for e in emps if e["id"] == emp["id"])
        self.assertEqual(me["full_name_genitive"], "Федоренко Николая Сергеевича")

    def test_02_suggest_endpoint(self):
        r = client.post("/api/employees/suggest-genitive",
                        json={"full_name": "Смирнова Анна Петровна"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["genitive"], "Смирновой Анны Петровны")
        # сотруднику недоступно (только менеджеры)
        r = self.emp_client.post("/api/employees/suggest-genitive", json={"full_name": "Иванов"})
        self.assertEqual(r.status_code, 403)

    def test_03_docx_render_genitive(self):
        """{full_name_genitive} в корпоративном шаблоне: ручной вариант приоритетен,
        при пустом поле — автоподстановка; разбитый на runs плейсхолдер склеивается."""
        from docx import Document

        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Заявление от {full_")
        p.add_run("name_genitive}, таб. № {tab_number}.")
        buf = io.BytesIO()
        doc.save(buf)
        tpl = buf.getvalue()

        r = self.admin.post("/api/docs/templates", params={"type": "vacation_paid"},
                            files={"file": ("gen.docx", tpl,
                                            "application/vnd.openxmlformats-officedocument."
                                            "wordprocessingml.document")})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("full_name_genitive", r.json()["placeholders"])
        info = self.admin.get("/api/docs/templates").json()
        self.assertIn("full_name_genitive", info["placeholders"])   # справка по плейсхолдерам

        emp_manual = self._create("Ручной Иван Петрович", "Ручного Ивана Петровича (сверено)")
        emp_auto = self._create("Авто Карл Марксович")
        d1, d2 = local_date().replace(day=1), local_date().replace(day=5)

        for emp, expected in ((emp_manual, "Ручного Ивана Петровича (сверено)"),
                              (emp_auto, "Авто Карла Марксовича")):
            r = self.admin.post("/api/docs/render", json={
                "type": "vacation_paid", "employee_id": emp["id"],
                "date_from": d1.isoformat(), "date_to": d2.isoformat(), "format": "docx"})
            self.assertEqual(r.status_code, 200, r.text)
            out = Document(io.BytesIO(r.content))
            text = "\n".join(p.text for p in out.paragraphs)
            self.assertIn(expected, text)
            self.assertNotIn("{full_name_genitive}", text)

        # текст-шаблон из настроек знает плейсхолдер
        settings = client.get("/api/settings").json()
        self.assertIn("{full_name_genitive}", settings["doc"]["placeholders"])
        self.assertIn("{full_name_genitive}", settings["doc"]["vacation"])

        # сетка графика отдаёт эффективное значение для клиентской печати
        today = local_date()
        grid = client.get("/api/schedule", params={"year": today.year, "month": today.month}).json()
        row = next(x for x in grid["rows"] if x["employee"]["id"] == emp_manual["id"])
        self.assertEqual(row["employee"]["full_name_genitive"], "Ручного Ивана Петровича (сверено)")
        row = next(x for x in grid["rows"] if x["employee"]["id"] == emp_auto["id"])
        self.assertEqual(row["employee"]["full_name_genitive"], "Авто Карла Марксовича")

        self.admin.delete("/api/docs/templates", params={"type": "vacation_paid"})

    def test_04_sample_docx(self):
        r = self.admin.get("/api/docs/sample", params={"type": "vacation_paid"})
        self.assertEqual(r.status_code, 200, r.text)
        from docx import Document
        out = Document(io.BytesIO(r.content))
        text = "\n".join(p.text for p in out.paragraphs)
        for ph in ("{director}", "{company}", "{full_name_genitive}", "{days_word}",
                   "{date_from_ru}", "{short_name}", "{today}"):
            self.assertIn(ph, text)
        r = self.admin.get("/api/docs/sample", params={"type": "vacation_unpaid"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("без сохранения", "\n".join(
            p.text for p in Document(io.BytesIO(r.content)).paragraphs))
        r = self.admin.get("/api/docs/sample", params={"type": "nope"})
        self.assertEqual(r.status_code, 422)
        r = self.emp_client.get("/api/docs/sample", params={"type": "vacation_paid"})
        self.assertEqual(r.status_code, 403)


if __name__ == "__main__":
    unittest.main(verbosity=2)
