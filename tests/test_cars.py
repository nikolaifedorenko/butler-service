"""Электрокары: операции парка, права, история, ночной обход, отдача фото.

Модуль добавлен в проект последним и не имел ни одного теста; здесь закрыты
основные сценарии ТЗ и две найденные регрессии:
  * NameError в cars.return_car() — условие `(by_employee_id or name)` обращалось
    к несуществующей переменной, из-за чего ЛЮБОЙ возврат кара без явного
    employee_id (обычный случай: сотрудник возвращает свой кар) падал в 500;
  * отсутствующий маршрут /api/photos/{id}/file — photos.photo_url() всегда
    отдавал эту ссылку, а маршрут существовал только как /api/night/photos/…,
    поэтому каждое загруженное фото показывалось как 404.

Отметка кара в ночном обходе (19.10.2026) приведена к окну возврата кара: та же
форма (место из справочника, заряд, чекбоксы состояния, комментарий), чекбокс
«Кар найден на месте» и фото-доказательство — поэтому эндпоинт принимает
multipart, а не JSON (POST /api/night/reports/{id}/cars/{car_id}/check).

Каждый тест возвращает парк в исходное состояние: демо-данные общие на прогон.
"""
from __future__ import annotations

import base64
import datetime as dt
from contextlib import contextmanager

RETURN_FORM = {"by_name": "", "location": "Парковка у главного входа", "charge": "half",
               "canopy": "true", "condition": "ok", "trash": "false", "clean": "true",
               "on_charge": "false", "comment": ""}

# отметка кара в ночном обходе: те же поля, что в окне возврата (строки — как шлёт FormData)
NIGHT_CHECK_FORM = {"found": "true", "location": "", "canopy": "true", "charge": "half",
                    "on_charge": "false", "condition": "ok", "trash": "false",
                    "clean": "true", "comment": ""}

PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFAAH/q842iQAAAABJRU5ErkJggg==")


@contextmanager
def _nochnaya_smena(db, freeze, shift_d: dt.date):
    """Открытый отчёт ночной смены с замороженным временем + уборка за тестом.

    Смену берём фиксированную (прошлую), чтобы тест не зависел от часа запуска
    прогона: иначе отчёт был бы закрыт и тест пропускался бы почти всегда.
    Демо-данные общие на прогон — отметки каров, их фото (и файлы на диске) и сам
    отчёт удаляем; журнал истории кара не редактируется по правилам ТЗ.
    """
    from sqlalchemy import select

    from app.models import CarNightCheck, NightReport, Photo
    from app.night import ensure_report
    from app.photos import unlink_photo

    with freeze(f"{shift_d.isoformat()} 23:00:00"):      # 02:00 МСК следующих суток — смена идёт
        rep = ensure_report(db, shift_d)
        assert rep is not None, "отчёт для идущей смены не создан"
        try:
            yield rep
        finally:
            checks = list(db.scalars(select(CarNightCheck)
                                     .where(CarNightCheck.report_id == rep.id)))
            ids = [c.id for c in checks]
            if ids:
                for ph in db.scalars(select(Photo).where(Photo.kind == "car_check",
                                                         Photo.ref_id.in_(ids))):
                    unlink_photo(ph)
                    db.delete(ph)
            for ch in checks:
                db.delete(ch)
            fresh = db.get(NightReport, rep.id)
            if fresh is not None:
                db.delete(fresh)          # каскад убирает области и пункты
            db.commit()


def _obhod(client, report_id: int) -> dict:
    return client.get(f"/api/night/reports/{report_id}/cars").json()


def _park(client) -> list[dict]:
    return client.get("/api/cars").json()["cars"]


def _car_by_number(client, number: str) -> dict:
    return next(c for c in _park(client) if c["number"] == number)


def _free_car(client) -> dict:
    cars = [c for c in _park(client) if c["status"] in ("free", "charging")]
    assert cars, "в демо-парке нет свободного кара"
    return cars[0]


def _return_car(client, car_id: int, **overrides) -> dict:
    form = {**RETURN_FORM, **overrides}
    r = client.post(f"/api/cars/{car_id}/return", data=form)
    return r


# ───────────────────────────── видимость и справочники ─────────────────────────────
def test_park_viden_vsem_rolyam(employee_client, manager_client):
    for c in (employee_client, manager_client):
        r = c.get("/api/cars")
        assert r.status_code == 200, r.text
        data = r.json()
        assert len(data["cars"]) >= 5, "демо-парк не загружен"
        assert {"id", "number", "status", "status_title", "location", "charge_title"} <= set(data["cars"][0])


def test_meta_soderzhit_slovari_i_prava(employee_client, manager_client,
                                        manager_role_client):
    """Права раздела 18.10.2026: use — сотрудник, issue — супервайзер, manage — менеджер."""
    emp = employee_client.get("/api/cars/meta").json()
    sup = manager_client.get("/api/cars/meta").json()          # gromova = supervisor
    mgr = manager_role_client.get("/api/cars/meta").json()     # smirnov = manager
    assert {s["code"] for s in emp["statuses"]} >= {"free", "busy", "charging", "maintenance"}
    assert (emp["can_use"], emp["can_issue"], emp["can_manage"]) == (True, False, False)
    assert (sup["can_use"], sup["can_issue"], sup["can_manage"]) == (True, True, False)
    assert (mgr["can_use"], mgr["can_issue"], mgr["can_manage"]) == (True, True, True)


def test_otkluchennye_kary_skryty_ot_sotrudnika(employee_client, admin_client):
    car = _car_by_number(admin_client, "5")
    r = admin_client.put(f"/api/cars/{car['id']}", json={"active": False})
    assert r.status_code == 200, r.text
    try:
        assert all(c["number"] != "5" for c in _park(employee_client)), \
            "отключённый кар виден сотруднику"
        assert any(c["number"] == "5" for c in _park(admin_client)), \
            "менеджмент должен видеть отключённые кары"
    finally:
        admin_client.put(f"/api/cars/{car['id']}", json={"active": True})


def test_mesta_stoyanki_crud(manager_role_client, admin_client):
    before = {loc["name"] for loc in manager_role_client.get("/api/cars/locations").json()["locations"]}
    r = manager_role_client.post("/api/cars/locations", json={"name": "Тестовый навес №42"})
    assert r.status_code == 200, r.text
    loc_id = r.json()["id"]
    try:
        names = {loc["name"] for loc in manager_role_client.get("/api/cars/locations").json()["locations"]}
        assert "Тестовый навес №42" in names and "Тестовый навес №42" not in before
        r = admin_client.put(f"/api/cars/locations/{loc_id}", json={"active": False})
        assert r.status_code == 200, r.text
    finally:
        admin_client.put(f"/api/cars/locations/{loc_id}", json={"active": False})


# ───────────────────────────── возврат кара (регрессия NameError) ─────────────────────────────
def test_vozvrat_svoego_kara_bez_by_poley(employee_client, manager_client):
    """Сотрудник возвращает СВОЙ кар, не указывая ни employee_id, ни ФИО.

    Именно этот путь падал с NameError (cars.py: `if (by_employee_id or name)`).
    """
    me = employee_client.get("/api/cars").json()["me"] or {}
    held_id = me.get("held")
    if not held_id:                       # если кар уже вернули предыдущие тесты — берём свободный
        car = _free_car(employee_client)
        r = employee_client.post(f"/api/cars/{car['id']}/take", json={"ack_assigned": False})
        assert r.status_code == 200, r.text
        held_id = car["id"]

    r = _return_car(employee_client, held_id, comment="оставил у главного входа")
    assert r.status_code == 200, f"возврат кара упал: {r.text}"
    car = r.json()["car"]
    assert car["status"] in ("free", "charging", "maintenance")
    assert car["holder_name"] == "", "держатель не очищен после возврата"
    assert car["holder_user_id"] is None


def test_vozvrat_s_zamechaniyami_perevodit_na_obsluzhivanie(manager_client, admin_client):
    car = _free_car(manager_client)
    r = manager_client.post(f"/api/cars/{car['id']}/take", json={"ack_assigned": False})
    assert r.status_code == 200, r.text
    try:
        r = _return_car(manager_client, car["id"], condition="bad", comment="спущено колесо")
        assert r.status_code == 200, r.text
        assert r.json()["car"]["status"] == "maintenance", \
            "замечания при возврате обязаны перевести кар в «на обслуживании»"
    finally:
        admin_client.post(f"/api/cars/{car['id']}/status", json={"status": "free", "note": "тест"})


def test_vozvrat_polnogo_kara_na_zaryadku_razreshen(manager_client, admin_client):
    """18.10.2026: оставить на зарядке можно кар с ЛЮБЫМ зарядом, включая полный."""
    car = _free_car(manager_client)
    r = manager_client.post(f"/api/cars/{car['id']}/take", json={"ack_assigned": False})
    assert r.status_code == 200, r.text
    try:
        r = _return_car(manager_client, car["id"], on_charge="true", charge="full")
        assert r.status_code == 200, f"заряженный кар можно оставлять на зарядке: {r.text}"
        body = r.json()["car"]
        assert body["status"] == "charging" and body["on_charge"] is True
    finally:
        _return_car(manager_client, car["id"], on_charge="false")


def test_veshchi_v_saloned_perevodyat_na_obsluzhivanie(manager_client, admin_client):
    """Чекбокс «В салоне вещи» (18.10.2026) — замечание: кар уходит на обслуживание."""
    car = _free_car(manager_client)
    r = manager_client.post(f"/api/cars/{car['id']}/take", json={"ack_assigned": False})
    assert r.status_code == 200, r.text
    r = _return_car(manager_client, car["id"], trash="true", condition="ok", clean="true")
    assert r.status_code == 200, r.text
    body = r.json()["car"]
    assert body["status"] == "maintenance", body
    assert body["trash"] is True
    r = admin_client.post(f"/api/cars/{car['id']}/status",
                          json={"status": "free", "note": "салон освобождён (тест)"})
    assert r.status_code == 200, r.text


def test_vozvrat_na_zaryadke_s_razryazhennym_zaryadom(manager_client, admin_client):
    car = _free_car(manager_client)
    r = manager_client.post(f"/api/cars/{car['id']}/take", json={"ack_assigned": False})
    assert r.status_code == 200, r.text
    try:
        r = _return_car(manager_client, car["id"], on_charge="true", charge="empty")
        assert r.status_code == 200, r.text
        assert r.json()["car"]["status"] == "charging"
    finally:
        admin_client.post(f"/api/cars/{car['id']}/status", json={"status": "free", "note": "тест"})


# ───────────────────────────── правила «один кар» и права ─────────────────────────────
def test_vtoroy_kar_vzyat_nelzya(employee_client):
    """«Один сотрудник — один кар»: пока есть держимый, второй взять нельзя."""
    me = employee_client.get("/api/cars").json()["me"] or {}
    held_id = me.get("held")
    if not held_id:                       # берём сами, чтобы тест не зависел от порядка
        car = _free_car(employee_client)
        r = employee_client.post(f"/api/cars/{car['id']}/take", json={"ack_assigned": False})
        assert r.status_code == 200, r.text
        held_id = car["id"]
    try:
        other = next(c for c in _park(employee_client)
                     if c["id"] != held_id and c["status"] in ("free", "charging"))
        r = employee_client.post(f"/api/cars/{other['id']}/take", json={"ack_assigned": False})
        assert r.status_code == 400, r.text
        detail = r.json()["detail"].lower()
        assert "один сотрудник" in detail or "верните" in detail
    finally:
        _return_car(employee_client, held_id)


def test_vzyat_kar_s_zaryadki_mozhno(manager_client, admin_client):
    car = next((c for c in _park(manager_client) if c["status"] == "charging"), None)
    created = False
    if car is None:                       # создаём условие сами: свободный кар → «на зарядке»
        car = _free_car(manager_client)
        r = admin_client.post(f"/api/cars/{car['id']}/status", json={"status": "charging", "note": "тест"})
        assert r.status_code == 200, r.text
        created = True
    r = manager_client.post(f"/api/cars/{car['id']}/take", json={"ack_assigned": False})
    assert r.status_code == 200, f"кар на зарядке должен быть доступен: {r.text}"
    try:
        assert r.json()["car"]["status"] == "busy"
    finally:
        _return_car(manager_client, car["id"])
        admin_client.post(f"/api/cars/{car['id']}/status",
                          json={"status": "free" if created else "charging", "note": "тест"})


def test_peredacha_drugomu_lichno(manager_client, admin_client):
    car = _free_car(manager_client)
    r = manager_client.post(f"/api/cars/{car['id']}/take", json={"ack_assigned": False})
    assert r.status_code == 200, r.text
    try:
        r = manager_client.post(f"/api/cars/{car['id']}/handover",
                                json={"employee_id": None, "name": "Внешний Тест Тестович",
                                      "with_key": True, "comment": "на время обхода"})
        assert r.status_code == 200, r.text
        after = r.json()["car"]
        assert after["holder_name"] == "Внешний Тест Тестович"
        assert after["status"] == "busy"
    finally:
        admin_client.post(f"/api/cars/{car['id']}/status", json={"status": "free", "note": "тест"})


def test_vydacha_i_zakreplenie_tolko_menedzhmentu(employee_client, manager_client):
    car = _free_car(employee_client)
    assert employee_client.post(f"/api/cars/{car['id']}/give",
                                json={"name": "Кто-то"}).status_code == 403
    assert employee_client.post(f"/api/cars/{car['id']}/assign",
                                json={"employee_id": None}).status_code == 403
    assert employee_client.post(f"/api/cars/{car['id']}/status",
                                json={"status": "free"}).status_code == 403
    r = manager_client.post(f"/api/cars/{car['id']}/assign", json={"employee_id": None})
    assert r.status_code == 200, r.text


def test_istoriya_ne_pusta_i_soderzhit_deystviya(manager_client, admin_client):
    car = _free_car(manager_client)
    r = manager_client.post(f"/api/cars/{car['id']}/take", json={"ack_assigned": False})
    assert r.status_code == 200, r.text
    try:
        _return_car(manager_client, car["id"])
        hist = manager_client.get(f"/api/cars/{car['id']}/history").json()["history"]
        actions = [h["action"] for h in hist]
        assert "take" in actions and "return" in actions, actions[:6]
        assert all({"ts", "action_title", "actor_name"} <= set(h) for h in hist)
    finally:
        admin_client.post(f"/api/cars/{car['id']}/status", json={"status": "free", "note": "тест"})


# ───────────────────────────── фото: маршрут отдачи существует ─────────────────────────────
def test_marshrut_otdachi_foto_sushchestvuet(employee_client):
    """/api/photos/{id}/file обязан существовать: на него ссылается photos.photo_url().

    До появления app/api/photos_api.py маршрут был только /api/night/photos/{id}/file,
    поэтому любая ссылка на фото отдавала 404 «Not Found» (маршрут не найден).
    Здесь различаем «маршрута нет» (detail == 'Not Found') и «фото не найдено».
    """
    r = employee_client.get("/api/photos/999999/file")
    assert r.status_code == 404, r.text
    assert r.json()["detail"] == "Фото не найдено", \
        f"похоже, маршрута /api/photos/{{id}}/file нет: {r.json()}"
    # алиас в ночном модуле продолжает работать
    r2 = employee_client.get("/api/night/photos/999999/file")
    assert r2.status_code == 404 and r2.json()["detail"] == "Фото не найдено"


def test_foto_bez_vhoda_ne_otdayotsya(client):
    assert client.get("/api/photos/1/file").status_code == 401


# ───────────────────────────── ночной обход электрокаров ─────────────────────────────
def test_nochnoy_obhod_karov_i_zavershenie_shaga(manager_client, db, freeze):
    """Шаг «Проверка электрокаров» внутри ночного отчёта: отметка по кару и завершение шага."""
    with _nochnaya_smena(db, freeze, dt.date(2026, 5, 12)) as rep:
        state = _obhod(manager_client, rep.id)
        assert state["total"] >= 1, "в обход не попал ни один активный кар"
        assert state["cars_step_done"] is False
        entry = next((e for e in state["list"] if not e["checked"]), state["list"][0])
        car = entry["car"]

        r = manager_client.post(f"/api/night/reports/{rep.id}/cars/{car['id']}/check",
                                data={**NIGHT_CHECK_FORM,
                                      "location": car["location"] or "Парковка у главного входа",
                                      "comment": "проверка тестом"})
        assert r.status_code == 200, r.text

        after = _obhod(manager_client, rep.id)
        checked = next(e for e in after["list"] if e["car"]["id"] == car["id"])
        assert checked["checked"] is True
        assert checked["check"]["comment"] == "проверка тестом"
        assert after["checked"] == state["checked"] + 1

        r = manager_client.post(f"/api/night/reports/{rep.id}/cars/finish", json={"ack": True})
        assert r.status_code == 200, r.text
        assert r.json()["cars_step_done"] is True
        # карточка кара обновлена отметкой обхода (общая база, а не снимок отчёта)
        fresh = next(c for c in _park(manager_client) if c["id"] == car["id"])
        assert fresh["charge"] == "half"
        assert fresh["last_checked_at"] is not None


def test_nochnaya_otmetka_kara_s_foto_i_mestom_iz_spravochnika(manager_client, db, freeze):
    """Отметка ночного обхода = окно возврата кара: место из справочника, чекбоксы
    состояния и фото-доказательство (multipart) — снимок виден в обходе и в DOCX."""
    shift_d = dt.date(2026, 5, 13)
    locations = manager_client.get("/api/cars/locations").json()["locations"]
    place = next(loc["name"] for loc in locations if loc["active"])

    with _nochnaya_smena(db, freeze, shift_d) as rep:
        car = _obhod(manager_client, rep.id)["list"][0]["car"]
        r = manager_client.post(
            f"/api/night/reports/{rep.id}/cars/{car['id']}/check",
            data={**NIGHT_CHECK_FORM, "found": "true", "location": place, "charge": "full",
                  "on_charge": "true", "canopy": "false", "trash": "true", "clean": "false",
                  "condition": "ok", "comment": "кар на месте, снят на камеру"},
            files={"file": ("kar.png", PNG_1PX, "image/png")})
        assert r.status_code == 200, r.text

        check = next(e for e in _obhod(manager_client, rep.id)["list"]
                     if e["car"]["id"] == car["id"])["check"]
        assert check["found"] is True and check["location"] == place
        assert check["charge"] == "full" and check["on_charge"] is True
        assert check["canopy"] is False and check["trash"] is True and check["clean"] is False
        assert len(check["photos"]) == 1, "фото отметки не вернулось в обход"
        photo = check["photos"][0]
        got = manager_client.get(photo["url"])
        assert got.status_code == 200 and got.content == PNG_1PX

        # карточка кара обновлена отметкой (общая база)
        fresh = next(c for c in _park(manager_client) if c["id"] == car["id"])
        assert fresh["location"] == place and fresh["charge"] == "full"
        assert fresh["on_charge"] is True and fresh["canopy"] is False

        # «в салоне вещи» и «не чистый» — факты карточки (в отличие от возврата кара,
        # где они уводят кар на обслуживание): статус ночная отметка не меняет
        assert fresh["trash"] is True and fresh["clean"] is False
        assert fresh["status"] in ("free", "charging", "busy")

        # повторное сохранение (last write wins): замечание по техсостоянию уводит кар
        # на обслуживание, а прежний снимок остаётся привязан к отметке
        r = manager_client.post(f"/api/night/reports/{rep.id}/cars/{car['id']}/check",
                                data={**NIGHT_CHECK_FORM, "location": place, "condition": "bad",
                                      "comment": "спущено колесо"})
        assert r.status_code == 200, r.text
        assert next(c for c in _park(manager_client) if c["id"] == car["id"])["status"] == "maintenance"
        state2 = _obhod(manager_client, rep.id)
        check2 = next(e for e in state2["list"] if e["car"]["id"] == car["id"])["check"]
        assert check2["condition"] == "bad" and check2["comment"] == "спущено колесо"
        assert len(check2["photos"]) == 1, "фото отметки потерялось при повторном сохранении"

        # выгрузка DOCX: блок обхода каров и снимок попадают в документ
        # (раньше render_night_report_docx() звал report_dict(rep) без db — блоки
        # «Проверка электрокаров» и «Перехваты областей» были пустыми)
        import io

        from docx import Document

        docx = manager_client.get(f"/api/night/reports/{rep.id}/export.docx")
        assert docx.status_code == 200, docx.text
        doc = Document(io.BytesIO(docx.content))
        text = "\n".join(par.text for par in doc.paragraphs) + "\n" + "\n".join(
            c.text for t in doc.tables for row in t.rows for c in row.cells)
        assert "Проверка электрокаров" in text
        assert any(t.rows[0].cells[0].text == "Время" for t in doc.tables), \
            "таблицы обхода каров нет в документе"
        assert car["number"] in text, "обход каров не попал в DOCX"
        assert "спущено колесо" in text and "ЕСТЬ ЗАМЕЧАНИЯ" in text
        assert len(doc.inline_shapes) >= 1, "фото отметки не вставлено в документ"


def test_otmetka_ne_nayden_ne_trogaet_kartochku_kara(manager_client, db, freeze):
    """«Кар найден на месте» снят: отметка уходит в замечания, карточка не меняется."""
    with _nochnaya_smena(db, freeze, dt.date(2026, 5, 14)) as rep:
        car = _obhod(manager_client, rep.id)["list"][0]["car"]
        before = next(c for c in _park(manager_client) if c["id"] == car["id"])

        r = manager_client.post(f"/api/night/reports/{rep.id}/cars/{car['id']}/check",
                                data={**NIGHT_CHECK_FORM, "found": "false", "location": "",
                                      "charge": "empty", "condition": "bad", "clean": "false",
                                      "comment": "кара нет на месте"})
        assert r.status_code == 200, r.text

        state = _obhod(manager_client, rep.id)
        check = next(e for e in state["list"] if e["car"]["id"] == car["id"])["check"]
        assert check["found"] is False and check["comment"] == "кара нет на месте"
        assert state["problems"] >= 1, "«не найден» не попал в замечания"

        after = next(c for c in _park(manager_client) if c["id"] == car["id"])
        for field in ("location", "charge", "condition", "clean", "status", "on_charge"):
            assert after[field] == before[field], f"карточка кара изменена при «не найден»: {field}"


def test_foto_otmetki_ne_kartinka_otklonyaetsya(manager_client, db, freeze):
    """В фото отметки принимаются только изображения (как в окне возврата кара)."""
    with _nochnaya_smena(db, freeze, dt.date(2026, 5, 15)) as rep:
        car = _obhod(manager_client, rep.id)["list"][0]["car"]
        r = manager_client.post(
            f"/api/night/reports/{rep.id}/cars/{car['id']}/check",
            data=NIGHT_CHECK_FORM,
            files={"file": ("zametka.txt", b"ne kartinka", "text/plain")})
        assert r.status_code == 400, r.text
        assert "изображения" in r.json()["detail"]
        # отметка не сохранена: ошибка до коммита
        entry = next(e for e in _obhod(manager_client, rep.id)["list"] if e["car"]["id"] == car["id"])
        assert entry["checked"] is False


# ───────────────────── права и удаление мест (18.10.2026) ─────────────────────
def test_prava_operacii_razgranicheny(employee_client, manager_client, manager_role_client):
    """Сотрудник не выдаёт и не правит парк; супервайзер не создаёт кары и места."""
    park = employee_client.get("/api/cars").json()["cars"]
    car = next(c for c in park if c["status"] == "free")
    r = employee_client.post(f"/api/cars/{car['id']}/give",
                             json={"employee_id": 1, "with_key": True})
    assert r.status_code == 403, r.text
    r = employee_client.post("/api/cars/locations", json={"name": "Нелегальный навес"})
    assert r.status_code == 403, r.text
    r = manager_client.post("/api/cars", json={"number": "SUP-1"})
    assert r.status_code == 403, "супервайзер без cars.manage не создаёт кары"
    r = manager_client.post("/api/cars/locations", json={"name": "Навес супервайзера"})
    assert r.status_code == 403, "супервайзер без cars.manage не создаёт места"
    r = manager_role_client.post("/api/cars", json={"number": "MGR-1"})
    assert r.status_code == 200, r.text


def test_udalenie_mesta_stoyanki(manager_role_client, employee_client):
    """Место удаляется, пока не указано у активных каров; иначе 409 с подсказкой."""
    r = manager_role_client.post("/api/cars/locations", json={"name": "Временный навес"})
    assert r.status_code == 200, r.text
    loc_id = r.json()["id"]
    r = employee_client.delete(f"/api/cars/locations/{loc_id}")
    assert r.status_code == 403, "удаление мест — право cars.manage"
    r = manager_role_client.delete(f"/api/cars/locations/{loc_id}")
    assert r.status_code == 200, r.text
    names = {x["name"] for x in manager_role_client.get("/api/cars/locations").json()["locations"]}
    assert "Временный навес" not in names
    # место активного кара удалить нельзя
    park = manager_role_client.get("/api/cars").json()["cars"]
    with_loc = next((c for c in park if c.get("location")), None)
    if with_loc:
        locs = manager_role_client.get("/api/cars/locations").json()["locations"]
        loc = next((x for x in locs if x["name"] == with_loc["location"]), None)
        if loc:
            r = manager_role_client.delete(f"/api/cars/locations/{loc['id']}")
            assert r.status_code == 409, r.text
