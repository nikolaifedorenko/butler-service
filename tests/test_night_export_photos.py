"""Фото в ночном отчёте — уменьшенная копия (аудит 22.10.2026): оригиналы с телефона до 12 МБ
делали DOCX и PDF отчёта на сотни мегабайт в памяти. Поворот по EXIF применяется сразу."""
from __future__ import annotations


def test_print_copy_is_small_and_upright(tmp_path):
    from PIL import Image

    from app.modules.night.application.night_export import PRINT_MAX_PX, _print_copy

    src = tmp_path / "phone.jpg"
    im = Image.new("RGB", (4000, 3000), (120, 30, 30))
    exif = Image.Exif()
    exif[0x0112] = 6                         # «повернуть на 90°» — так пишут камеры телефонов
    im.save(src, "JPEG", quality=95, exif=exif.tobytes())

    out = Image.open(_print_copy(src))
    assert max(out.size) <= PRINT_MAX_PX
    assert out.size[1] > out.size[0], "поворот по EXIF не применён: снимок лежит"
    assert len(_print_copy(src).getvalue()) < src.stat().st_size


def test_video_is_not_a_picture(tmp_path):
    from docx import Document

    from app.modules.night.application.night_export import _image_cell

    video = tmp_path / "clip.mp4"
    video.write_bytes(b"\x00\x00\x00\x18ftypmp42" + b"\x00" * 64)
    cell = Document().add_table(rows=1, cols=1).cell(0, 0)
    assert _image_cell(cell, {"path": str(video)}) is False
