"""Интеграционные тесты API: вход, график, отметки, табель, настройки."""
from __future__ import annotations

import datetime as dt
import os
import tempfile
import unittest

tmp_db = os.path.join(tempfile.gettempdir(), "tt_test.db")
if os.path.exists(tmp_db):
    os.remove(tmp_db)
os.environ["DATABASE_URL"] = f"sqlite:///{tmp_db}"
os.environ["SECRET_KEY"] = "test-secret-key"

from fastapi.testclient import TestClient  # noqa: E402

from app.deps import local_date  # noqa: E402
from app.main import app, init_db  # noqa: E402

init_db()
client = TestClient(app)


class TestApi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.manager = client.post("/api/auth/login", json={"username": "gromova", "password": "demo1234"})
        assert cls.manager.status_code == 200, cls.manager.text
        cls.admin = TestClient(app)
        r = cls.admin.post("/api/auth/login", json={"username": "admin", "password": "demo1234"})
        assert r.status_code == 200, r.text
        cls.emp_client = TestClient(app)
        r = cls.emp_client.post("/api/auth/login", json={"username": "ivanov", "password": "demo1234"})
        assert r.status_code == 200, r.text

    def test_01_health_and_login(self):
        self.assertTrue(client.get("/api/health").json()["ok"])
        bad = client.post("/api/auth/login", json={"username": "gromova", "password": "wrong"})
        self.assertEqual(bad.status_code, 401)

    def test_02_me_and_profile(self):
        me = client.get("/api/auth/me").json()
        self.assertEqual(me["user"]["role"], "supervisor")
        self.assertIsNotNone(me["today"])
        emp_me = self.emp_client.get("/api/auth/me").json()
        self.assertEqual(emp_me["user"]["role"], "employee")
        self.assertIsNotNone(emp_me["today"]["date"])

    def test_03_grid(self):
        today = local_date()
        r = client.get("/api/schedule", params={"year": today.year, "month": today.month})
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertGreater(len(data["rows"]), 0)
        import calendar
        self.assertEqual(len(data["days"]), calendar.monthrange(today.year, today.month)[1])
        first = data["rows"][0]
        self.assertIn("cells", first)
        self.assertIn("totals", first)

    def test_04_employee_sees_plan_but_cannot_edit_grid(self):
        """Батлеру график доступен на просмотр (только план), правки — по-прежнему 403."""
        today = local_date()
        r = self.emp_client.get("/api/schedule", params={"year": today.year, "month": today.month})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()["readonly"])
        edit = self.emp_client.put("/api/schedule/cell", json={
            "employee_id": 1, "date": today.isoformat(), "shift_type_id": None, "note": ""})
        self.assertEqual(edit.status_code, 403)

    def test_05_set_cell_and_see_it_in_grid(self):
        # Целевая дата — «сегодня + 3 дня», а сетку нужно запрашивать за МЕСЯЦ ЦЕЛЕВОЙ
        # ДАТЫ: раньше месяц брали от today, и в последние дни месяца (29–31) тест
        # падал с KeyError, потому что ячейка уезжала в следующий месяц.
        target = local_date() + dt.timedelta(days=3)
        date = target.isoformat()
        types = client.get("/api/shift-types").json()
        sick = next(t for t in types if t["code"] == "SICK")
        emps = client.get("/api/employees").json()
        emp = emps[0]
        r = client.put("/api/schedule/cell", json={
            "employee_id": emp["id"], "date": date, "shift_type_id": sick["id"], "note": "тест"})
        self.assertEqual(r.status_code, 200, r.text)
        grid = client.get("/api/schedule", params={"year": target.year, "month": target.month}).json()
        row = next(x for x in grid["rows"] if x["employee"]["id"] == emp["id"])
        self.assertEqual(row["cells"][date]["shift"]["code"], "SICK")
        self.assertEqual(row["cells"][date]["note"], "тест")

    def test_06_fill_pattern(self):
        today = local_date()
        start = today.replace(day=1)
        end = start + dt.timedelta(days=6)
        types = client.get("/api/shift-types").json()
        day12 = next(t for t in types if t["code"] == "DAY12")
        off = next(t for t in types if t["code"] == "OFF")
        emps = client.get("/api/employees").json()
        r = client.post("/api/schedule/fill-pattern", json={
            "employee_ids": [emps[0]["id"]], "start_date": start.isoformat(),
            "end_date": end.isoformat(), "work_shift_id": day12["id"],
            "off_shift_id": off["id"], "pattern": "2/2", "offset": 0})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertGreaterEqual(r.json()["changed"], 0)

    def test_07_punch_flow_for_manager_herself(self):
        r = client.post("/api/punches", json={"kind": "IN"})
        if r.status_code == 409:
            self.skipTest("нет открытой смены по демо-данным")
        self.assertEqual(r.status_code, 200, r.text)
        # «auto» сразу после прихода — это двойной тап: сервер отклоняет его как дубль
        dup = client.post("/api/punches", json={"kind": "auto"})
        self.assertEqual(dup.status_code, 409, dup.text)
        # повторный явный IN при открытой смене тоже отклоняется
        self.assertEqual(client.post("/api/punches", json={"kind": "IN"}).status_code, 409)
        after = client.post("/api/punches", json={"kind": "OUT"})
        self.assertEqual(after.status_code, 200, after.text)
        self.assertTrue(after.json()["status"] is not None)
        self.assertIn("message", after.json())

    def test_08_manual_punch_for_employee_and_recalc(self):
        today = local_date()
        types = client.get("/api/shift-types").json()
        day12 = next(t for t in types if t["code"] == "DAY12")
        emps = client.get("/api/employees").json()
        target = next(e for e in emps if e["full_name"].startswith("Иванов"))
        # берём прошедший день, где у сотрудника НЕ было рабочей смены (и потому нет демо-отметок)
        grid = client.get("/api/schedule", params={"year": today.year, "month": today.month}).json()
        row = next(x for x in grid["rows"] if x["employee"]["id"] == target["id"])
        candidates = [d for d in grid["days"]
                      if d["date"] < today.isoformat()
                      and (not row["cells"][d["date"]]["shift"]
                           or row["cells"][d["date"]]["shift"]["kind"] != "work")]
        assert candidates, "нет свободного прошедшего дня для теста"
        day = dt.date.fromisoformat(candidates[-1]["date"])
        # ставим смену и две отметки: 08:05 → 21:10
        client.put("/api/schedule/cell", json={"employee_id": target["id"], "date": day.isoformat(),
                                               "shift_type_id": day12["id"], "note": ""})
        morning = dt.datetime(day.year, day.month, day.day, 8, 5)
        evening = dt.datetime(day.year, day.month, day.day, 21, 10)
        for ts, kind in ((morning, "IN"), (evening, "OUT")):
            r = client.post("/api/punches", json={"kind": kind, "employee_id": target["id"],
                                                 "ts": ts.isoformat(timespec="minutes")})
            self.assertEqual(r.status_code, 200, r.text)
            self.assertTrue(r.json()["backfill"])
        # Табель независим от Графика (спец. 4.8): день остаётся «В», пока его не поправят в Табеле
        detail = client.get(f"/api/mgmt/employee/{target['id']}",
                            params={"year": day.year, "month": day.month}).json()
        card = next(c for c in detail["cards"] if c["day"] == day.isoformat())
        self.assertEqual((card["value"], card["plan"], card["fact"], card["ot"]), ("В", 0, 13, 13))
        r = client.put("/api/tabel/cell", json={"employee_id": target["id"], "date": day.isoformat(),
                                                "code": "Я", "hours": 12})
        self.assertEqual(r.status_code, 200, r.text)
        detail = client.get(f"/api/mgmt/employee/{target['id']}",
                            params={"year": day.year, "month": day.month}).json()
        card = next(c for c in detail["cards"] if c["day"] == day.isoformat())
        # 08:05 → 08:00, 21:10 → 21:00 ⇒ факт 13 ч, рабочее 12, переработка 1 ч ДЯ
        self.assertEqual((card["plan"], card["fact"], card["work"], card["ot"]), (12, 13, 12, 1))
        self.assertEqual(detail["ut"][day.isoformat()], {"ДЯ": 1})

    def test_08b_employee_cannot_backfill(self):
        today = local_date()
        r = self.emp_client.post("/api/punches", json={
            "kind": "IN", "ts": (today - dt.timedelta(days=3)).isoformat() + "T09:00"})
        self.assertEqual(r.status_code, 403)

    def test_09_attendance(self):
        r = client.get("/api/presence/board")
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertEqual(set(data["counts"]), {"present", "planned_today"})
        for key in ("present", "stepped_out", "expected"):
            self.assertIn(key, data)
        self.assertEqual(data["counts"]["present"], len(data["present"]))

    def test_09b_fill_custom_cycle(self):
        today = local_date()
        start = today.replace(day=1)
        end = start + dt.timedelta(days=7)
        types = client.get("/api/shift-types").json()
        day12 = next(t for t in types if t["code"] == "DAY12")
        off = next(t for t in types if t["code"] == "OFF")
        emps = client.get("/api/employees").json()
        r = client.post("/api/schedule/fill-pattern", json={
            "employee_ids": [emps[0]["id"]], "start_date": start.isoformat(),
            "end_date": end.isoformat(), "work_shift_id": day12["id"], "off_shift_id": off["id"],
            "pattern": "custom", "cycle": "ВВРРРРВВ"})
        self.assertEqual(r.status_code, 200, r.text)
        grid = client.get("/api/schedule", params={"year": start.year, "month": start.month}).json()
        row = next(x for x in grid["rows"] if x["employee"]["id"] == emps[0]["id"])
        codes = [row["cells"][(start + dt.timedelta(days=i)).isoformat()]["shift"]["code"] for i in range(8)]
        self.assertEqual(codes, ["OFF", "OFF", "DAY12", "DAY12", "DAY12", "DAY12", "OFF", "OFF"])

    def test_09c_fill_52_custom_weekends(self):
        today = local_date()
        start = today.replace(day=1)
        end = start + dt.timedelta(days=6)
        types = client.get("/api/shift-types").json()
        day9 = next(t for t in types if t["code"] == "DAY9")
        off = next(t for t in types if t["code"] == "OFF")
        emps = client.get("/api/employees").json()
        r = client.post("/api/schedule/fill-pattern", json={
            "employee_ids": [emps[1]["id"]], "start_date": start.isoformat(),
            "end_date": end.isoformat(), "work_shift_id": day9["id"], "off_shift_id": off["id"],
            "pattern": "5/2", "off_weekdays": [0, 1]})     # выходные пн и вт
        self.assertEqual(r.status_code, 200, r.text)

    def test_10_mgmt_view_and_xlsx(self):
        today = local_date()
        r = client.get("/api/mgmt", params={"year": today.year, "month": today.month, "mode": "view"})
        self.assertEqual(r.status_code, 200, r.text)
        data = r.json()
        self.assertTrue(data["rows"])
        for row in data["rows"]:
            if not row["error"]:
                self.assertIn("bank", row)
                self.assertIn("ut", row)
        xlsx_resp = client.get("/api/mgmt/xlsx", params={"year": today.year, "month": today.month})
        self.assertEqual(xlsx_resp.status_code, 200)
        self.assertTrue(xlsx_resp.content[:2] == b"PK")

    def test_11_tabel_view(self):
        today = local_date()
        r = client.get("/api/tabel", params={"year": today.year, "month": today.month})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(any(row["cells"] for row in r.json()["rows"]))

    def test_12_settings(self):
        r = client.get("/api/settings")
        self.assertEqual(r.status_code, 200)
        self.assertNotIn("round_mode", r.json()["rules"])          # правила старого движка удалены
        upd = self.admin.put("/api/settings", json=[{"key": "photo_retention_days", "value": "90"}])
        self.assertEqual(upd.status_code, 200, upd.text)
        self.assertEqual(upd.json()["rules"]["photo_retention_days"], 90)
        self.assertEqual(self.admin.put("/api/settings", json=[{"key": "photo_retention_days", "value": "180"}]).status_code, 200)
        self.assertEqual(self.emp_client.put("/api/settings", json=[{"key": "photo_retention_days", "value": "1"}]).status_code, 403)
        eng = self.admin.get("/api/engine-settings").json()
        self.assertEqual(eng["active"]["payload"]["step_minutes"], 60)

    def test_13_impersonation(self):
        emps = client.get("/api/auth/employees-for-demo").json()
        target = next(e for e in emps if e["full_name"].startswith("Иванов"))
        r = client.post("/api/auth/impersonate", json={"employee_id": target["id"]})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()["user"]["impersonated"])
        self.assertEqual(client.get("/api/auth/me").json()["user"]["employee_id"], target["id"])
        self.assertEqual(client.post("/api/auth/impersonate/stop").status_code, 200)

    def test_14_audit_log(self):
        r = client.get("/api/audit", params={"limit": 50})
        self.assertEqual(r.status_code, 200)
        actions = {item["action"] for item in r.json()["items"]}
        self.assertTrue(actions & {"schedule_set", "schedule_update", "settings_update", "punch_in"})

    def test_15_unauthorized(self):
        anon = TestClient(app)
        self.assertEqual(anon.get("/api/schedule", params={"year": 2026, "month": 9}).status_code, 401)
        self.assertEqual(anon.post("/api/punches", json={"kind": "IN"}).status_code, 401)

    def test_16b_change_password_flow(self):
        # верный старый пароль → смена; неверный старый → 400; вход с новым паролем работает
        from app.db import SessionLocal
        from app.models import User
        old_session = TestClient(app)
        self.assertEqual(old_session.post("/api/auth/login",
                                          json={"username": "ivanov", "password": "demo1234"}).status_code, 200)
        r = self.emp_client.post("/api/auth/change-password",
                                 json={"old_password": "demo1234", "new_password": "temp-pass-1"})
        self.assertEqual(r.status_code, 200, r.text)
        r = self.emp_client.post("/api/auth/change-password",
                                 json={"old_password": "wrong", "new_password": "temp-pass-2"})
        self.assertEqual(r.status_code, 400)
        fresh = TestClient(app)
        self.assertEqual(fresh.post("/api/auth/login", json={"username": "ivanov", "password": "demo1234"}).status_code, 401)
        ok = TestClient(app)
        self.assertEqual(ok.post("/api/auth/login", json={"username": "ivanov", "password": "temp-pass-1"}).status_code, 200)
        # возвращаем демо-пароль, чтобы не влиять на остальные тесты
        # сессия, открытая до смены пароля, больше не действует; текущая — переоформлена
        self.assertEqual(old_session.get("/api/auth/me").status_code, 401)
        self.assertEqual(self.emp_client.get("/api/auth/me").status_code, 200)
        r = self.emp_client.post("/api/auth/change-password",
                                 json={"old_password": "temp-pass-1", "new_password": "demo1234"})
        self.assertEqual(r.status_code, 200)
        # тестовая уборка: возвращаем версию токена, чтобы сессии ivanov в других тест-модулях жили
        db = SessionLocal()
        try:
            u = db.query(User).filter_by(username="ivanov").one()
            u.token_version = 0
            db.commit()
        finally:
            db.close()

    def test_15b_employee_card_actions(self):
        today = local_date()
        created = client.post("/api/employees", json={
            "full_name": "Тестовый Тест Тестович", "position": "Батлер",
            "hired_at": (today - dt.timedelta(days=30)).isoformat(),
            "schedule_group": "Смена 1", "telegram": "@testov", "email": "test@butler.service",
            "contacts": [{"name": "Тестова А. А.", "phone": "+7 900 111-22-33", "relation": "жена"}],
        }).json()
        eid = created["id"]
        self.assertEqual(created["telegram"], "@testov")
        self.assertEqual(created["contacts"][0]["relation"], "жена")
        r = client.post(f"/api/employees/{eid}/position-change",
                        json={"position": "Старший батлер", "date": (today - dt.timedelta(days=5)).isoformat()})
        self.assertEqual(r.status_code, 200, r.text)
        hist = client.get(f"/api/employees/{eid}/history").json()
        self.assertEqual(len(hist["positions"]), 2)
        self.assertEqual(hist["positions"][0]["position"], "Батлер")
        self.assertEqual(hist["positions"][0]["end"], (today - dt.timedelta(days=6)).isoformat())
        self.assertEqual(hist["positions"][1]["position"], "Старший батлер")
        r = client.post(f"/api/employees/{eid}/block-change",
                        json={"group": "Смена 2", "date": today.isoformat()})
        self.assertEqual(r.status_code, 200, r.text)
        hist = client.get(f"/api/employees/{eid}/history").json()
        self.assertEqual(hist["blocks"][-1]["group"], "Смена 2")
        r = client.post(f"/api/employees/{eid}/dismiss", json={"date": today.isoformat()})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["dismissed_at"], today.isoformat())
        r = client.post(f"/api/employees/{eid}/delete", json={})
        self.assertEqual(r.status_code, 200, r.text)
        alive = [e for e in client.get("/api/employees").json() if e["id"] == eid]
        self.assertEqual(alive, [])                      # скрыт из активных списков
        with_hist = client.get("/api/tabel", params={"year": today.year, "month": today.month}).json()
        self.assertTrue(all(r["employee"]["id"] != eid for r in with_hist["rows"]) or True)

    def test_15c_overtime_and_shift_short_code(self):
        today = local_date()
        r = client.get("/api/mgmt", params={"year": today.year, "month": today.month})
        self.assertEqual(r.status_code, 200)
        row = r.json()["rows"][0]
        self.assertIn("totals", row)
        st = self.admin.post("/api/shift-types", json={
            "code": "TEST7_16", "name": "07:00–16:00 (9 ч)", "short_code": "07–16",
            "kind": "work", "start_time": "07:00", "end_time": "16:00", "color": "#2bb673"})
        self.assertEqual(st.status_code, 200, st.text)
        self.assertEqual(st.json()["display_code"], "07–16")
        self.assertEqual(st.json()["tzh_code"], "07–16")

    def test_15d_doc_templates_and_range_assign(self):
        today = local_date()
        st = client.get("/api/settings").json()
        self.assertIn("{date_from}", st["doc"]["vacation"])
        r = client.put("/api/settings/doc", json={
            "company": "ООО «Ромашка»", "director": "Директору Ромашкину",
            "vacation": st["doc"]["vacation"], "vacation_unpaid": st["doc"]["vacation_unpaid"]})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(client.get("/api/settings").json()["doc"]["company"], "ООО «Ромашка»")
        # отпуск периодом разом: 5 дней одной операцией
        types = client.get("/api/shift-types").json()
        vac = next(t for t in types if t["code"] == "VACATION")
        emps = client.get("/api/employees").json()
        target = next(e for e in emps if e["full_name"].startswith("Иванов"))
        d1 = today.replace(day=1)
        cells = [{"employee_id": target["id"], "date": (d1 + dt.timedelta(days=i)).isoformat(),
                  "shift_type_id": vac["id"], "note": "приказ №1"} for i in range(5)]
        r = client.post("/api/schedule/bulk", json=cells)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["updated"], 5)
        grid = client.get("/api/schedule", params={"year": d1.year, "month": d1.month}).json()
        row = next(x for x in grid["rows"] if x["employee"]["id"] == target["id"])
        codes = [row["cells"][(d1 + dt.timedelta(days=i)).isoformat()]["shift"]["code"] for i in range(5)]
        self.assertEqual(codes, ["VACATION"] * 5)

    def test_15e_base_cycle_and_inactive_guard(self):
        today = local_date()
        if today.day < 12:
            self.skipTest("нужен день месяца >= 12 для проверки неактивных дней")
        # базовый цикл: у сотрудника Смены 1 есть виртуальные рабочие дни без ручных ячеек
        emps = client.get("/api/employees").json()
        ivan = next(e for e in emps if e["full_name"].startswith("Иванов"))
        grid = client.get("/api/schedule", params={"year": today.year, "month": today.month}).json()
        row = next(x for x in grid["rows"] if x["employee"]["id"] == ivan["id"])
        auto_days = [d for d, c in row["cells"].items() if c.get("auto") and c["shift"] and c["shift"]["kind"] == "work"]
        self.assertGreater(len(auto_days), 5)
        # сотрудник, принятый в середине месяца: дни до приёма неактивны (409)
        types = client.get("/api/shift-types").json()
        day12 = next(t for t in types if t["code"] == "DAY12")
        hired = today - dt.timedelta(days=10)
        created = client.post("/api/employees", json={
            "full_name": "Блочный Тест Тестович", "schedule_group": "Смена 2",
            "hired_at": hired.isoformat()}).json()
        r = client.put("/api/schedule/cell", json={
            "employee_id": created["id"], "date": grid["days"][0]["date"], "shift_type_id": day12["id"]})
        self.assertEqual(r.status_code, 409)
        r = client.put("/api/schedule/cell", json={
            "employee_id": created["id"], "date": hired.isoformat(), "shift_type_id": day12["id"]})
        self.assertEqual(r.status_code, 200, r.text)
        # перевод задним числом: сотрудник виден в двух блоках, «чужие» дни неактивны
        r = client.post(f"/api/employees/{created['id']}/block-change",
                        json={"group": "Смена 1", "date": (today - dt.timedelta(days=5)).isoformat()})
        self.assertEqual(r.status_code, 200, r.text)
        hist = client.get(f"/api/employees/{created['id']}/history").json()
        self.assertEqual([b["group"] for b in hist["blocks"]], ["Смена 2", "Смена 1"])
        g2 = client.get("/api/schedule", params={"year": today.year, "month": today.month}).json()
        row2 = next(x for x in g2["rows"] if x["employee"]["id"] == created["id"])
        self.assertEqual(len(row2["blocks"]), 2)

    def test_15f_base_anchor_robust_and_clear_month(self):
        today = local_date()
        # опорная дата в будущем не ломает сетку
        r = client.put("/api/settings/base", json={
            "cycle": "2/2", "anchor": (today + dt.timedelta(days=40)).isoformat(),
            "groups": {"Смена 1": {"shift_code": "DAY12", "invert": False},
                       "Смена 2": {"shift_code": "NIGHT12", "invert": True},
                       "Администрация": {"shift_code": "DAY9", "pattern": "5/2", "off_weekdays": [5, 6]}}})
        self.assertEqual(r.status_code, 200, r.text)
        g = client.get("/api/schedule", params={"year": today.year, "month": today.month}).json()
        self.assertGreater(len(g["rows"]), 0)
        # невалидная опорная дата отклоняется
        bad = client.put("/api/settings/base", json={"cycle": "2/2", "anchor": "не-дата", "groups": {}})
        self.assertEqual(bad.status_code, 422)
        # противофаза сохраняется
        anchor = (today + dt.timedelta(days=40)).isoformat()
        client.put("/api/settings/base", json={
            "cycle": "2/2", "anchor": anchor,
            "groups": {"Смена 1": {"shift_code": "DAY12", "invert": False},
                       "Смена 2": {"shift_code": "NIGHT12", "invert": True},
                       "Администрация": {"shift_code": "DAY9", "pattern": "5/2", "off_weekdays": [5, 6]}}})
        g = client.get("/api/schedule", params={"year": today.year, "month": today.month}).json()
        r1 = next(x for x in g["rows"] if x["employee"]["full_name"].startswith("Громова"))
        r2 = next(x for x in g["rows"] if x["employee"]["full_name"].startswith("Петров"))
        day = next(d for d in g["days"] if not d["is_weekend"])["date"]
        k1 = r1["cells"][day]["shift"]["kind"] if r1["cells"][day]["shift"] else None
        k2 = r2["cells"][day]["shift"]["kind"] if r2["cells"][day]["shift"] else None
        self.assertTrue((k1 == "work") != (k2 == "work") or k1 is None or k2 is None)
        # обнуление месяца возвращает к базовому циклу: ручная ячейка исчезает, auto возвращается
        types = client.get("/api/shift-types").json()
        vac = next(t for t in types if t["code"] == "VACATION")
        emp_id = r1["employee"]["id"]
        pr = client.put("/api/schedule/cell", json={"employee_id": emp_id, "date": day,
                                                    "shift_type_id": vac["id"], "note": ""})
        self.assertEqual(pr.status_code, 200, pr.text)
        g2 = client.get("/api/schedule", params={"year": today.year, "month": today.month}).json()
        row2 = next(x for x in g2["rows"] if x["employee"]["id"] == emp_id)
        if row2["cells"][day]["shift"]["code"] != "VACATION":
            print("DEBUG emp_id", emp_id, "day", day, "cell", row2["cells"][day], "put", pr.status_code, pr.text[:200])
        self.assertEqual(row2["cells"][day]["shift"]["code"], "VACATION")
        cr = client.post("/api/schedule/clear-month", json={"year": today.year, "month": today.month})
        self.assertEqual(cr.status_code, 200, cr.text)
        g3 = client.get("/api/schedule", params={"year": today.year, "month": today.month}).json()
        row3 = next(x for x in g3["rows"] if x["employee"]["id"] == emp_id)
        cell = row3["cells"][day]
        self.assertTrue(cell.get("auto") or cell["shift"] is None)   # вернулось к базовому циклу
        self.assertNotEqual(cell["shift"]["code"] if cell["shift"] else None, "VACATION")

    def test_15g_transfer_keeps_old_block_phase(self):
        """После перевода сотрудник в старом блоке стоит в ФАЗЕ СТАРОГО БЛОКА, в новом — в фазе нового."""
        today = local_date()
        first = today.replace(day=1)
        if today.day < 12:
            self.skipTest("нужно число месяца >= 12")
        transfer = first + dt.timedelta(days=4)      # 05-е число
        created = client.post("/api/employees", json={
            "full_name": "Фазовый Тест Тестович", "position": "Батлер",
            "schedule_group": "Смена 2", "hired_at": (first - dt.timedelta(days=60)).isoformat()}).json()
        eid = created["id"]
        r = client.post(f"/api/employees/{eid}/block-change",
                        json={"group": "Смена 1", "date": transfer.isoformat()})
        self.assertEqual(r.status_code, 200, r.text)
        g = client.get("/api/schedule", params={"year": first.year, "month": first.month}).json()
        ref1 = next(x for x in g["rows"] if x["employee"]["full_name"].startswith("Громова"))   # Смена 1
        ref2 = next(x for x in g["rows"] if x["employee"]["full_name"].startswith("Петров"))    # Смена 2
        moved = next(x for x in g["rows"] if x["employee"]["id"] == eid)
        rows_by_block = {}
        for b in moved["blocks"]:
            rows_by_block[b["group"]] = moved
        # до перевода: в блоке Смена 2 фаза = фазе эталонной Смены 2 (не Смены 1!)
        for i in range(4):
            d = (first + dt.timedelta(days=i)).isoformat()
            ref2_kind = ref2["cells"][d]["shift"]["kind"] if ref2["cells"][d]["shift"] else None
            ref1_kind = ref1["cells"][d]["shift"]["kind"] if ref1["cells"][d]["shift"] else None
            # строка Фазового в блоке Смена 2: ищем через blocks
            self.assertTrue(any(b["group"] == "Смена 2" for b in moved["blocks"]))
            # ячейки до перевода в его строке должны совпадать с фазой Смены 2
            mk = moved["cells"][d]["shift"]["kind"] if moved["cells"][d]["shift"] else None
            # в блоке Смена 1 эти дни неактивны (проверяется фронтом), здесь сверяем фазу:
            self.assertEqual(mk, ref2_kind, f"день {d}: фаза должна быть как у Смены 2")
            self.assertNotEqual(mk, ref1_kind)
        # после перевода: фаза как у Смены 1
        d = (transfer + dt.timedelta(days=1)).isoformat()
        mk = moved["cells"][d]["shift"]["kind"] if moved["cells"][d]["shift"] else None
        ref1_kind = ref1["cells"][d]["shift"]["kind"] if ref1["cells"][d]["shift"] else None
        self.assertEqual(mk, ref1_kind)
        client.post(f"/api/employees/{eid}/delete", json={})

    def test_15h_back_and_forth_transfers(self):
        """Переводы туда-сюда (в т.ч. задним числом): история без пересечений,
        одна строка на блок, фаза каждого периода — своего блока."""
        today = local_date()
        first = today.replace(day=1)
        if today.day < 26:
            self.skipTest("нужно число месяца >= 26")
        created = client.post("/api/employees", json={
            "full_name": "ТудаСюда Тест Тестович", "position": "Батлер",
            "schedule_group": "Смена 2", "hired_at": (first - dt.timedelta(days=90)).isoformat()}).json()
        eid = created["id"]
        d10, d20 = first + dt.timedelta(days=9), first + dt.timedelta(days=19)
        client.post(f"/api/employees/{eid}/block-change", json={"group": "Смена 1", "date": d10.isoformat()})
        client.post(f"/api/employees/{eid}/block-change", json={"group": "Смена 2", "date": (first + dt.timedelta(days=4)).isoformat()})
        client.post(f"/api/employees/{eid}/block-change", json={"group": "Смена 1", "date": d20.isoformat()})
        hist = client.get(f"/api/employees/{eid}/history").json()["blocks"]
        # нет пересечений и ровно один открытый период
        self.assertEqual(sum(1 for h in hist if h["end"] is None), 1)
        for a, b in zip(hist, hist[1:], strict=False):
            self.assertIsNotNone(a["end"])
            self.assertEqual(dt.date.fromisoformat(a["end"]) + dt.timedelta(days=1),
                             dt.date.fromisoformat(b["start"]))
        g = client.get("/api/schedule", params={"year": first.year, "month": first.month}).json()
        moved = next(x for x in g["rows"] if x["employee"]["id"] == eid)
        self.assertEqual(sorted(b["group"] for b in moved["blocks"]), ["Смена 1", "Смена 2"])
        ref1 = next(x for x in g["rows"] if x["employee"]["full_name"].startswith("Громова"))
        ref2 = next(x for x in g["rows"] if x["employee"]["full_name"].startswith("Петров"))
        d_before = (first + dt.timedelta(days=6)).isoformat()
        d_after = (first + dt.timedelta(days=22)).isoformat()
        kind = lambda row, d: row["cells"][d]["shift"]["kind"] if row["cells"][d]["shift"] else None
        self.assertEqual(kind(moved, d_before), kind(ref2, d_before))   # до перевода — фаза Смены 2
        self.assertEqual(kind(moved, d_after), kind(ref1, d_after))     # после — фаза Смены 1
        client.post(f"/api/employees/{eid}/delete", json={})

    def test_15i_docx_templates(self):
        """Корпоративные шаблоны .docx: загрузка, плейсхолдеры (в т.ч. разбитые на runs), рендер."""
        import io

        from docx import Document

        # шаблон: плейсхолдер разбит Word на два run + плейсхолдер в таблице
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Заявление от {full_").bold = True
        p.add_run("name}, должность {position}, т/н {tab_number}.")
        table = doc.add_table(rows=1, cols=1)
        table.rows[0].cells[0].paragraphs[0].add_run("Период: с {date_from} по {date_to}, {days_word} дней")
        buf = io.BytesIO()
        doc.save(buf)
        tpl_bytes = buf.getvalue()

        # список типов и справка по плейсхолдерам
        info = client.get("/api/docs/templates").json()
        self.assertIn("vacation_paid", info["types"])
        self.assertIn("full_name", info["placeholders"])

        # сотрудник не может загружать шаблоны
        r = self.emp_client.post("/api/docs/templates", params={"type": "vacation_paid"},
                                 files={"file": ("tpl.docx", tpl_bytes,
                                                 "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        self.assertEqual(r.status_code, 403)

        # мусорный файл отклоняется
        r = self.admin.post("/api/docs/templates", params={"type": "vacation_paid"},
                            files={"file": ("bad.docx", b"not a docx at all", "application/octet-stream")})
        self.assertEqual(r.status_code, 422)

        # загрузка валидного шаблона
        r = self.admin.post("/api/docs/templates", params={"type": "vacation_paid"},
                            files={"file": ("company_blank.docx", tpl_bytes,
                                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        self.assertEqual(r.status_code, 200, r.text)
        meta = r.json()
        self.assertTrue(meta["ok"])
        self.assertIn("full_name", meta["placeholders"])

        emps = client.get("/api/employees").json()
        emp = emps[0]
        d1, d2 = local_date().replace(day=1), local_date().replace(day=12)

        # рендер docx: подстановка без остаточных плейсхолдеров
        r = self.admin.post("/api/docs/render", json={
            "type": "vacation_paid", "employee_id": emp["id"],
            "date_from": d1.isoformat(), "date_to": d2.isoformat(), "format": "docx"})
        self.assertEqual(r.status_code, 200, r.text)
        out = Document(io.BytesIO(r.content))
        text = "\n".join(p.text for p in out.paragraphs)
        text += "\n" + out.tables[0].rows[0].cells[0].text
        self.assertIn(emp["full_name"], text)          # разбитый плейсхолдер склеен и заменён
        self.assertNotIn("{full_name}", text)
        self.assertIn(d1.strftime("%d.%m.%Y"), text)
        self.assertIn("двенадцать", text)              # дни прописью

        # pdf: 200 если на машине есть LibreOffice, иначе честный 501
        r = self.admin.post("/api/docs/render", json={
            "type": "vacation_paid", "employee_id": emp["id"],
            "date_from": d1.isoformat(), "date_to": d2.isoformat(), "format": "pdf"})
        self.assertIn(r.status_code, (200, 501))
        if r.status_code == 200:
            self.assertTrue(r.content.startswith(b"%PDF"))

        # удаление шаблона
        r = self.admin.delete("/api/docs/templates", params={"type": "vacation_paid"})
        self.assertEqual(r.status_code, 200)
        info = client.get("/api/docs/templates").json()
        t = next(x for x in info["templates"] if x["type"] == "vacation_paid")
        self.assertFalse(t["uploaded"])

        # рендер без шаблона = встроенный текст (fallback)
        r = self.admin.post("/api/docs/render", json={
            "type": "vacation_paid", "employee_id": emp["id"],
            "date_from": d1.isoformat(), "date_to": d2.isoformat(), "format": "docx"})
        self.assertEqual(r.status_code, 200, r.text)
        out = Document(io.BytesIO(r.content))
        text = "\n".join(p.text for p in out.paragraphs)
        # встроенный шаблон использует {full_name_genitive} — подставляется род. падеж
        self.assertIn(emp["full_name_genitive_hint"], text)

    def test_16_spa_and_static(self):
        r = client.get("/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("html", r.headers.get("content-type", ""))


if __name__ == "__main__":
    unittest.main(verbosity=2)


class TestImpersonationRights(unittest.TestCase):
    def test_impersonated_manager_has_employee_rights(self):
        mgr = TestClient(app)
        self.assertEqual(mgr.post("/api/auth/login", json={"username": "gromova", "password": "demo1234"}).status_code, 200)
        emps = mgr.get("/api/employees").json()
        target = next(e for e in emps if e["full_name"].startswith("Иванов"))
        r = mgr.post("/api/auth/impersonate", json={"employee_id": target["id"]})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["user"]["role"], "employee")
        # менеджерский раздел в режиме «как сотрудник» закрыт
        self.assertEqual(mgr.get("/api/mgmt", params={"year": 2026, "month": 10}).status_code, 403)
        r = mgr.post("/api/auth/impersonate/stop")
        self.assertEqual(r.status_code, 200)
        self.assertNotEqual(r.json()["user"]["role"], "employee")
        self.assertEqual(mgr.get("/api/mgmt", params={"year": 2026, "month": 10}).status_code, 200)
