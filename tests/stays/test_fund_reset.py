"""Начальная загрузка фонда — «заменить всё» (спецификация stays-fund-versions §9.3, §13).

Замена стирает данные шахматки во всей тестовой базе, а она общая для прогона: фикстура
_restore снимает все строки таблиц шахматки до теста и возвращает их после — следующие
тесты видят демо-данные как были. Копии пишутся во временный каталог."""
from __future__ import annotations

import io
import os

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, insert, select, update

from app.core.audit import AuditLog
from app.core.db import Base
from app.modules.stays.adapters.orm import (
    CustomField,
    Escort,
    FundVersion,
    GroupState,
    Label,
    StaysSettings,
    UnitState,
)
from app.modules.stays.application import fund_reset as fr
from app.modules.stays.application import fund_xlsx as fx
from app.modules.stays.application.fund_drafts import FundFileError
from app.modules.stays.application.fund_versions import INITIAL_DATE

from .helpers import API

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ── фикстуры ───────────────────────────────────────────────────────────────


def _rows(s, model) -> list[dict]:
    return [dict(r._mapping) for r in s.execute(select(model.__table__))]


@pytest.fixture(autouse=True)
def _restore(app_stack, tmp_path, monkeypatch):
    """Строки всех таблиц шахматки до теста → после теста ровно они (и in_production)."""
    from app.core.db import SessionLocal

    monkeypatch.setattr(fr, "BACKUP_DIR", tmp_path / "backups")
    with SessionLocal() as s:
        saved = {m: _rows(s, m) for m in fr.WIPE}
        in_prod = s.get(StaysSettings, 1).in_production
    yield
    with SessionLocal() as s:
        fr.wipe(s)
        for m in reversed(fr.WIPE):          # сначала те, на кого ссылаются
            rows = saved[m]
            if not rows:
                continue
            if m is GroupState:
                s.execute(insert(m.__table__), [{**r, "parent_id": None} for r in rows])
                for r in rows:
                    if r["parent_id"] is not None:
                        s.execute(update(GroupState).where(GroupState.version_id == r["version_id"],
                                                           GroupState.group_id == r["group_id"])
                                  .values(parent_id=r["parent_id"]))
            elif m is FundVersion:
                s.execute(insert(m.__table__), [{**r, "base_version_id": None} for r in rows])
                for r in rows:
                    if r["base_version_id"] is not None:
                        s.execute(update(FundVersion).where(FundVersion.id == r["id"])
                                  .values(base_version_id=r["base_version_id"]))
            else:
                s.execute(insert(m.__table__), rows)
        s.get(StaysSettings, 1).in_production = in_prod
        s.commit()


def _export(c) -> bytes:
    r = c.get(f"{API}/fund/export", params={"date": fr.local_date().isoformat()})
    assert r.status_code == 200, r.text
    return r.content


def _edit(data: bytes, fn) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    fn(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _no_ids(data: bytes) -> bytes:
    def clear(wb):
        for name in (fx.SHEET_CATEGORIES, fx.SHEET_TREE, fx.SHEET_FLAGS, fx.SHEET_UNITS):
            ws = wb[name]
            for r in range(2, ws.max_row + 1):
                ws.cell(r, 1).value = None
    return _edit(data, clear)


def _preview(c, data: bytes):
    return c.post(f"{API}/fund/reset", params={"preview": 1}, files={"file": ("fund.xlsx", data, XLSX)})


def _reset(c, data: bytes, fingerprint: str, confirm: str = fr.CONFIRM_WORD):
    return c.post(f"{API}/fund/reset", data={"confirm": confirm, "fingerprint": fingerprint},
                  files={"file": ("real-fund.xlsx", data, XLSX)})


def _counts(db) -> dict:
    db.expire_all()
    return fr.reset_counts(db)


def _backups() -> list[str]:
    d = fr.BACKUP_DIR
    return sorted(os.listdir(d)) if d.exists() else []


# ── права и состояние модуля ───────────────────────────────────────────────


def test_reset_tables_cover_whole_stays_schema():
    """Каждая таблица stays_* решена явно: данные шахматки (удаляются) или настройки (остаются)."""
    stays = {t for t in Base.metadata.tables if t.startswith("stays_")}
    wipe = {m.__tablename__ for m in fr.WIPE}
    keep = {m.__tablename__ for m in fr.KEEP}
    assert not wipe & keep
    assert wipe | keep == stays, stays ^ (wipe | keep)
    assert keep == {"stays_settings", "stays_labels", "stays_fields", "stays_escorts"}


def test_reset_requires_stays_reset_perm(manager_role_client, admin_client, db):
    """Менеджер (stays.setup есть, stays.reset — только у администратора) — 403 везде."""
    data = _export(manager_role_client)
    before = _counts(db)
    c = manager_role_client
    assert _preview(c, data).status_code == 403
    assert _reset(c, data, fr.fingerprint(before)).status_code == 403
    assert c.put(f"{API}/fund/in-production", json={"value": True}).status_code == 403
    assert c.get(f"{API}/fund/backups/stays-20261028-120000.xlsx").status_code == 403
    assert _counts(db) == before
    assert _preview(admin_client, data).status_code == 200


def test_reset_blocked_in_production_409(admin_client, db):
    c = admin_client
    data = _export(c)
    fp = _preview(c, data).json()["fingerprint"]
    r = c.put(f"{API}/fund/in-production", json={"value": True})
    assert r.status_code == 200 and r.json() == {"in_production": True}
    assert c.get(f"{API}/fund").json()["in_production"] is True
    before = _counts(db)
    assert _preview(c, data).status_code == 409
    r = _reset(c, data, fp)
    assert r.status_code == 409 and "эксплуатации" in r.json()["detail"]
    assert _counts(db) == before and _backups() == []
    # выключить обратно — тоже администратор, с записью в аудит
    assert c.put(f"{API}/fund/in-production", json={"value": False}).json() == {"in_production": False}
    logs = db.scalars(select(AuditLog.payload_json).where(AuditLog.action == "stays_in_production")
                      .order_by(AuditLog.id.desc()).limit(2)).all()
    assert logs == ['{"from": true, "to": false}', '{"from": false, "to": true}']
    assert _preview(c, data).status_code == 200


def test_reset_wrong_confirm_422(admin_client, db):
    c = admin_client
    data = _export(c)
    fp = _preview(c, data).json()["fingerprint"]
    before = _counts(db)
    for word in ("", "заменить", "ЗАМЕНИТЬ ", "REPLACE"):
        r = _reset(c, data, fp, confirm=word)
        assert r.status_code == 422, (word, r.text)
        assert "ЗАМЕНИТЬ" in r.json()["detail"]
    assert _counts(db) == before and _backups() == []


def test_reset_checks_under_module_lock(admin_client, db, monkeypatch):
    """Слово и отпечаток сверяются после захвата замка модуля: две вкладки не проскочат вместе."""
    calls: list[str] = []
    real_lock, real_counts = fr.lock_module, fr.reset_counts
    monkeypatch.setattr(fr, "lock_module", lambda s: (calls.append("lock"), real_lock(s)))
    monkeypatch.setattr(fr, "reset_counts", lambda s: (calls.append("counts"), real_counts(s))[1])
    parsed, errors = fx.parse_fund(_export(admin_client))
    assert not errors
    with pytest.raises(fr.FundError, match="ЗАМЕНИТЬ"):
        fr.initial_load(db, parsed, None, real_counts(db), confirm="нет")
    assert calls == ["lock"]
    calls.clear()
    with pytest.raises(fr.FundConflict):
        fr.initial_load(db, parsed, None, {"units": -1}, confirm=fr.CONFIRM_WORD)
    assert calls == ["lock", "counts"]
    db.rollback()


# ── плохой файл ────────────────────────────────────────────────────────────


def test_reset_bad_file_deletes_nothing(admin_client, db):
    c = admin_client
    data = _export(c)
    fp = _preview(c, data).json()["fingerprint"]
    before = _counts(db)

    def bad_category(wb):
        wb[fx.SHEET_UNITS].cell(2, 3).value = "НЕТ-ТАКОЙ"
    bad = _edit(data, bad_category)
    for r in (_preview(c, bad), _reset(c, bad, fp), _reset(c, b"not an xlsx", fp)):
        assert r.status_code == 422, r.text
        assert r.json()["errors"]
    errs = _reset(c, bad, fp).json()["errors"]
    assert errs[0]["sheet"] == fx.SHEET_UNITS and errs[0]["row"] == 2 and "НЕТ-ТАКОЙ" in errs[0]["message"]

    def bad_weight(wb):
        wb[fx.SHEET_CATEGORIES].cell(2, 5).value = 5000
    r = _reset(c, _edit(data, bad_weight), fp)
    assert r.status_code == 422 and r.json()["errors"][0]["sheet"] == fx.SHEET_CATEGORIES
    assert _counts(db) == before and _backups() == []


def test_reset_plan_rejects_unknown_refs_without_keyerror(admin_client, db):
    """Разобранный файл, собранный не parse_fund: неизвестные категория, узел и флаг объекта —
    ошибки файла (FileError), а не KeyError в seed_initial; ничего не удалено."""
    parsed, _ = fx.parse_fund(_export(admin_client))
    u = parsed.units[0]
    u.category_code, u.group_path, u.flags = "НЕТ", ("Нет", "узла"), frozenset({"нет-флага"})
    before = _counts(db)
    with pytest.raises(FundFileError) as e:
        fr.initial_load(db, parsed, None, before, confirm=fr.CONFIRM_WORD)
    msgs = " ".join(x.message for x in e.value.errors)
    assert "категория «НЕТ»" in msgs and "узел «Нет / узла»" in msgs and "флаг «нет-флага»" in msgs
    db.rollback()
    assert _counts(db) == before and _backups() == []


# ── замена ─────────────────────────────────────────────────────────────────


def test_reset_backup_written_before_delete(admin_client, db, monkeypatch):
    c = admin_client
    data = _export(c)
    before = _counts(db)
    assert before["reservations"] and before["guests"]
    fp = _preview(c, data).json()["fingerprint"]
    seen: dict = {}
    real_wipe = fr.wipe

    def checked_wipe(s):
        names = _backups()
        assert len(names) == 1 and fr.BACKUP_NAME.fullmatch(names[0]), names   # без «.part»
        wb = load_workbook(fr.BACKUP_DIR / names[0], read_only=True)
        assert wb.sheetnames == ["Категории", "Дерево", "Флаги", "Объекты", "Брони", "Гости"]
        seen["reservations"] = wb["Брони"].max_row - 1
        seen["guests"] = wb["Гости"].max_row - 1
        seen["units"] = wb["Объекты"].max_row - 1
        wb.close()
        seen["db"] = fr.reset_counts(s)          # данные ещё на месте
        real_wipe(s)
    monkeypatch.setattr(fr, "wipe", checked_wipe)
    r = _reset(c, data, fp)
    assert r.status_code == 200, r.text
    body = r.json()
    assert seen["db"] == before
    assert seen["reservations"] == before["reservations"] and seen["guests"] == before["guests"]
    assert seen["units"] == len(fx.parse_fund(data)[0].units)
    assert body["deleted"] == before and fr.BACKUP_NAME.fullmatch(body["backup"])
    assert body["backup_url"] == f"{API}/fund/backups/{body['backup']}"
    got = c.get(body["backup_url"])
    assert got.status_code == 200 and got.content == (fr.BACKUP_DIR / body["backup"]).read_bytes()
    log = db.scalars(select(AuditLog).where(AuditLog.action == "stays_reset", AuditLog.target == "stays")
                     .order_by(AuditLog.id.desc())).first()
    assert body["backup"] in log.payload_json and "real-fund.xlsx" in log.payload_json


def _other_tables(db) -> dict[str, int]:
    return {t: db.scalar(select(func.count()).select_from(Base.metadata.tables[t]))
            for t in Base.metadata.tables if not t.startswith("stays_") and t != "audit_log"}


def test_reset_keeps_employees_schedule_tabel_audit_and_module_settings(admin_client, db):
    """Удалены только данные шахматки: сотрудники, График, Табель, журнал и прочие таблицы
    приложения на месте; настройки модуля (метки, поля, сопровождение, надбавки и лимит) — тоже.
    Фонд — из файла исходной версией; брони и гости — пусто; шахматка открывается."""
    c = admin_client
    data = _no_ids(_export(c))
    parsed = fx.parse_fund(data)[0]
    others = _other_tables(db)
    assert others["employees"] and others["schedule_shifts"]
    keep = {m: _rows(db, m) for m in (StaysSettings, Label, CustomField, Escort)}
    assert keep[Label] and keep[Escort]
    audit_max = db.scalar(select(func.max(AuditLog.id)))
    audit_n = db.scalar(select(func.count()).select_from(AuditLog))
    fp = _preview(c, data).json()["fingerprint"]
    r = _reset(c, data, fp)
    assert r.status_code == 200, r.text
    db.expire_all()
    assert _other_tables(db) == others
    assert {m: _rows(db, m) for m in keep} == keep
    assert db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.id <= audit_max)) == audit_n
    after = _counts(db)
    assert {k: after[k] for k in ("reservations", "guests", "coverage", "holds")} == dict.fromkeys(
        ("reservations", "guests", "coverage", "holds"), 0)
    assert after["versions"] == 1
    assert r.json()["loaded"] == {"categories": len(parsed.categories), "groups": len(parsed.groups),
                                  "flags": len(parsed.flags), "units": len(parsed.units)}
    v = db.scalar(select(FundVersion))
    assert v.valid_from == INITIAL_DATE and v.status == "published"
    snap = c.get(f"{API}/fund/snapshot", params={"date": fr.local_date().isoformat()}).json()
    assert sorted(u["label"] for u in snap["units"]) == sorted(u.label for u in parsed.units)
    flagged = {u.label: u.flags for u in parsed.units}
    titles = {f["id"]: f["title"] for f in snap["flags"]}
    assert all({titles[f] for f in u["flags"]} == flagged[u["label"]] for u in snap["units"])
    today = fr.local_date().isoformat()
    board = c.get(f"{API}/board/units", params={"first": today, "last": today})
    assert board.status_code == 200 and board.json()["bars"] == []
    assert len(board.json()["units"]) == len(parsed.units)
    # файл без правок снова принимается — копия пригодна и для повторной загрузки
    assert _preview(c, _export(c)).json()["will_load"]["units"] == len(parsed.units)


def test_reset_twice_second_conflicts(admin_client, db):
    """Повторное «Заменить всё» с отпечатком того же предпросмотра (вторая вкладка) — 409."""
    c = admin_client
    data = _export(c)
    fp = _preview(c, data).json()["fingerprint"]
    assert _reset(c, data, fp).status_code == 200
    after = _counts(db)
    r = _reset(c, data, fp)
    assert r.status_code == 409 and "предпросмотр" in r.json()["detail"]
    assert _counts(db) == after
    assert len(_backups()) == 1


def test_reset_ignores_ids_with_warning(admin_client, db):
    c = admin_client
    data = _export(c)
    p = _preview(c, data).json()
    assert p["warnings"] == ["ID в файле будут проигнорированы"]
    assert set(p["counts"]) == {"units", "categories", "groups", "flags", "versions", "reservations",
                                "guests", "coverage", "holds"}
    assert p["will_load"]["units"] == len(fx.parse_fund(data)[0].units)
    assert _preview(c, _no_ids(data)).json()["warnings"] == []
    old_max = db.scalar(select(func.max(UnitState.unit_id)))
    assert _reset(c, data, p["fingerprint"]).status_code == 200
    db.expire_all()
    ids = set(db.scalars(select(UnitState.unit_id)))
    assert ids and min(ids) > old_max           # id из файла не взяты: все объекты новые


def test_backup_download_rejects_path_traversal(admin_client, tmp_path):
    c = admin_client
    fr.BACKUP_DIR.mkdir(parents=True)
    (fr.BACKUP_DIR / "stays-20261028-120000.xlsx").write_bytes(b"ok")
    (fr.BACKUP_DIR / "stays-20261028-120000.xlsx.part").write_bytes(b"part")
    secret = tmp_path / "secret.xlsx"
    secret.write_bytes(b"secret")
    (fr.BACKUP_DIR / "stays-20261028-120001.xlsx").symlink_to(secret)
    assert c.get(f"{API}/fund/backups/stays-20261028-120000.xlsx").content == b"ok"
    for name in ("stays-20261028-120000.xlsx.part", "..%2Fsecret.xlsx", "%2e%2e%2fsecret.xlsx",
                 "stays-20261028-120001.xlsx", "stays-2026102-120000.xlsx",
                 "stays-20261028-120000.xlsx%00", "x%2Fstays-20261028-120000.xlsx",
                 "STAYS-20261028-120000.XLSX", "stays-20261028-120000.xlsx%2F..%2F..%2Fsecret.xlsx"):
        r = c.get(f"{API}/fund/backups/{name}")
        assert r.status_code == 404, (name, r.status_code)
        assert b"secret" not in r.content and b"part" not in r.content
    # «..» клиент схлопывает ещё до запроса — проверяем сам разбор имени
    for name in ("..", "../secret.xlsx", "/etc/passwd", "stays-20261028-120000.xlsx/..", "stays-20261028-120001.xlsx"):
        assert fr.backup_path(name) is None
