"""Загрузка файла фонда черновиком с даты (спецификация stays-fund-versions §9.2–9.3)."""
from __future__ import annotations

import io
import json

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.audit import AuditLog
from app.modules.stays.adapters.orm import Category, FundVersion, Unit, UnitFlag, UnitLineage
from app.modules.stays.application import fund_drafts as fd
from app.modules.stays.application import fund_xlsx as fx

from .helpers import API, FundDraft, kind_id, uniq
from .test_fund_versions_api import _clean  # noqa: F401 — autouse: снять версии и сущности теста

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _today() -> str:
    return fd.local_date().isoformat()


def _seed(c) -> dict:
    """Своя категория (дробный вес), два уровня дерева с порядком, два объекта, флаг."""
    d = FundDraft(c)
    code = "T" + uniq()
    cat = d.call("POST", "categories", code=code, title="Кат " + code, kind_id=kind_id("villa"), weight=1.5,
                 escorted=True, sort_order=7)["id"]
    top = d.call("POST", "groups", title="Корпус " + code, sort_order=3)["id"]
    sub = d.call("POST", "groups", title="Блок", parent_id=top, sort_order=2)["id"]
    u1 = d.call("POST", "units", label=code + "-1", category_id=cat, group_id=sub, floor="2")["id"]
    u2 = d.call("POST", "units", label=code + "-2", category_id=cat)["id"]
    flag = c.post(f"{API}/fund/flags", json={"title": "апи-флаг-" + uniq(), "color": "#112233"}).json()["id"]
    d.call("PUT", "flags", flag_id=flag, unit_ids=[u1], value=True)
    d.publish()
    return {"code": code, "cat": cat, "top": top, "sub": sub, "u1": u1, "u2": u2, "flag": flag}


def _export(c) -> bytes:
    r = c.get(f"{API}/fund/export", params={"date": _today()})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(XLSX)
    return r.content


def _upload(c, data: bytes, day: str | None = None):
    return c.post(f"{API}/fund/import", data={"valid_from": day or _today()},
                  files={"file": ("fund.xlsx", data, XLSX)})


def _edit(data: bytes, fn) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    fn(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _row(ws, value) -> int:
    """Номер строки, где в колонке ID стоит value."""
    return next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == value)


def _draft(db) -> FundVersion | None:
    db.expire_all()
    return db.scalar(select(FundVersion).where(FundVersion.status == "draft"))


def _units(diff: dict, kind: str) -> list[int]:
    return [e["id"] for e in diff["units"] if e["kind"] == kind]


def test_roundtrip_is_empty_draft(manager_role_client, db):
    """Выгрузили и загрузили без правок — черновик без единого изменения: веса (float против
    Decimal), порядок дерева и категорий, флаги, этажи и номера."""
    c = manager_role_client
    s = _seed(c)
    r = _upload(c, _export(c))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["diff"] == {"units": [], "categories": [], "groups": [], "flags": []}
    assert body["retired"] == []
    draft = _draft(db)
    assert draft is not None and body["id"] == draft.id and body["rev"] == draft.rev
    mine = c.get(f"{API}/fund/snapshot", params={"version_id": draft.id}).json()
    base = c.get(f"{API}/fund/snapshot", params={"version_id": draft.base_version_id}).json()
    for k in ("categories", "groups", "units", "flags"):
        assert mine[k] == base[k], k
    unit = next(u for u in mine["units"] if u["id"] == s["u1"])
    assert unit["flags"] == [s["flag"]] and unit["floor"] == "2" and unit["path"] == [s["top"], s["sub"]]
    assert next(x for x in mine["categories"] if x["id"] == s["cat"])["weight"] == 1.5
    log = db.scalars(select(AuditLog).where(AuditLog.action == "stays_fund_import")
                     .order_by(AuditLog.id.desc())).first()
    assert log is not None and log.target == f"fund_draft:{draft.id}"
    payload = json.loads(log.payload_json)
    assert payload["filename"] == "fund.xlsx" and payload["units_flags_changed"] == 0
    assert payload["flags"] == body["flags"] == {"created": [], "renamed": [], "recoloured": []}


def test_missing_row_retires(manager_role_client, db):
    c = manager_role_client
    s = _seed(c)
    data = _edit(_export(c), lambda wb: wb["Объекты"].delete_rows(_row(wb["Объекты"], s["u2"])))
    r = _upload(c, data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert _units(body["diff"], "retired") == [s["u2"]]
    assert body["retired"] == [{"entity": "unit", "id": s["u2"], "name": s["code"] + "-2"}]
    assert not _units(body["diff"], "added") and not _units(body["diff"], "changed")


def test_empty_id_creates(manager_role_client, db):
    c = manager_role_client
    s = _seed(c)
    flag_title = "апи-флаг-" + uniq()

    def add(wb):
        ws = wb["Объекты"]
        col = ws.max_column + 1
        ws.cell(1, col, flag_title)                       # новая колонка — новый флаг
        ws.append([None, s["code"] + "-3", s["code"], f"Корпус {s['code']} / Блок", "5", None])
        ws.cell(ws.max_row, col, "да")
        ws = wb["Дерево"]
        ws.append([None, "Флигель", f"Корпус {s['code']}", 9])
    r = _upload(c, _edit(_export(c), add))
    assert r.status_code == 200, r.text
    diff = r.json()["diff"]
    added = [e for e in diff["units"] if e["kind"] == "added"]
    assert len(added) == 1 and not _units(diff, "retired")
    new = added[0]
    assert new["after"]["label"] == s["code"] + "-3" and new["after"]["category_id"] == s["cat"]
    assert new["after"]["group_id"] == s["sub"] and new["after"]["floor"] == "5"
    flag_id = db.scalar(select(UnitFlag.id).where(UnitFlag.title == flag_title))
    assert flag_id is not None and new["after"]["flags"] == [flag_id]
    assert r.json()["flags"] == {"created": [{"id": flag_id, "title": flag_title}], "renamed": [], "recoloured": []}
    groups = [e for e in diff["groups"] if e["kind"] == "added"]
    assert len(groups) == 1 and groups[0]["after"] == {"parent_id": s["top"], "title": "Флигель", "sort_order": 9}


def test_from_column_lineage(manager_role_client, db):
    """Деление файлом: строка 12134 удалена, добавлены две строки «из 12134»."""
    c = manager_role_client
    s = _seed(c)

    def split(wb):
        ws = wb["Объекты"]
        ws.delete_rows(_row(ws, s["u2"]))
        ws.append([None, s["code"] + "-2a", s["code"], None, None, str(s["u2"])])
        ws.append([None, s["code"] + "-2b", s["code"], None, None, f"{s['u2']}, {s['u1']}"])
    r = _upload(c, _edit(_export(c), split))
    assert r.status_code == 200, r.text
    diff = r.json()["diff"]
    assert _units(diff, "retired") == [s["u2"]]
    new = {e["after"]["label"]: e["id"] for e in diff["units"] if e["kind"] == "added"}
    assert set(new) == {s["code"] + "-2a", s["code"] + "-2b"}
    draft = _draft(db)
    rows = set(db.execute(select(UnitLineage.unit_id, UnitLineage.predecessor_id)
                          .where(UnitLineage.version_id == draft.id)))
    assert rows == {(new[s["code"] + "-2a"], s["u2"]), (new[s["code"] + "-2b"], s["u2"]),
                    (new[s["code"] + "-2b"], s["u1"])}
    shown = {e["after"]["label"]: e.get("from") for e in c.get(f"{API}/fund/draft/diff").json()["units"]
             if e["kind"] == "added"}
    assert shown[s["code"] + "-2a"] == [s["code"] + "-2"]


def test_from_column_unknown_id_422(manager_role_client, db):
    c = manager_role_client
    s = _seed(c)
    data = _edit(_export(c), lambda wb: wb["Объекты"].append(
        [None, s["code"] + "-9", s["code"], None, None, "999999999"]))
    r = _upload(c, data)
    assert r.status_code == 422, r.text
    assert any(e["sheet"] == "Объекты" and "999999999" in e["message"] for e in r.json()["errors"])
    assert _draft(db) is None


def test_unknown_id_422_nothing_applied(manager_role_client, db):
    """Неизвестный ID на каждом из четырёх листов — ошибка с листом и строкой; не применяется
    ничего, даже черновик не остаётся."""
    c = manager_role_client
    _seed(c)
    flag_title = "апи-флаг-" + uniq()
    tops = {m: db.scalar(select(func.max(m.id))) for m in (Category, Unit, UnitFlag)}

    def spoil(wb):
        for name in ("Категории", "Дерево", "Флаги", "Объекты"):
            wb[name].cell(2, 1, 999999999)
        ws = wb["Объекты"]
        ws.cell(1, ws.max_column + 1, flag_title)
        ws.append([None, "NEW-" + uniq(), wb["Категории"].cell(3, 2).value, None, None, None])
    r = _upload(c, _edit(_export(c), spoil))
    assert r.status_code == 422, r.text
    errors = r.json()["errors"]
    assert {e["sheet"] for e in errors} == {"Категории", "Дерево", "Флаги", "Объекты"}
    assert all(e["row"] == 2 and "999999999" in e["message"] for e in errors)
    assert _draft(db) is None
    assert {m: db.scalar(select(func.max(m.id))) for m in tops} == tops
    assert db.scalar(select(UnitFlag.id).where(UnitFlag.title == flag_title)) is None


def test_bad_file_422_no_draft(manager_role_client, db):
    r = _upload(manager_role_client, b"not an xlsx")
    assert r.status_code == 422, r.text
    assert r.json()["errors"][0]["sheet"] == ""
    assert _draft(db) is None


def test_too_big_file_rejected(manager_role_client, db, monkeypatch):
    monkeypatch.setattr(fx, "MAX_FILE_BYTES", 100)
    r = _upload(manager_role_client, b"x" * 101)
    assert r.status_code == 413, r.text
    assert _draft(db) is None


def test_import_with_existing_draft_409(manager_role_client, db):
    c = manager_role_client
    _seed(c)
    data = _export(c)
    d = FundDraft(c)                                      # черновик уже открыт
    r = _upload(c, data)
    assert r.status_code == 409, r.text
    assert _draft(db).rev == d.rev                        # черновик не тронут


def test_code_change_with_same_id_is_rename(manager_role_client, db):
    c = manager_role_client
    s = _seed(c)
    new_code = s["code"] + "X"

    def rename(wb):
        ws = wb["Категории"]
        ws.cell(_row(ws, s["cat"]), 2, new_code)
        ws = wb["Объекты"]
        for uid in (s["u1"], s["u2"]):
            ws.cell(_row(ws, uid), 3, new_code)
    r = _upload(c, _edit(_export(c), rename))
    assert r.status_code == 200, r.text
    diff = r.json()["diff"]
    assert diff["units"] == [] and diff["groups"] == []
    assert [(e["id"], e["kind"], e["before"]["code"], e["after"]["code"]) for e in diff["categories"]] == [
        (s["cat"], "changed", s["code"], new_code)]
    assert r.json()["retired"] == []


def test_import_and_export_require_setup(manager_client, names_client):
    for cl in (manager_client, names_client):
        assert cl.get(f"{API}/fund/export", params={"date": _today()}).status_code == 403
        assert _upload(cl, b"x").status_code == 403


# ── справочник флагов через файл ───────────────────────────────────────────


def _col(ws, title) -> int:
    return next(c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value == title)


def _title(db, fid) -> str:
    db.expire_all()
    return db.get(UnitFlag, fid).title


def _last_import(db) -> dict:
    db.expire_all()
    log = db.scalars(select(AuditLog).where(AuditLog.action == "stays_fund_import")
                     .order_by(AuditLog.id.desc())).first()
    return json.loads(log.payload_json)


def test_flag_rename_via_file_with_header_renamed(manager_role_client, db):
    c = manager_role_client
    s = _seed(c)
    old, new = _title(db, s["flag"]), _title(db, s["flag"]) + "Z"

    def rename(wb):
        wb["Флаги"].cell(_row(wb["Флаги"], s["flag"]), 2, new)
        wb["Объекты"].cell(1, _col(wb["Объекты"], old), new)
    r = _upload(c, _edit(_export(c), rename))
    assert r.status_code == 200, r.text
    assert r.json()["diff"]["flags"] == [] and r.json()["diff"]["units"] == []
    assert r.json()["flags"]["renamed"] == [{"id": s["flag"], "from": old, "to": new}]
    assert _title(db, s["flag"]) == new
    assert _last_import(db)["flags"]["renamed"] == [{"id": s["flag"], "from": old, "to": new}]


def test_flag_rename_via_sheet_only_keeps_units(manager_role_client, db):
    """Переименовали на «Флаги», колонку «Объектов» — нет: колонка со старым названием — тот же
    флаг, дубля под старым названием не появляется, отметки объектов остаются."""
    c = manager_role_client
    s = _seed(c)
    old, new = _title(db, s["flag"]), _title(db, s["flag"]) + "Z"
    top = db.scalar(select(func.max(UnitFlag.id)))
    r = _upload(c, _edit(_export(c), lambda wb: wb["Флаги"].cell(_row(wb["Флаги"], s["flag"]), 2, new)))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["diff"]["flags"] == [] and body["diff"]["units"] == []
    assert body["flags"] == {"created": [], "renamed": [{"id": s["flag"], "from": old, "to": new}], "recoloured": []}
    assert db.scalar(select(func.max(UnitFlag.id))) == top
    assert db.scalar(select(UnitFlag.id).where(UnitFlag.title == old)) is None


def test_flag_rename_with_both_headers_is_error(manager_role_client, db):
    c = manager_role_client
    s = _seed(c)
    old, new = _title(db, s["flag"]), _title(db, s["flag"]) + "Z"

    def both(wb):
        wb["Флаги"].cell(_row(wb["Флаги"], s["flag"]), 2, new)
        ws = wb["Объекты"]
        ws.cell(1, ws.max_column + 1, new)
    r = _upload(c, _edit(_export(c), both))
    assert r.status_code == 422, r.text
    e = r.json()["errors"]
    assert len(e) == 1 and e[0]["sheet"] == "Объекты" and e[0]["row"] == 1 and "не совпадает" in e[0]["message"]
    assert _draft(db) is None and _title(db, s["flag"]) == old


def test_flag_rename_with_unknown_header_is_error(manager_role_client, db):
    """Флаг переименован на листе, а колонка названа ни так, ни так: новый ли это флаг —
    неоднозначно, файл не применяется."""
    c = manager_role_client
    s = _seed(c)
    old = _title(db, s["flag"])

    def typo(wb):
        wb["Флаги"].cell(_row(wb["Флаги"], s["flag"]), 2, old + "Z")
        wb["Объекты"].cell(1, _col(wb["Объекты"], old), old + "Q")
    r = _upload(c, _edit(_export(c), typo))
    assert r.status_code == 422, r.text
    assert [(e["sheet"], e["row"]) for e in r.json()["errors"]] == [("Объекты", 1)]
    assert db.scalar(select(UnitFlag.id).where(UnitFlag.title == old + "Q")) is None


def test_flag_recolour_via_file(manager_role_client, db):
    c = manager_role_client
    s = _seed(c)
    r = _upload(c, _edit(_export(c), lambda wb: wb["Флаги"].cell(_row(wb["Флаги"], s["flag"]), 3, "#abcdef")))
    assert r.status_code == 200, r.text
    assert r.json()["flags"]["recoloured"] == [{"id": s["flag"], "from": "#112233", "to": "#abcdef"}]
    db.expire_all()
    assert db.get(UnitFlag, s["flag"]).color == "#abcdef"
    assert c.delete(f"{API}/fund/versions/{r.json()['id']}").status_code == 200
    data = _export(c)
    row = _row(load_workbook(io.BytesIO(data))["Флаги"], s["flag"])
    bad = _upload(c, _edit(data, lambda wb: wb["Флаги"].cell(row, 3, "red")))
    assert bad.status_code == 422, bad.text
    assert [(e["sheet"], e["row"]) for e in bad.json()["errors"]] == [("Флаги", row)]


def test_flag_swap_titles_via_file(manager_role_client, db):
    """Два флага обменялись названиями (на листе и в заголовках): UNIQUE не мешает, отметки
    объектов остаются за теми же флагами."""
    c = manager_role_client
    s = _seed(c)
    other = c.post(f"{API}/fund/flags", json={"title": "апи-флаг-" + uniq()}).json()["id"]
    a, b = _title(db, s["flag"]), _title(db, other)

    def swap(wb):
        fl, ws = wb["Флаги"], wb["Объекты"]
        fl.cell(_row(fl, s["flag"]), 2, b)
        fl.cell(_row(fl, other), 2, a)
        ca, cb = _col(ws, a), _col(ws, b)
        ws.cell(1, ca, b)
        ws.cell(1, cb, a)
    r = _upload(c, _edit(_export(c), swap))
    assert r.status_code == 200, r.text
    assert r.json()["diff"]["flags"] == [] and r.json()["diff"]["units"] == []
    assert sorted(x["id"] for x in r.json()["flags"]["renamed"]) == sorted([s["flag"], other])
    assert (_title(db, s["flag"]), _title(db, other)) == (b, a)


def test_flag_rename_to_taken_title_error_at_row(manager_role_client, db):
    """Новое название занято флагом, которого на листе нет (архивный неиспользуемый не
    выгружается): ошибка на строке листа «Флаги», а не «строка 0»."""
    c = manager_role_client
    s = _seed(c)
    hidden = "апи-флаг-" + uniq()
    hid = c.post(f"{API}/fund/flags", json={"title": hidden}).json()["id"]
    assert c.put(f"{API}/fund/flags/{hid}/archive", json={"archived": True}).status_code == 200
    data = _export(c)
    row = _row(load_workbook(io.BytesIO(data))["Флаги"], s["flag"])
    r = _upload(c, _edit(data, lambda wb: wb["Флаги"].cell(row, 2, hidden)))
    assert r.status_code == 422, r.text
    assert [(e["sheet"], e["row"]) for e in r.json()["errors"]] == [("Флаги", row)]
    assert str(hid) in r.json()["errors"][0]["message"]


def test_archived_flag_renamed_still_checked(manager_role_client, db):
    c = manager_role_client
    s = _seed(c)
    assert c.put(f"{API}/fund/flags/{s['flag']}/archive", json={"archived": True}).status_code == 200
    old = _title(db, s["flag"])

    def rename_and_set(wb):
        wb["Флаги"].cell(_row(wb["Флаги"], s["flag"]), 2, old + "Z")
        ws = wb["Объекты"]
        col = _col(ws, old)
        ws.cell(1, col, old + "Z")
        ws.cell(_row(ws, s["u2"]), col, "да")
    r = _upload(c, _edit(_export(c), rename_and_set))
    assert r.status_code == 422, r.text
    assert any(e["sheet"] == "Объекты" and "в архиве" in e["message"] for e in r.json()["errors"])
