"""Справочник видов категорий (спецификация 2026-10-28-stays-part1-addendum §1): CRUD и архив,
использованный вид — только в архив (409 на удаление), архивный вид у прежней категории
показывается и фильтруется, а новой категории не даётся; фильтр kind_ids; колонка «вид» файла
фонда — по названию; ревизия 0012 на данных."""
from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, select, update

from app.modules.stays.adapters.orm import CategoryKind, CategoryState

from .helpers import API, FundDraft, kind_id, make_category, uniq
from .test_fund_import import _draft, _edit, _export, _row, _upload
from .test_fund_schema import iso_engine  # noqa: F401 — фикстура миграционного теста
from .test_fund_versions_api import _clean  # noqa: F401 — autouse: снять версии и сущности теста

PREFIX = "тест-вид-"
MONTH = {"first": "2037-02-01", "last": "2037-02-28"}


@pytest.fixture(autouse=True)
def _kinds(db):
    """Виды теста снимаются; категории, ещё ссылающиеся на них, переводятся на «номер»."""
    yield
    db.rollback()
    mine = select(CategoryKind.id).where(CategoryKind.title.like(PREFIX + "%"))
    db.execute(update(CategoryState).where(CategoryState.kind_id.in_(mine)).values(kind_id=kind_id("room")))
    db.execute(delete(CategoryKind).where(CategoryKind.title.like(PREFIX + "%")))
    db.commit()


def _new_kind(c, **body) -> int:
    r = c.post(f"{API}/fund/kinds", json={"title": PREFIX + uniq(), **body})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _snapshot(c, **params) -> dict:
    r = c.get(f"{API}/fund/snapshot", params=params or {"date": _today()})
    assert r.status_code == 200, r.text
    return r.json()


def _today() -> str:
    from app.modules.stays.application import fund_drafts as fd
    return fd.local_date().isoformat()


def _sheet(data: bytes):
    return load_workbook(io.BytesIO(data))["Категории"]


def _kind_row(c, kid: int) -> dict | None:
    return next((k for k in _snapshot(c)["kinds"] if k["id"] == kid), None)


def test_kind_crud(manager_role_client, manager_client, occupancy_client):
    c = manager_role_client
    title = PREFIX + uniq()
    r = c.post(f"{API}/fund/kinds", json={"title": title, "sort_order": 5, "expanded": True})
    assert r.status_code == 200, r.text
    kid = r.json()["id"]
    assert _kind_row(occupancy_client, kid) == {"id": kid, "title": title, "sort_order": 5, "expanded": True,
                                                "archived": False}
    assert c.post(f"{API}/fund/kinds", json={"title": title}).status_code == 409             # название уникально
    assert c.post(f"{API}/fund/kinds", json={"title": ""}).status_code == 422
    assert c.post(f"{API}/fund/kinds", json={"title": "x" * 41}).status_code == 422
    assert c.put(f"{API}/fund/kinds/{kid}", json={"title": title + "!", "sort_order": 7}).status_code == 200
    assert _kind_row(c, kid) == {"id": kid, "title": title + "!", "sort_order": 7, "expanded": False,
                                 "archived": False}
    assert c.put(f"{API}/fund/kinds/{kid}/archive", json={"archived": True}).status_code == 200
    assert _kind_row(c, kid)["archived"] is True
    assert c.put(f"{API}/fund/kinds/999999", json={"title": PREFIX + uniq()}).status_code == 422
    for other in (manager_client, occupancy_client):             # без stays.setup — только чтение
        assert other.post(f"{API}/fund/kinds", json={"title": PREFIX + uniq()}).status_code == 403
        assert other.delete(f"{API}/fund/kinds/{kid}").status_code == 403
    assert c.delete(f"{API}/fund/kinds/{kid}").status_code == 200                           # не используется
    assert _kind_row(c, kid) is None
    assert c.delete(f"{API}/fund/kinds/{kid}").status_code == 422


def test_delete_used_kind_409_archive_ok(manager_role_client, db):
    c = manager_role_client
    kid = _new_kind(c)
    d = FundDraft(c)
    d.call("POST", "categories", code="T" + uniq(), title="Кат", kind_id=kid, weight=1)
    # вид стоит у категории черновика — уже «используется»
    assert c.delete(f"{API}/fund/kinds/{kid}").status_code == 409
    d.publish()
    r = c.delete(f"{API}/fund/kinds/{kid}")
    assert r.status_code == 409 and "архив" in r.json()["detail"]
    assert c.put(f"{API}/fund/kinds/{kid}/archive", json={"archived": True}).status_code == 200
    assert db.get(CategoryKind, kid) is not None


def test_category_requires_known_kind(manager_role_client):
    c = manager_role_client
    d = FundDraft(c)
    base = {"rev": d.rev, "code": "T" + uniq(), "title": "Кат", "weight": 1}
    assert c.post(f"{API}/fund/draft/categories", json=base).status_code == 422             # вид обязателен
    assert c.post(f"{API}/fund/draft/categories", json={**base, "kind_id": 999999}).status_code == 422
    assert c.post(f"{API}/fund/draft/categories", json={**base, "kind": "room"}).status_code == 422
    d.call("POST", "categories", code=base["code"], title="Кат", kind_id=kind_id("villa"), weight=1)
    d.publish()


def test_archived_kind_still_shown(manager_role_client, manager_client):
    """Review Focus 4: вид в архиве у существующей категории — категория показывается и
    фильтруется, правится без смены вида; новой категории и смене вида архивный вид не даётся."""
    c = manager_role_client
    kid = _new_kind(c, expanded=True)
    cat = make_category(c, n_units=2, kind="room")
    other = make_category(c, n_units=1, kind="room")
    d = FundDraft(c)
    d.call("PUT", f"categories/{cat['id']}", code=cat["code"], title=cat["code"], kind_id=kid, weight=1)
    d.publish()
    assert c.put(f"{API}/fund/kinds/{kid}/archive", json={"archived": True}).status_code == 200

    snap = _snapshot(manager_client)
    assert next(x for x in snap["categories"] if x["id"] == cat["id"])["kind_id"] == kid
    assert next(k for k in snap["kinds"] if k["id"] == kid)["archived"] is True
    r = manager_client.get(f"{API}/board/units", params={**MONTH, "kind_ids": [kid]})
    assert r.status_code == 200, r.text
    assert {u["unit_id"] for u in r.json()["units"]} == {u["id"] for u in cat["units"]}

    d = FundDraft(c)
    r = c.post(f"{API}/fund/draft/categories", json={"rev": d.rev, "code": "T" + uniq(), "title": "Нов",
                                                     "kind_id": kid, "weight": 1})
    assert r.status_code == 422 and "архив" in r.json()["detail"]
    r = c.put(f"{API}/fund/draft/categories/{other['id']}", json={"rev": d.rev, "code": other["code"],
                                                                  "title": other["code"], "kind_id": kid, "weight": 1})
    assert r.status_code == 422 and "архив" in r.json()["detail"]
    d.call("PUT", f"categories/{cat['id']}", code=cat["code"], title="Новое имя", kind_id=kid, weight=2)
    d.publish()


def test_filter_by_kind_ids(manager_role_client, manager_client):
    c = manager_role_client
    a, b = _new_kind(c), _new_kind(c)
    ca = make_category(c, n_units=2, kind="room")
    cb = make_category(c, n_units=1, kind="room")
    d = FundDraft(c)
    d.call("PUT", f"categories/{ca['id']}", code=ca["code"], title=ca["code"], kind_id=a, weight=1)
    d.call("PUT", f"categories/{cb['id']}", code=cb["code"], title=cb["code"], kind_id=b, weight=1)
    d.publish()

    def rows(**params) -> set[int]:
        r = manager_client.get(f"{API}/board/units", params={**MONTH, **params})
        assert r.status_code == 200, r.text
        return {u["unit_id"] for u in r.json()["units"]}

    ids_a, ids_b = {u["id"] for u in ca["units"]}, {u["id"] for u in cb["units"]}
    assert rows(kind_ids=[a]) == ids_a
    assert rows(kind_ids=[a, b]) == ids_a | ids_b                           # внутри измерения — «или»
    assert rows(kind_ids=[a], category_ids=[cb["id"]]) == set()             # между измерениями — «и»
    r = manager_client.get(f"{API}/board/units", params={**MONTH, "kinds": ["room"]})
    assert r.status_code == 422 and "kind_ids" in r.json()["detail"]        # прежний параметр — не молча


def test_xlsx_kind_by_title(manager_role_client, db):
    c = manager_role_client
    kid = _new_kind(c)
    title = next(k["title"] for k in _snapshot(c)["kinds"] if k["id"] == kid)
    cat = make_category(c, n_units=1, kind="villa")
    data = _export(c)
    ws = _sheet(data)
    assert ws.cell(1, 4).value == "вид"
    assert ws.cell(_row(ws, cat["id"]), 4).value == "вилла"             # название, а не код

    r = _upload(c, _edit(data, lambda wb: wb["Категории"].cell(_row(wb["Категории"], cat["id"]), 4,
                                                              title.upper())))   # регистр не важен
    assert r.status_code == 200, r.text
    (entry,) = [e for e in r.json()["diff"]["categories"] if e["id"] == cat["id"]]
    assert entry["kind"] == "changed" and entry["before"]["kind_id"] == kind_id("villa")
    assert entry["after"]["kind_id"] == kid
    draft = _draft(db)
    c.delete(f"{API}/fund/versions/{draft.id}")


def test_xlsx_unknown_kind_error(manager_role_client, db):
    c = manager_role_client
    cat = make_category(c, n_units=1, kind="villa")
    archived = _new_kind(c)
    assert c.put(f"{API}/fund/kinds/{archived}/archive", json={"archived": True}).status_code == 200
    archived_title = next(k["title"] for k in _snapshot(c)["kinds"] if k["id"] == archived)

    def bad(wb):
        ws = wb["Категории"]
        ws.cell(_row(ws, cat["id"]), 4, "замок")
        ws.append([None, "N" + uniq(), "Новая", "", 1, "нет", 1])             # новой категории вид обязателен
        ws.append([None, "A" + uniq(), "Архивная", archived_title, 1, "нет", 1])
    r = _upload(c, _edit(_export(c), bad))
    assert r.status_code == 422, r.text
    errors = r.json()["errors"]
    row = _row(_sheet(_export(c)), cat["id"])
    assert any(e["sheet"] == "Категории" and e["row"] == row and "замок" in e["message"]
               and "вилла" in e["message"] for e in errors), errors
    assert any(e["sheet"] == "Категории" and "не указан вид" in e["message"] for e in errors), errors
    assert any(e["sheet"] == "Категории" and "в архиве" in e["message"] for e in errors), errors
    assert _draft(db) is None                                                # ничего не осталось


def test_xlsx_roundtrip_with_kinds(manager_role_client, db):
    """Выгрузили и загрузили без правок — пустой черновик, в том числе с архивным видом у
    категории (он уже стоит у неё — файл его не меняет)."""
    c = manager_role_client
    kid = _new_kind(c)
    cat = make_category(c, n_units=1, kind="room")
    d = FundDraft(c)
    d.call("PUT", f"categories/{cat['id']}", code=cat["code"], title=cat["code"], kind_id=kid, weight=1)
    d.publish()
    assert c.put(f"{API}/fund/kinds/{kid}/archive", json={"archived": True}).status_code == 200
    r = _upload(c, _export(c))
    assert r.status_code == 200, r.text
    assert r.json()["diff"] == {"units": [], "categories": [], "groups": [], "flags": []}
    c.delete(f"{API}/fund/versions/{_draft(db).id}")


def test_kinds_migration_roundtrip(iso_engine):  # noqa: F811
    """0011 → 0012 на данных: четыре прежних вида в справочнике, строки — их id; откат
    возвращает строку kind."""
    from alembic import command
    from sqlalchemy import inspect, text

    from app.bootstrap.schema_version import config

    engine = iso_engine
    command.downgrade(config(), "0011")
    try:
        with engine.begin() as c:
            v = c.execute(text("SELECT id FROM stays_fund_versions WHERE valid_from = '2000-01-01'")).scalar_one()
            for cid, code, kind in ((9301, "MV", "villa"), (9302, "MR", "room"), (9303, "MA", "apartment")):
                c.execute(text("INSERT INTO stays_categories (id) VALUES (:i)"), {"i": cid})
                c.execute(text("INSERT INTO stays_category_states (version_id, category_id, code, title, kind, "
                               "weight, escorted, sort_order) VALUES (:v, :i, :c, :c, :k, 1, 0, 1)"),
                          {"v": v, "i": cid, "c": code, "k": kind})
        command.upgrade(config(), "head")
        assert "kind" not in {col["name"] for col in inspect(engine).get_columns("stays_category_states")}
        with engine.connect() as c:
            kinds = {t: (bool(e), bool(a)) for t, e, a in c.execute(text(
                "SELECT title, expanded, archived FROM stays_category_kinds"))}
            assert kinds == {"вилла": (True, False), "имение": (True, False), "апартаменты": (True, False),
                             "номер": (False, False)}
            got = dict(c.execute(text(
                "SELECT s.code, k.title FROM stays_category_states s JOIN stays_category_kinds k "
                "ON k.id = s.kind_id WHERE s.category_id >= 9300")).all())
            assert got == {"MV": "вилла", "MR": "номер", "MA": "апартаменты"}
        command.downgrade(config(), "0011")
        assert "stays_category_kinds" not in inspect(engine).get_table_names()
        with engine.connect() as c:
            got = dict(c.execute(text("SELECT code, kind FROM stays_category_states WHERE category_id >= 9300")).all())
            assert got == {"MV": "villa", "MR": "room", "MA": "apartment"}
    finally:
        command.upgrade(config(), "head")


def test_reset_plan_kind_by_title(db):
    """Начальная загрузка: вид — по названию из справочника; пустой, неизвестный и архивный — ошибка
    файла с листом и строкой (справочник видов загрузка не трогает — он в KEEP)."""
    from app.modules.stays.application import fund_reset as fr
    from app.modules.stays.application import fund_xlsx as fx

    kinds = [CategoryKind(id=1, title="вилла", archived=False), CategoryKind(id=2, title="старое", archived=True)]
    cats = [fx.ParsedCategory(row_no=n, id=None, code=f"C{n}", title="К", kind=k, weight=1, escorted=False,
                              sort_order=n) for n, k in ((2, "ВИЛЛА"), (3, ""), (4, "замок"), (5, "старое"))]
    plan, errors = fr.plan_load(fx.ParsedFund(categories=cats, groups=[], flags=[], units=[]), kinds)
    assert plan["categories"][0]["kind_id"] == 1
    assert {(e.sheet, e.row) for e in errors} == {("Категории", 3), ("Категории", 4), ("Категории", 5)}
    assert any("замок" in e.message and "вилла" in e.message and "старое" not in e.message for e in errors)
