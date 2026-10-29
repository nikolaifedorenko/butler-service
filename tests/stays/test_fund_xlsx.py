import datetime as dt
import io

import pytest
from openpyxl import Workbook, load_workbook

from app.modules.stays.adapters.orm import UnitFlag
from app.modules.stays.application import fund_xlsx as fx
from app.modules.stays.application.fund_versions import CategoryInfo, FundSnapshot, GroupInfo, UnitInfo


def snap() -> FundSnapshot:
    return FundSnapshot(
        version_id=1, valid_from=dt.date(2026, 1, 1),
        categories={
            1: CategoryInfo(1, "DLX", "Делюкс", 4, 1.5, True, 10, "номер"),
            2: CategoryInfo(2, "CONN", "Коннект", 4, 2.0, False, 20, "номер"),
        },
        groups={
            10: GroupInfo(10, None, "Корпус", 1, (10,)),
            11: GroupInfo(11, 10, "Блок B", 2, (10, 11)),
            12: GroupInfo(12, 11, "Вид на море", 3, (10, 11, 12)),
        },
        units={
            57: UnitInfo(57, "12134", 2, 12, "1", frozenset({5})),
            58: UnitInfo(58, "1215", 1, None, "", frozenset()),
        },
    )


def flags() -> list[UnitFlag]:
    return [UnitFlag(id=5, title="Курящий", color="#112233", sort_order=1, archived=False),
            UnitFlag(id=6, title="Тихий", color="#445566", sort_order=2, archived=False)]


def book(cats=None, tree=None, flg=None, units=None, ucols=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Категории"
    ws.append(["ID", "код", "название", "вид", "вес", "сопр", "порядок"])
    for r in (cats if cats is not None else [[None, "DLX", "Делюкс", "room", 1, "да", 1]]):
        ws.append(r)
    ws = wb.create_sheet("Дерево")
    ws.append(["ID", "путь", "порядок"])
    for r in (tree if tree is not None else [[None, "Корпус", 1], [None, "Корпус / Блок", 2]]):
        ws.append(r)
    ws = wb.create_sheet("Флаги")
    ws.append(["ID", "название", "цвет"])
    for r in (flg or []):
        ws.append(r)
    ws = wb.create_sheet("Объекты")
    ws.append(["ID", "номер", "категория", "узел", "этаж", "из чего", *(ucols or [])])
    for r in (units if units is not None else [[None, "101", "DLX", "Корпус / Блок", "1", None]]):
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def errs(data, **kw):
    parsed, e = fx.parse_fund(data, **kw)
    assert parsed is None
    out: dict = {}
    for x in e:
        k = (x.sheet, x.row)
        out[k] = f"{out[k]}; {x.message}" if k in out else x.message
    return out


def test_export_has_four_sheets_and_id_column():
    wb = load_workbook(io.BytesIO(fx.export_fund(snap(), flags())))
    assert wb.sheetnames == ["Категории", "Дерево", "Флаги", "Объекты"]
    ws = wb["Объекты"]
    assert [c.value for c in ws[1]][1:] == ["номер", "категория (код)", "узел (путь)", "этаж",
                                            "из чего (ID через запятую)", "Курящий", "Тихий"]
    assert ws["A2"].value == 57 and ws["D2"].value == "Корпус / Блок B / Вид на море"
    assert ws["G2"].value == "да" and ws["H2"].value == "нет"
    assert ws["A2"].protection.locked and not ws["B2"].protection.locked
    assert ws["A2"].fill.fgColor.rgb.endswith("D9D9D9") and ws.protection.sheet
    assert wb["Дерево"]["B4"].value == "Корпус / Блок B / Вид на море"


def test_parse_roundtrip_equals_snapshot():
    s = snap()
    parsed, e = fx.parse_fund(fx.export_fund(s, flags()))
    assert e == [] and parsed is not None
    assert {(c.id, c.code, c.title, c.kind, c.weight, c.escorted, c.sort_order) for c in parsed.categories} == {
        (c.id, c.code, c.title, c.kind_title, c.weight, c.escorted, c.sort_order) for c in s.categories.values()}
    assert {(g.id, g.path, g.sort_order) for g in parsed.groups} == {
        (g.id, s.path_titles(g.id), g.sort_order) for g in s.groups.values()}
    assert {(f.id, f.title, f.color) for f in parsed.flags} == {(5, "Курящий", "#112233"), (6, "Тихий", "#445566")}
    u = {x.id: x for x in parsed.units}
    assert u[57].label == "12134" and u[57].category_code == "CONN" and u[57].floor == "1"
    assert u[57].group_path == ("Корпус", "Блок B", "Вид на море") and u[57].flags == {"Курящий"}
    assert u[58].group_path is None and u[58].flags == frozenset() and u[57].predecessors == []


def test_new_flag_column_becomes_flag():
    parsed, e = fx.parse_fund(book(units=[[None, "101", "DLX", None, "1", None, "да"]], ucols=["Балкон"]))
    assert e == []
    assert [(f.id, f.title) for f in parsed.flags] == [(None, "Балкон")]
    assert parsed.units[0].flags == {"Балкон"}


def test_predecessors_and_yes_no_variants():
    parsed, e = fx.parse_fund(book(
        cats=[[None, "DLX", "Д", "room", 1, "YES", 1]],
        units=[[None, "1", "DLX", None, "", "57, 58", "TRUE", "0"]], ucols=["А", "Б"]))
    assert e == [] and parsed.units[0].predecessors == [57, 58]
    assert parsed.units[0].flags == {"А"} and parsed.categories[0].escorted


def test_duplicate_id_number_code():
    e = errs(book(cats=[[1, "A", "a", "k", 1, "да", 1], [1, "A", "b", "k", 1, "да", 2]],
                  units=[[7, "1", "A", None, "", None], [7, "1", "A", None, "", None]]))
    assert "ID 1 повторяется" in e[("Категории", 3)]
    assert ("Объекты", 3) in e
    parsed_msgs = " ".join(e.values())
    assert "код «A» повторяется" in parsed_msgs and "номер «1» повторяется" in parsed_msgs


def test_path_without_parent():
    e = errs(book(tree=[[None, "Корпус / Блок / Вид", 1]]))
    assert "нет родителя" in e[("Дерево", 2)]


def test_cycle_is_unrepresentable_but_repeating_node_is_ok_when_parent_exists():
    # В файле дерево задано путями, цикл в нём невозможен; «А / Б / А» — три разных узла.
    parsed, e = fx.parse_fund(book(tree=[[None, "А", 1], [None, "А / Б", 1], [None, "А / Б / А", 1]], units=[]))
    assert e == [] and len(parsed.groups) == 3


def test_yes_no_typo():
    e = errs(book(units=[[None, "1", "DLX", None, "", None, "дда"]], ucols=["Балкон"]))
    assert "ожидается да или нет" in e[("Объекты", 2)]
    e = errs(book(cats=[[None, "DLX", "Д", "room", 1, "возможно", 1]]))
    assert ("Категории", 2) in e


def test_unknown_category_and_node():
    e = errs(book(units=[[None, "1", "ZZZ", None, "", None], [None, "2", "DLX", "Нет такого", "", None]]))
    assert "неизвестная категория" in e[("Объекты", 2)] and "неизвестный узел" in e[("Объекты", 3)]


def test_unknown_predecessor_id_and_bad_format():
    data = book(units=[[None, "1", "DLX", None, "", "99"], [None, "2", "DLX", None, "", "x"]])
    e = errs(data, known_unit_ids={57})
    assert "неизвестный ID 99" in e[("Объекты", 2)] and "не целое" in e[("Объекты", 3)]


def test_bad_inputs_do_not_raise():
    assert fx.parse_fund(b"not a zip")[1][0].message.startswith("не файл Excel")
    assert fx.parse_fund(b"")[0] is None
    wb = Workbook()
    wb.active.title = "Категории"
    b = io.BytesIO()
    wb.save(b)
    msgs = [x.message for x in fx.parse_fund(b.getvalue())[1]]
    assert any("нет листа «Дерево»" in m for m in msgs)


def test_size_and_row_caps(monkeypatch):
    monkeypatch.setattr(fx, "MAX_FILE_BYTES", 100)
    assert "больше" in fx.parse_fund(book())[1][0].message
    monkeypatch.undo()
    monkeypatch.setattr(fx, "MAX_ROWS", 2)
    units = [[None, str(i), "DLX", None, "", None] for i in range(5)]
    parsed, e = fx.parse_fund(book(units=units))
    assert parsed is None and "слишком много строк" in e[0].message


def test_formula_guard_roundtrip():
    s = snap()
    s2 = FundSnapshot(s.version_id, s.valid_from, s.categories, s.groups,
                      {1: UnitInfo(1, "=1+1", 1, None, "", frozenset())})
    parsed, e = fx.parse_fund(fx.export_fund(s2, flags()))
    assert e == [] and parsed.units[0].label == "=1+1"


@pytest.mark.parametrize("val", ["1", 1, 1.0])
def test_numeric_cells_as_text(val):
    parsed, e = fx.parse_fund(book(units=[[None, val, "DLX", None, 3, None]]))
    assert e == [] and parsed.units[0].label == "1" and parsed.units[0].floor == "3"


def test_formula_flag_title_exported_as_text_and_parsed_back():
    fl = [UnitFlag(id=5, title="=1+1", color="#112233", sort_order=1, archived=False)]
    data = fx.export_fund(snap(), fl)
    ws = load_workbook(io.BytesIO(data))["Объекты"]
    assert ws.cell(row=1, column=7).value == "'=1+1"
    parsed, e = fx.parse_fund(data)
    assert e == [] and {f.title for f in parsed.flags} == {"=1+1"}
    assert [u.flags for u in parsed.units if u.id == 57] == [frozenset({"=1+1"})]


def test_too_many_flag_columns():
    cols = [f"F{i}" for i in range(fx.MAX_FLAG_COLUMNS + 1)]
    e = errs(book(units=[], ucols=cols))
    assert "слишком много колонок" in e[("Объекты", 1)]


def test_predecessors_on_row_with_id_is_error():
    e = errs(book(units=[[7, "1", "DLX", None, "", "57"]]))
    assert "только у новых" in e[("Объекты", 2)]
