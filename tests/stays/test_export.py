"""Шахматка, задача 18: выгрузка в Excel — обезличенная и полная (§8), журнал аудита."""
from __future__ import annotations

import io
import json

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.audit import AuditLog

from .helpers import API, create_res, make_category, make_guest

Q = {"first": "2037-02-01", "last": "2037-02-07"}


def _setup(sup, mgr):
    cat = make_category(sup, n_units=2)
    g = make_guest(mgr, "Выгрузкин")
    res = create_res(mgr, cat, "2037-02-02", "2037-02-05", [(g["id"], "main", None, None)],
                     note="=HYPERLINK(\"http://x\")", customer="Заказчиков")
    return cat, g, res


def _cells(content: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(content))
    return [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]


def test_anon_export_has_no_guest_data(occupancy_client, manager_role_client, manager_client, db):
    cat, g, res = _setup(manager_role_client, manager_client)
    r = occupancy_client.get(f"{API}/export", params={**Q, "kind": "anon"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxml")
    cells = _cells(r.content)
    assert not any("Выгрузкин" in c or "Заказчиков" in c or res["number"] in c for c in cells)
    assert any(cat["units"][0]["label"] == c for c in cells) and "занято" in cells
    wb = load_workbook(io.BytesIO(r.content))
    assert [ws.title for ws in wb.worksheets] == ["Шахматка", "Брони", "Сводка"]
    row = db.scalars(select(AuditLog).where(AuditLog.action == "stays_export").order_by(AuditLog.id.desc())).first()
    assert json.loads(row.payload_json) == {"kind": "anon", **Q}


def test_full_export_requires_export_and_guests(manager_client, manager_role_client, occupancy_client):
    assert manager_client.get(f"{API}/export", params={**Q, "kind": "full"}).status_code == 403    # старший
    assert occupancy_client.get(f"{API}/export", params={**Q, "kind": "full"}).status_code == 403


def test_full_export_has_names_and_escapes_formulas(manager_role_client, manager_client):
    cat, g, res = _setup(manager_role_client, manager_client)
    r = manager_role_client.get(f"{API}/export", params={**Q, "kind": "full"})
    assert r.status_code == 200
    cells = _cells(r.content)
    assert any("Выгрузкин" in c for c in cells) and any(res["number"] == c for c in cells)
    assert "'=HYPERLINK(\"http://x\")" in cells and "=HYPERLINK(\"http://x\")" not in cells


def test_export_window_limit(manager_role_client):
    r = manager_role_client.get(f"{API}/export", params={"first": "2037-01-01", "last": "2037-06-01", "kind": "anon"})
    assert r.status_code == 422
