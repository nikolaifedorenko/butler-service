"""Шахматка, задача 8: брони — документ, состав во времени, группы, статусы, no-show, версии."""
from __future__ import annotations

import json

from sqlalchemy import select

from app.core.audit import AuditLog

from .helpers import API, create_res, make_category, make_guest, res_body, units_of


def _put(c, res, **changes):
    body = {**{k: res[k] for k in ("number", "check_in", "check_out", "category_id", "early", "late",
                                   "group_id", "note", "declined_butler")},
            "moves": res["moves"], "occupants": res["occupants"], "companions": res["companions"],
            "logistics": res["logistics"], "version": res["version"], **changes}
    return c.put(f"{API}/reservations/{res['id']}", json=body)


def _status(c, res, status, **kw):
    return c.post(f"{API}/reservations/{res['id']}/status",
                  json={"status": status, "version": res["version"], **kw})


def test_create_reservation_gets_unit(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    g = make_guest(manager_client)
    res = create_res(manager_client, cat, "2027-03-02", "2027-03-06", [(g["id"], "main", None, None)])
    assert units_of(res) == [cat["units"][0]["id"]] and res["version"] == 1
    assert res["status"] == "confirmed"


def test_two_mains_same_day_refused(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    a, b = make_guest(manager_client), make_guest(manager_client)
    r = manager_client.post(f"{API}/reservations", json=res_body(
        cat, "2027-03-02", "2027-03-06", [(a["id"], "main", None, None), (b["id"], "main", "2027-03-04", None)]))
    assert r.status_code == 422 and r.json()["detail"] == "На 04.03 два главных гостя"


def test_gap_without_main_refused(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    a = make_guest(manager_client)
    r = manager_client.post(f"{API}/reservations", json=res_body(
        cat, "2027-03-02", "2027-03-06", [(a["id"], "main", "2027-03-02", "2027-03-04")]))
    assert r.status_code == 422 and r.json()["detail"] == "На 04.03 нет главного гостя"


def test_main_guest_changes_midstay(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    a, b = make_guest(manager_client), make_guest(manager_client)
    res = create_res(manager_client, cat, "2027-03-10", "2027-03-18", [
        (a["id"], "main", "2027-03-10", "2027-03-14"), (b["id"], "main", "2027-03-14", "2027-03-18")])
    assert [(o["guest_id"], o["date_from"], o["date_to"]) for o in res["occupants"]] == [
        (a["id"], "2027-03-10", "2027-03-14"), (b["id"], "2027-03-14", "2027-03-18")]


def test_occupant_outside_dates_refused(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    a = make_guest(manager_client)
    r = manager_client.post(f"{API}/reservations", json=res_body(
        cat, "2027-03-02", "2027-03-06", [(a["id"], "main", "2027-03-01", None)]))
    assert r.status_code == 422 and r.json()["detail"] == "Даты гостя вне дат брони"


def test_bad_dates_refused(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    r = manager_client.post(f"{API}/reservations", json=res_body(cat, "2027-03-06", "2027-03-06"))
    assert r.status_code == 422 and r.json()["detail"] == "Выезд должен быть позже заезда"


def test_stale_version_conflict(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    res = create_res(manager_client, cat, "2027-03-02", "2027-03-06")
    assert _put(manager_client, res, note="первый").status_code == 200
    r = _put(manager_client, res, note="второй")
    assert r.status_code == 409 and r.json()["detail"] == "Бронь изменилась — обновите"


def test_occupant_ids_kept_on_update(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    a = make_guest(manager_client)
    res = create_res(manager_client, cat, "2027-03-02", "2027-03-06", [(a["id"], "main", None, None)])
    r = _put(manager_client, res, note="правка")
    assert [o["id"] for o in r.json()["result"]["occupants"]] == [o["id"] for o in res["occupants"]]


def test_cancel_frees_unit(manager_role_client, manager_client):
    cat = make_category(manager_role_client, n_units=1)
    a = create_res(manager_client, cat, "2027-04-02", "2027-04-06")
    b = create_res(manager_client, cat, "2027-04-03", "2027-04-05")
    assert units_of(b) == [None]
    assert _status(manager_client, a, "cancelled").status_code == 200
    b2 = manager_client.get(f"{API}/reservations/{b['id']}").json()
    assert units_of(b2) == [cat["units"][0]["id"]]
    a2 = manager_client.get(f"{API}/reservations/{a['id']}").json()
    assert units_of(a2) == [None]


def test_bad_transition_refused(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    a = create_res(manager_client, cat, "2027-04-02", "2027-04-06")
    r = _status(manager_client, a, "checked_out")
    assert r.status_code == 422 and r.json()["detail"] == "Нельзя перевести бронь из «подтверждена» в «выехал»"


def test_no_show_before_check_in_refused(manager_role_client, manager_client, today):
    cat = make_category(manager_role_client)
    a = create_res(manager_client, cat, "2027-05-10", "2027-05-14")
    today("2027-05-09")
    r = _status(manager_client, a, "no_show")
    assert r.status_code == 422 and r.json()["detail"] == "Отметить «не заехал» можно с даты заезда"


def test_no_show_frees_unit_and_return_with_new_date(manager_role_client, manager_client, today):
    cat = make_category(manager_role_client, n_units=1)
    a = create_res(manager_client, cat, "2027-05-10", "2027-05-14")
    today("2027-05-10")
    r = _status(manager_client, a, "no_show")
    assert r.status_code == 200 and units_of(r.json()["result"]) == [None]
    back = _status(manager_client, r.json()["result"], "confirmed")
    assert back.status_code == 422 and back.json()["detail"] == "Укажите новую дату заезда"
    back = _status(manager_client, r.json()["result"], "confirmed", new_check_in="2027-05-11")
    assert back.status_code == 200
    res = back.json()["result"]
    assert res["check_in"] == "2027-05-11" and units_of(res) == [cat["units"][0]["id"]]


def test_group_one_unit(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    gid = manager_client.post(f"{API}/groups", json={"title": "Семья"}).json()["id"]
    a = create_res(manager_client, cat, "2027-06-01", "2027-06-05", group_id=gid)
    b = create_res(manager_client, cat, "2027-06-02", "2027-06-05", group_id=gid)
    a = manager_client.get(f"{API}/reservations/{a['id']}").json()
    assert units_of(a) == units_of(b) == [cat["units"][0]["id"]]
    assert a["warnings"] == [] and b["warnings"] == []


def test_group_no_show_one_of_two_keeps_unit(manager_role_client, manager_client, today):
    cat = make_category(manager_role_client, n_units=1)
    gid = manager_client.post(f"{API}/groups", json={"title": "Пара"}).json()["id"]
    a = create_res(manager_client, cat, "2027-06-10", "2027-06-14", group_id=gid)
    b = create_res(manager_client, cat, "2027-06-10", "2027-06-14", group_id=gid)
    today("2027-06-10")
    assert _status(manager_client, a, "no_show").status_code == 200
    b2 = manager_client.get(f"{API}/reservations/{b['id']}").json()
    assert units_of(b2) == [cat["units"][0]["id"]]
    r = _status(manager_client, b2, "no_show")
    assert units_of(r.json()["result"]) == [None]
    other = create_res(manager_client, cat, "2027-06-11", "2027-06-13")
    assert units_of(other) == [cat["units"][0]["id"]]


def test_move_gets_two_segments(manager_role_client, manager_client):
    cat = make_category(manager_role_client, n_units=2)
    res = create_res(manager_client, cat, "2027-07-01", "2027-07-09",
                     moves=[{"move_date": "2027-07-05", "category_id": cat["id"]}])
    assert units_of(res) == [cat["units"][0]["id"], cat["units"][1]["id"]]


def test_pin_respected(manager_role_client, manager_client):
    cat = make_category(manager_role_client, n_units=2)
    res = create_res(manager_client, cat, "2027-08-01", "2027-08-05", pins={"0": cat["units"][1]["id"]})
    assert units_of(res) == [cat["units"][1]["id"]] and res["segments"][0]["pinned"] is True


def test_hold_triggers_reallocation(manager_role_client, manager_client):
    cat = make_category(manager_role_client, n_units=2)
    res = create_res(manager_client, cat, "2027-08-10", "2027-08-14")
    assert units_of(res) == [cat["units"][0]["id"]]
    h = manager_client.post(f"{API}/ooo", json={"unit_id": cat["units"][0]["id"], "start": "2027-08-09",
                                                  "end": "2027-08-20", "start_half": "AM", "end_half": "PM", "reason": "Ремонт"})
    assert h.status_code == 200, h.text
    assert units_of(manager_client.get(f"{API}/reservations/{res['id']}").json()) == [cat["units"][1]["id"]]
    assert manager_client.delete(f"{API}/ooo/{h.json()['result']['id']}", params={"version": 1}).status_code == 200


def test_hold_bad_dates_refused(manager_role_client, manager_client):
    cat = make_category(manager_role_client, n_units=1)
    r = manager_client.post(f"{API}/ooo", json={"unit_id": cat["units"][0]["id"], "start": "2027-08-09",
                                                  "end": "2027-08-01", "start_half": "AM", "end_half": "PM", "reason": ""})
    assert r.status_code == 422 and r.json()["detail"] == "Конец OOO раньше начала"


def test_reservation_audit_has_diff(manager_role_client, manager_client, db):
    cat = make_category(manager_role_client)
    res = create_res(manager_client, cat, "2027-11-01", "2027-11-05")
    assert _put(manager_client, res, check_out="2027-11-07").status_code == 200
    row = db.scalars(select(AuditLog).where(AuditLog.action == "stays_reservation",
                                            AuditLog.target == f"reservation:{res['id']}")
                     .order_by(AuditLog.id.desc())).first()
    assert json.loads(row.payload_json)["diff"]["check_out"] == ["2027-11-05", "2027-11-07"]


def test_edit_requires_perm(manager_role_client, manager_client, employee_client, mgmt_client):
    cat = make_category(manager_role_client)
    body = res_body(cat, "2027-09-01", "2027-09-03")
    assert employee_client.post(f"{API}/reservations", json=body).status_code == 403
    assert mgmt_client.post(f"{API}/reservations", json=body).status_code == 403


def _fund_with_new_category(setup_client):
    """Категория, появившаяся в фонде только с версии «сегодня» — позже дат тестовой брони в прошлом."""
    from .helpers import FundDraft, kind_id, uniq
    d = FundDraft(setup_client)
    code = "P" + uniq()
    cid = d.call("POST", "categories", code=code, title=code, kind_id=kind_id("villa"), escorted=False,
                 weight=1)["id"]
    d.call("POST", "units", label=code + "-1", category_id=cid)
    d.publish()
    return {"id": cid, "code": code}


def test_paid_category_defaults_to_booking_category(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    g = make_guest(manager_client)
    res = create_res(manager_client, cat, "2027-05-02", "2027-05-05", [(g["id"], "main", None, None)])
    assert res["paid_category_id"] == cat["id"]
    other = make_category(manager_role_client)
    res2 = create_res(manager_client, cat, "2027-06-02", "2027-06-05", [(g["id"], "main", None, None)],
                      paid_category_id=other["id"])
    assert res2["paid_category_id"] == other["id"]
    r = _put(manager_client, res2, paid_category_id=None)
    assert r.status_code == 200 and r.json()["result"]["paid_category_id"] == cat["id"]


def test_paid_category_must_exist_on_check_in(manager_role_client, manager_client):
    cat = make_category(manager_role_client)
    g = make_guest(manager_client)
    late = _fund_with_new_category(manager_role_client)
    r = manager_client.post(f"{API}/reservations", json=res_body(
        cat, "2020-05-02", "2020-05-05", [(g["id"], "main", None, None)], paid_category_id=late["id"]))
    assert r.status_code == 422, r.text
    r = manager_client.post(f"{API}/reservations", json=res_body(
        cat, "2027-05-02", "2027-05-05", [(g["id"], "main", None, None)], paid_category_id=9_999_999))
    assert r.status_code == 422


def test_upgrade_mark_in_card_and_bookings(manager_role_client, manager_client):
    cat, paid = make_category(manager_role_client), make_category(manager_role_client)
    g = make_guest(manager_client)
    res = create_res(manager_client, cat, "2027-07-02", "2027-07-05", [(g["id"], "main", None, None)],
                     paid_category_id=paid["id"])
    card = manager_client.get(f"{API}/reservations/{res['id']}/card").json()
    assert card["upgrade"] == f"апгрейд: оплачена {paid['code']}" and card["paid_category"] == paid["code"]
    rows = manager_client.get(f"{API}/bookings", params={"first": "2027-07-01", "last": "2027-07-10"}).json()["rows"]
    assert [x["upgrade"] for x in rows if x["id"] == res["id"]] == [f"апгрейд: оплачена {paid['code']}"]
    plain = create_res(manager_client, cat, "2027-08-02", "2027-08-05", [(g["id"], "main", None, None)])
    assert manager_client.get(f"{API}/reservations/{plain['id']}/card").json()["upgrade"] == ""


def test_purpose_removed_from_api_and_export(manager_role_client, manager_client):
    import io

    from openpyxl import load_workbook
    cat = make_category(manager_role_client)
    g = make_guest(manager_client)
    res = create_res(manager_client, cat, "2027-09-02", "2027-09-05", [(g["id"], "main", None, None)],
                     purpose="отпуск", paid_by="гость")
    assert "purpose" not in res and "paid_by" not in res
    card = manager_client.get(f"{API}/reservations/{res['id']}/card").json()
    assert "purpose" not in card and "paid_by" not in card
    r = manager_role_client.get(f"{API}/export", params={"first": "2027-09-01", "last": "2027-09-10", "kind": "full"})
    heads = [c.value for ws in load_workbook(io.BytesIO(r.content)).worksheets for row in ws.iter_rows(max_row=1)
             for c in row]
    assert not any("повод" in str(h).lower() for h in heads)


def _upgrade(c, rid) -> str:
    return c.get(f"{API}/reservations/{rid}/card").json()["upgrade"]


def test_paid_category_follows_category_change(manager_role_client, manager_client):
    cat, other = make_category(manager_role_client), make_category(manager_role_client)
    g = make_guest(manager_client)
    res = create_res(manager_client, cat, "2027-10-02", "2027-10-05", [(g["id"], "main", None, None)])
    r = _put(manager_client, res, category_id=other["id"])        # эхо старой нагрузки: оплата не указана
    assert r.status_code == 200, r.text
    assert r.json()["result"]["paid_category_id"] == other["id"] and _upgrade(manager_client, res["id"]) == ""


def test_group_join_keeps_default_paid(manager_role_client, manager_client):
    cat, other, third = (make_category(manager_role_client) for _ in range(3))
    g = make_guest(manager_client)
    gid = manager_client.post(f"{API}/groups", json={"title": "Семья"}).json()["id"]
    create_res(manager_client, cat, "2027-11-02", "2027-11-05", group_id=gid)
    b = create_res(manager_client, other, "2027-11-02", "2027-11-05", [(g["id"], "main", None, None)],
                   group_id=gid)                                   # категория группы перекрывает свою
    assert b["category_id"] == cat["id"] and _upgrade(manager_client, b["id"]) == ""
    c = create_res(manager_client, other, "2027-11-02", "2027-11-05", group_id=gid, paid_category_id=third["id"])
    assert _upgrade(manager_client, c["id"]) == f"апгрейд: оплачена {third['code']}"


def test_paid_category_before_fund_start_refused(manager_role_client, manager_client):
    cat, other = make_category(manager_role_client), make_category(manager_role_client)
    g = make_guest(manager_client)
    r = manager_client.post(f"{API}/reservations", json=res_body(
        cat, "1990-05-02", "1990-05-05", [(g["id"], "main", None, None)], paid_category_id=other["id"]))
    assert r.status_code == 422, r.text


def test_group_join_form_payload_no_false_upgrade(manager_role_client, manager_client):
    cat, other = make_category(manager_role_client), make_category(manager_role_client)
    g = make_guest(manager_client)
    gid = manager_client.post(f"{API}/groups", json={"title": "Форма"}).json()["id"]
    create_res(manager_client, cat, "2028-01-02", "2028-01-05", group_id=gid)
    # как форма: оплата явно равна выбранной категории брони
    b = create_res(manager_client, other, "2028-01-02", "2028-01-05", [(g["id"], "main", None, None)],
                   group_id=gid, paid_category_id=other["id"])
    assert b["category_id"] == cat["id"] and _upgrade(manager_client, b["id"]) == ""
    from app.core.db import SessionLocal
    from app.modules.stays.adapters.orm import Reservation
    with SessionLocal() as s:
        assert s.get(Reservation, b["id"]).paid_category_id is None
    # правка: бронь без группы, затем вступление в группу с той же формой
    solo = create_res(manager_client, other, "2028-02-02", "2028-02-05", [(g["id"], "main", None, None)],
                      paid_category_id=other["id"])
    r = _put(manager_client, solo, group_id=gid, category_id=other["id"], paid_category_id=other["id"])
    assert r.status_code == 200, r.text
    assert _upgrade(manager_client, solo["id"]) == ""
