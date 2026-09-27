"""Расширяемые шаблоны заявлений (statement_kinds): встроенные виды, CRUD своих,
привязка к смене словаря, рендер и образец документа по собственному тексту."""
from __future__ import annotations

import io
import os
import tempfile
import unittest

tmp_db = os.path.join(tempfile.gettempdir(), "tt_test_kinds.db")
if os.path.exists(tmp_db):
    os.remove(tmp_db)
os.environ.setdefault("DATABASE_URL", f"sqlite:///{tmp_db}")
os.environ.setdefault("SECRET_KEY", "test-secret-key")

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.deps import local_date  # noqa: E402
from app.main import app, init_db  # noqa: E402

init_db()
client = TestClient(app)
admin = TestClient(app)
emp = TestClient(app)
assert client.post("/api/auth/login", json={"username": "gromova", "password": "demo1234"}).status_code == 200
assert admin.post("/api/auth/login", json={"username": "admin", "password": "demo1234"}).status_code == 200
assert emp.post("/api/auth/login", json={"username": "ivanov", "password": "demo1234"}).status_code == 200

CODE = "mat_aid_smoke"
SHIFT_CODE = "MAT_AID_SM"
TEXT = ("{director}\n{company}\nот {full_name_genitive}\n\nЗАЯВЛЕНИЕ\n\n"
        "Прошу оказать мне материальную помощь в связи с {reason}.\n\n"
        "{today}\t\t____________ / {short_name} /")


def docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


class TestStatementKinds(unittest.TestCase):
    kind_id = None
    shift_id = None

    def test_01_builtin_seeded(self):
        d = client.get("/api/docs/kinds").json()
        codes = {k["code"]: k for k in d["kinds"]}
        for built in ("vacation_paid", "vacation_unpaid", "day_off_hours", "time_off_request"):
            self.assertIn(built, codes)
            self.assertTrue(codes[built]["builtin"])
        # текст встроенного вида унаследован из настроек/дефолтов
        self.assertIn("{full_name_genitive}", codes["vacation_paid"]["text"])
        self.assertIn("full_name", d["placeholders"])

    def test_02_create_validation(self):
        # плохие коды
        for bad in ("1bad", "код", "a", "A" * 41, "with space"):
            r = admin.post("/api/docs/kinds", json={"code": bad, "name": "x"})
            self.assertEqual(r.status_code, 422, f"code={bad!r}: {r.text}")
        # пустое название
        r = admin.post("/api/docs/kinds", json={"code": CODE, "name": "  "})
        self.assertEqual(r.status_code, 422)
        # нормальное создание (код в верхнем регистре нормализуется)
        r = admin.post("/api/docs/kinds", json={
            "code": CODE.upper(), "name": "Заявление на материальную помощь", "text": TEXT})
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        self.assertEqual(d["code"], CODE)
        self.assertFalse(d["builtin"])
        TestStatementKinds.kind_id = d["id"]
        # дубль
        r = admin.post("/api/docs/kinds", json={"code": CODE, "name": "ещё один"})
        self.assertEqual(r.status_code, 409)

    def test_03_shift_binding(self):
        r = admin.post("/api/shift-types", json={
            "code": SHIFT_CODE, "name": "Матпомощь (тест)", "short_code": "МП",
            "kind": "absence", "is_working": False, "counts_as_worked": False,
            "doc_type": CODE})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["doc_type"], CODE)
        TestStatementKinds.shift_id = r.json()["id"]

    def test_04_render_custom_kind(self):
        emp_id = emp.get("/api/auth/me").json()["user"]["employee_id"]
        d1 = local_date().replace(day=1)
        r = admin.post("/api/docs/render", json={
            "type": CODE, "employee_id": emp_id,
            "date_from": d1.isoformat(), "date_to": d1.isoformat(),
            "format": "docx", "reason": "тяжёлыми семейными обстоятельствами"})
        self.assertEqual(r.status_code, 200, r.text)
        text = docx_text(r.content)
        self.assertIn("материальную помощь", text)
        self.assertIn("тяжёлыми семейными обстоятельствами", text)
        self.assertNotIn("{reason}", text)
        self.assertNotIn("{full_name_genitive}", text)   # плейсхолдеры подставлены
        # неизвестный вид
        r = admin.post("/api/docs/render", json={
            "type": "no_such_kind", "employee_id": emp_id,
            "date_from": d1.isoformat(), "date_to": d1.isoformat()})
        self.assertEqual(r.status_code, 422)

    def test_05_sample_custom_kind(self):
        r = admin.get("/api/docs/sample", params={"type": CODE})
        self.assertEqual(r.status_code, 200)
        self.assertIn("материальную помощь", docx_text(r.content))
        # встроенный вид — эталонный образец с богатыми плейсхолдерами
        r = admin.get("/api/docs/sample", params={"type": "vacation_paid"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("{days_word}", docx_text(r.content))

    def test_06_update_kind(self):
        r = admin.put(f"/api/docs/kinds/{self.kind_id}", json={
            "name": "Матпомощь (обновлено)", "text": TEXT.replace("материальную помощь", "выплату"),
            "active": True, "sort_order": 50})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["name"], "Матпомощь (обновлено)")
        d = client.get("/api/docs/kinds").json()
        k = next(x for x in d["kinds"] if x["code"] == CODE)
        self.assertIn("выплату", k["text"])

    def test_07_delete_clears_shift_refs(self):
        r = admin.delete(f"/api/docs/kinds/{self.kind_id}")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["shifts_cleared"], 1)
        shifts = client.get("/api/shift-types", params={"include_archived": True}).json()
        st = next(s for s in shifts if s["code"] == SHIFT_CODE)
        self.assertEqual(st["doc_type"], "")
        # рендер удалённого вида больше невозможен
        r = admin.post("/api/docs/render", json={
            "type": CODE, "employee_id": 4, "date_from": "2026-01-01", "date_to": "2026-01-02"})
        self.assertEqual(r.status_code, 422)

    def test_08_permissions(self):
        self.assertEqual(emp.get("/api/docs/kinds").status_code, 403)
        self.assertEqual(emp.post("/api/docs/kinds",
                                  json={"code": "hack_me", "name": "x"}).status_code, 403)

    def test_99_cleanup(self):
        r = admin.delete(f"/api/shift-types/{self.shift_id}")
        self.assertEqual(r.status_code, 200, r.text)


if __name__ == "__main__":
    unittest.main()
