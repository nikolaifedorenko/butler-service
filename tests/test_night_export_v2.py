"""Ночной отчёт v2 — этап 6: выгрузки DOCX/PDF со службами, недостатками и журналом.

PDF собирается из того же DOCX (app/pdf_render.docx_to_pdf), поэтому содержимое
достаточно проверить один раз — на DOCX.

Что проверяется:
  * у пункта проверки в документе видна ответственная служба (колонка «Служба»);
  * блок «Недостатки смены» содержит «Не ОК» и «Не проверено» (решение владельца:
    неотмеченное включается в отчёт) и статус отправки;
  * блок «Отправки службам» — строки общего журнала (кто, кому, канал, статус);
  * CSV общего журнала выгружается и доступен менеджменту;
  * исправленное после отправки помечается в документе, повторно не отправляется.
"""
from __future__ import annotations

import io

from docx import Document

from tests.test_night_dispatch import (  # noqa: F401 — фикстуры транспорта и каналов
    TECH_IDX,
    _answer,
    _report_item_ids,
    channels,
    fake_smtp,
    fake_telegram,
)


def _docx_bytes(client, rep) -> bytes:
    r = client.get(f"/api/night/reports/{rep.id}/export.docx")
    assert r.status_code == 200, r.text
    return r.content


def _doc_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for row in t.rows for c in row.cells)


def test_v_docx_vidna_sluzhba_u_punkta(employee_client, open_shift):
    with open_shift() as rep:
        text = _doc_text(_docx_bytes(employee_client, rep))
        assert "Пункт проверки" in text and "Служба" in text
        assert "Техническая служба" in text and "Хозяйственная служба" in text
        assert "Семейные виллы" in text          # группировка по зонам (этап 1) сохранена


def _counts(employee_client, rep) -> dict:
    r = employee_client.get(f"/api/night/reports/{rep.id}/defects")
    assert r.status_code == 200, r.text
    return r.json()["counts"]


def _journal(employee_client, rep) -> list[dict]:
    r = employee_client.get(f"/api/night/reports/{rep.id}/dispatches")
    assert r.status_code == 200, r.text
    return r.json()["dispatches"]


def test_blok_nedostatkov_i_otpravok_v_docx(employee_client, admin_client, open_shift,
                                            channels, fake_telegram, fake_smtp):
    # отчёт за сегодняшнюю смену может быть общим с другими тестами прогона,
    # поэтому сверяемся с фактическими счётчиками, а не с абсолютными числами
    with open_shift() as rep:
        items = _report_item_ids(employee_client, rep, "Вилла 4003")
        item_id = items[TECH_IDX]["id"]
        note = "Свет не выключен в холле (проверка DOCX)"
        _answer(employee_client, item_id, "bad", note)

        counts = _counts(employee_client, rep)
        journal_before = _journal(employee_client, rep)
        text = _doc_text(_docx_bytes(employee_client, rep))
        assert "Недостатки смены" in text
        assert f"«Не ОК» {counts['bad']}" in text, counts
        assert f"«Не проверено» {counts['unchecked']}" in text       # неотмеченное включено в отчёт
        assert note in text and "не отправлено" in text
        assert "Отправки службам" in text
        if not journal_before:
            assert "В эту смену ничего не отправлялось" in text

        r = employee_client.post(f"/api/night/reports/{rep.id}/dispatch",
                                 json={"mode": "selected", "item_ids": [item_id]})
        assert r.status_code == 200, r.text
        assert r.json()["sent"] >= 1, r.json()

        text = _doc_text(_docx_bytes(employee_client, rep))
        assert note in text and "отправлено" in text
        # журнал отправок: служба, канал, получатель, статус
        assert "чат Telegram" in text and str(channels["tech"]) in text
        assert "Техническая служба" in text
        assert len(_journal(employee_client, rep)) > len(journal_before)

        # исправили после отправки: в документе видно, повторно не отправляем
        _answer(employee_client, item_id, "ok")
        text = _doc_text(_docx_bytes(employee_client, rep))
        assert "исправлено после отправки" in text
        r2 = employee_client.post(f"/api/night/reports/{rep.id}/dispatch",
                                  json={"mode": "selected", "item_ids": [item_id]})
        assert r2.status_code == 422, r.text


def test_csv_obshchego_zhurnala(admin_client, employee_client, manager_client, open_shift,
                                channels, fake_telegram, fake_smtp):
    # батлеру общий журнал приложения не доступен (он видит журнал своей смены)
    assert employee_client.get("/api/notify/dispatches/export.csv").status_code == 403

    with open_shift() as rep:
        items = _report_item_ids(employee_client, rep, "Вилла 4003")
        item_id = items[TECH_IDX]["id"]
        _answer(employee_client, item_id, "bad", "не горит")
        assert employee_client.post(f"/api/night/reports/{rep.id}/dispatch",
                                    json={"mode": "selected", "item_ids": [item_id]}).status_code == 200

    r = admin_client.get("/api/notify/dispatches/export.csv")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/csv")
    body = r.content.decode("utf-8-sig")
    head, *rows = [line for line in body.splitlines() if line.strip()]
    assert head.startswith("Время;Раздел;Кто отправил;Получатель;Канал")
    night = [line for line in rows if ";night;" in line]
    assert night and any("Техническая служба" in line for line in night)
    assert any("не горит" in line for line in night)          # состав отправки снимком
    assert any("отправлено" in line for line in night)
    # супервайзеру выгрузка журнала доступна (require_manager)
    assert manager_client.get("/api/notify/dispatches/export.csv").status_code == 200
